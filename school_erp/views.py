import json
import datetime
import random
import csv
import io
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.http import HttpResponse
from django.db.models import Sum, Count, Q
from django.utils import timezone
from django.core.management import call_command
from functools import wraps
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.models import User
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from django.http import HttpResponse, FileResponse, JsonResponse
from .models import (
    SchoolProfile, AcademicSession, ClassLevel, Section, Subject,
    Teacher, Student, Attendance, Timetable, FeeCategory, ClassFeeMaster,
    FeeInvoice, FeePayment, FeeReceipt, FeeReceiptItem, Exam, ExamResult, Notice,
    StudentSessionEnrollment, UserProfile, FeeHead, FeeStructure, FeeStructureItem,
    FeeInstalmentSchedule, TransportRoute, HostelFacility, StudentFacilityAllocation,
    StudentConcession, FeeRefund, StudentFeeLedger, FeeReminderLog,
    Role, RolePermission, AuditLog, Branch, BackupConfiguration, BackupLog,
    SecurityConfiguration, LoginAttemptLog
)
from .forms import (
    StudentForm, TeacherForm, FeePaymentForm, FeeInvoiceForm, NoticeForm,
    SchoolSettingsForm, ExamForm, TimetableForm, FeeHeadForm, FeeRefundForm,
    StudentConcessionForm, TransportRouteForm, HostelFacilityForm,
    RoleForm, AuditApprovalForm, UserCreateUpdateForm, BranchForm,
    BackupConfigurationForm, SecurityConfigurationForm
)
from .backup_utils import generate_system_backup, verify_restore_archive, execute_test_environment_restore
from .student_utils import (
    find_duplicate_student, normalize_student_name, normalize_parent_name,
    parse_date_of_birth, get_row_field, generate_student_admission_number,
    generate_unique_school_code
)
from .context_processors import get_active_session, get_current_school


def get_school_context(request=None):
    if request:
        profile = get_current_school(request)
    else:
        profile = SchoolProfile.objects.first()
    return {'school': profile}


# ==============================================================================
# MENU PERMISSIONS & ACCESS CONTROL DECORATORS
# ==============================================================================

def get_client_ip(request):
    if not request:
        return None
    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded_for:
        return x_forwarded_for.split(',')[0].strip()
    return request.META.get('REMOTE_ADDR')


def log_audit(request, action, module, object_type, object_id='', object_repr='',
              description='', reason='', changes=None, requires_approval=False, status='SUCCESS'):
    """
    Records an immutable audit trail entry for security, tracking, and compliance.
    """
    try:
        user = request.user if (request and getattr(request, 'user', None) and request.user.is_authenticated) else None
        school = get_current_school(request) if request else None
        user_role = ''
        if user:
            profile = getattr(user, 'profile', None)
            if profile:
                user_role = (profile.role_obj.name if profile.role_obj else profile.role) or ''
            elif user.is_superuser:
                user_role = 'Super Admin'
        ip_addr = get_client_ip(request)

        return AuditLog.objects.create(
            school=school,
            user=user,
            user_role=user_role,
            action=action,
            module=module,
            object_type=object_type,
            object_id=str(object_id) if object_id else '',
            object_repr=str(object_repr)[:250] if object_repr else '',
            description=description,
            changes_json=changes,
            requires_approval=requires_approval,
            status=status,
            reason=reason or '',
            ip_address=ip_addr
        )
    except Exception as e:
        print(f"Error logging audit: {e}")
        return None


def ensure_roles_and_permissions():
    """
    Seeds and ensures the 8 standard roles and their 10 module permissions exist.
    """
    standard_roles = [
        ('Super Admin', 'SUPER_ADMIN', 'Platform & School Super Administrator with full privileges', True),
        ('School Admin', 'SCHOOL_ADMIN', 'School Head Administrator managing school operations', True),
        ('Principal', 'PRINCIPAL', 'Academic head, approves marks, results & general reports', True),
        ('Teacher', 'TEACHER', 'Classroom faculty, marks entry & attendance for assigned classes', True),
        ('Accountant', 'ACCOUNTANT', 'Fee collection, ledgers, concessions, refunds & finance', True),
        ('HR', 'HR', 'Staff records, teacher onboarding & human resources', True),
        ('Parent', 'PARENT', 'Guardian viewing records, fees & marks for linked children', True),
        ('Student', 'STUDENT', 'Enrolled student viewing own timetable, results & fee status', True),
    ]

    modules = [
        ('STUDENTS', 'Students & SIS'),
        ('FEES', 'Fees & Accounts'),
        ('ATTENDANCE', 'Attendance'),
        ('EXAMS', 'Examinations & Marks'),
        ('TIMETABLE', 'Timetable & Routines'),
        ('SESSIONS', 'Academic Sessions & Promotion'),
        ('STAFF_HR', 'Teachers & Staff / HR'),
        ('USERS_ROLES', 'User Management & Roles'),
        ('AUDIT_LOGS', 'Audit Trail & Logs'),
        ('SETTINGS', 'School Profile & Settings'),
    ]

    role_objs = {}
    for name, code, desc, is_sys in standard_roles:
        r, _ = Role.objects.get_or_create(
            code=code,
            defaults={'name': name, 'description': desc, 'is_system_role': is_sys}
        )
        role_objs[code] = r

    # Default permission policies per role:
    # Tuple: (can_view, can_add, can_edit, can_delete, can_approve, can_export)
    perm_matrix = {
        'SUPER_ADMIN': {m[0]: (True, True, True, True, True, True) for m in modules},
        'SCHOOL_ADMIN': {m[0]: (True, True, True, True, True, True) for m in modules},
        'PRINCIPAL': {
            'STUDENTS': (True, True, True, False, True, True),
            'FEES': (True, False, False, False, True, True),
            'ATTENDANCE': (True, True, True, False, True, True),
            'EXAMS': (True, True, True, False, True, True),
            'TIMETABLE': (True, True, True, False, True, True),
            'SESSIONS': (True, True, True, False, True, True),
            'STAFF_HR': (True, True, True, False, True, True),
            'USERS_ROLES': (True, False, False, False, False, True),
            'AUDIT_LOGS': (True, False, False, False, True, True),
            'SETTINGS': (True, False, True, False, True, True),
        },
        'TEACHER': {
            'STUDENTS': (True, False, False, False, False, False),
            'FEES': (False, False, False, False, False, False),
            'ATTENDANCE': (True, True, True, False, False, False),
            'EXAMS': (True, True, True, False, False, False),
            'TIMETABLE': (True, False, False, False, False, False),
            'SESSIONS': (False, False, False, False, False, False),
            'STAFF_HR': (False, False, False, False, False, False),
            'USERS_ROLES': (False, False, False, False, False, False),
            'AUDIT_LOGS': (False, False, False, False, False, False),
            'SETTINGS': (False, False, False, False, False, False),
        },
        'ACCOUNTANT': {
            'STUDENTS': (True, False, False, False, False, True),
            'FEES': (True, True, True, False, False, True),
            'ATTENDANCE': (False, False, False, False, False, False),
            'EXAMS': (False, False, False, False, False, False),
            'TIMETABLE': (False, False, False, False, False, False),
            'SESSIONS': (True, False, False, False, False, False),
            'STAFF_HR': (False, False, False, False, False, False),
            'USERS_ROLES': (False, False, False, False, False, False),
            'AUDIT_LOGS': (True, False, False, False, False, True),
            'SETTINGS': (False, False, False, False, False, False),
        },
        'HR': {
            'STUDENTS': (False, False, False, False, False, False),
            'FEES': (False, False, False, False, False, False),
            'ATTENDANCE': (True, False, False, False, False, False),
            'EXAMS': (False, False, False, False, False, False),
            'TIMETABLE': (True, False, False, False, False, False),
            'SESSIONS': (False, False, False, False, False, False),
            'STAFF_HR': (True, True, True, False, False, True),
            'USERS_ROLES': (True, False, False, False, False, False),
            'AUDIT_LOGS': (True, False, False, False, False, False),
            'SETTINGS': (False, False, False, False, False, False),
        },
        'PARENT': {
            'STUDENTS': (True, False, False, False, False, False),
            'FEES': (True, False, False, False, False, False),
            'ATTENDANCE': (True, False, False, False, False, False),
            'EXAMS': (True, False, False, False, False, False),
            'TIMETABLE': (True, False, False, False, False, False),
            'SESSIONS': (False, False, False, False, False, False),
            'STAFF_HR': (False, False, False, False, False, False),
            'USERS_ROLES': (False, False, False, False, False, False),
            'AUDIT_LOGS': (False, False, False, False, False, False),
            'SETTINGS': (False, False, False, False, False, False),
        },
        'STUDENT': {
            'STUDENTS': (True, False, False, False, False, False),
            'FEES': (True, False, False, False, False, False),
            'ATTENDANCE': (True, False, False, False, False, False),
            'EXAMS': (True, False, False, False, False, False),
            'TIMETABLE': (True, False, False, False, False, False),
            'SESSIONS': (False, False, False, False, False, False),
            'STAFF_HR': (False, False, False, False, False, False),
            'USERS_ROLES': (False, False, False, False, False, False),
            'AUDIT_LOGS': (False, False, False, False, False, False),
            'SETTINGS': (False, False, False, False, False, False),
        },
    }

    for r_code, m_perms in perm_matrix.items():
        role = role_objs.get(r_code)
        if not role:
            continue
        for m_code, m_name in modules:
            p_tuple = m_perms.get(m_code, (False, False, False, False, False, False))
            RolePermission.objects.get_or_create(
                role=role,
                module_code=m_code,
                defaults={
                    'module_name': m_name,
                    'can_view': p_tuple[0],
                    'can_add': p_tuple[1],
                    'can_edit': p_tuple[2],
                    'can_delete': p_tuple[3],
                    'can_approve': p_tuple[4],
                    'can_export': p_tuple[5],
                }
            )


def has_module_perm(user, module_code, action='view'):
    """
    Check if user has specific action permission for a module.
    Actions: 'view', 'add', 'edit', 'delete', 'approve', 'export'
    """
    if not user.is_authenticated:
        return False
    if user.is_superuser:
        return True
    profile = getattr(user, 'profile', None)
    if not profile:
        return False

    role = profile.role_obj
    if not role and profile.role:
        role = Role.objects.filter(Q(name__iexact=profile.role) | Q(code__iexact=profile.role)).first()
        if role:
            profile.role_obj = role
            profile.save(update_fields=['role_obj'])

    if role:
        if role.code in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
            return True
        perm = RolePermission.objects.filter(role=role, module_code=module_code).first()
        if perm:
            action_field = f"can_{action}"
            return getattr(perm, action_field, False)

    # Fallback to boolean flags on UserProfile
    legacy_map = {
        'STUDENTS': 'can_manage_students',
        'FEES': 'can_manage_fees',
        'ATTENDANCE': 'can_manage_attendance',
        'EXAMS': 'can_manage_exams',
        'TIMETABLE': 'can_manage_timetable',
        'SESSIONS': 'can_manage_sessions',
        'USERS_ROLES': 'can_manage_users',
    }
    flag_name = legacy_map.get(module_code)
    if flag_name:
        return getattr(profile, flag_name, False)
    return False


LEGACY_PERM_MODULE_MAP = {
    'can_manage_students': ('STUDENTS', 'view'),
    'can_manage_fees': ('FEES', 'view'),
    'can_manage_attendance': ('ATTENDANCE', 'view'),
    'can_manage_exams': ('EXAMS', 'view'),
    'can_manage_timetable': ('TIMETABLE', 'view'),
    'can_manage_classes': ('SETTINGS', 'view'),
    'can_manage_sessions': ('SESSIONS', 'view'),
    'can_manage_users': ('USERS_ROLES', 'view'),
}

def check_menu_permission(request, perm_name):
    """Check if current user has permission for a specific menu item."""
    user = request.user
    if not user.is_authenticated:
        return True
    if user.is_superuser:
        return True
    profile = getattr(user, 'profile', None)
    if not profile:
        return False

    # Check via modern RolePermission matrix first
    if perm_name in LEGACY_PERM_MODULE_MAP:
        mod_code, act = LEGACY_PERM_MODULE_MAP[perm_name]
        return has_module_perm(user, mod_code, act)

    return getattr(profile, perm_name, False)


def require_menu_permission(perm_name, perm_label="this section"):
    """View decorator ensuring user has required menu permission."""
    def decorator(view_func):
        @wraps(view_func)
        def _wrapped(request, *args, **kwargs):
            if request.user.is_authenticated and not check_menu_permission(request, perm_name):
                messages.error(
                    request,
                    f"Access Restricted: You do not have permission to access '{perm_label}'. Please contact your administrator."
                )
                return redirect('dashboard')
            return view_func(request, *args, **kwargs)
        return _wrapped
    return decorator


def is_teacher_user(user):
    if not user.is_authenticated:
        return False
    if user.is_superuser:
        return False
    profile = getattr(user, 'profile', None)
    if not profile:
        return False
    role_code = (getattr(profile.role_obj, 'code', None) or '').upper()
    role_name = (profile.role or '').upper()
    return 'TEACHER' in role_code or 'TEACHER' in role_name or bool(profile.assigned_teacher)


def is_parent_user(user):
    if not user.is_authenticated:
        return False
    if user.is_superuser:
        return False
    profile = getattr(user, 'profile', None)
    if not profile:
        return False
    role_code = (getattr(profile.role_obj, 'code', None) or '').upper()
    role_name = (profile.role or '').upper()
    return 'PARENT' in role_code or 'PARENT' in role_name


def get_teacher_assigned_classes(user):
    """
    Returns (class_queryset, subject_queryset) assigned to teacher user.
    """
    profile = getattr(user, 'profile', None)
    if not profile:
        return ClassLevel.objects.none(), Subject.objects.none()

    classes = set(profile.assigned_classes.all())
    subjects = set()

    teacher = profile.assigned_teacher
    if not teacher and hasattr(user, 'teacher_profile'):
        teacher = user.teacher_profile
    if not teacher:
        teacher = Teacher.objects.filter(Q(user=user) | Q(email=user.email)).first()

    if teacher:
        ct_sections = Section.objects.filter(class_teacher=teacher)
        for sec in ct_sections:
            classes.add(sec.class_level)

        tt_entries = Timetable.objects.filter(teacher=teacher)
        for tt in tt_entries:
            classes.add(tt.class_level)
            if tt.subject:
                subjects.add(tt.subject)

        for sub in Subject.objects.filter(teachers=teacher):
            subjects.add(sub)

    class_ids = [c.id for c in classes]
    class_qs = ClassLevel.objects.filter(id__in=class_ids).order_by('order', 'name')
    if subjects:
        subject_ids = [s.id for s in subjects]
        subject_qs = Subject.objects.filter(id__in=subject_ids).order_by('name')
    else:
        subject_qs = Subject.objects.all()

    return class_qs, subject_qs


def get_parent_allowed_students(user):
    """
    Returns queryset of students that a parent user is authorized to view.
    """
    profile = getattr(user, 'profile', None)
    if not profile:
        return Student.objects.none()
    student_ids = set(profile.linked_students.values_list('id', flat=True))
    if profile.phone:
        by_phone = Student.objects.filter(parent_phone=profile.phone).values_list('id', flat=True)
        student_ids.update(by_phone)
    if user.email:
        by_email = Student.objects.filter(parent_email=user.email).values_list('id', flat=True)
        student_ids.update(by_email)
    return Student.objects.filter(id__in=student_ids)



def dashboard_view(request):
    if not request.user.is_authenticated:
        return redirect('user_login')
    ctx = get_school_context(request)
    school = ctx['school']
    active_session = get_active_session(request, school=school)
    
    # Filter students for active session & current school
    students_in_session = Student.objects.filter(school=school, is_active=True)
    if active_session:
        students_in_session = students_in_session.filter(
            Q(academic_session=active_session) | Q(session_enrollments__academic_session=active_session)
        ).distinct()
        
    total_students = students_in_session.count()
    total_teachers = Teacher.objects.filter(school=school, is_active=True).count()
    total_classes = ClassLevel.objects.filter(school=school).count()
    
    # Fees summary for active session & current school
    invoices = FeeInvoice.objects.filter(school=school)
    if active_session:
        invoices = invoices.filter(academic_session=active_session)
    total_invoiced = sum((inv.total_payable for inv in invoices), Decimal('0.00'))
    if active_session:
        total_collected = FeePayment.objects.filter(invoice__school=school, invoice__academic_session=active_session).aggregate(total=Sum('amount_paid'))['total'] or Decimal('0.00')
    else:
        total_collected = FeePayment.objects.filter(invoice__school=school).aggregate(total=Sum('amount_paid'))['total'] or Decimal('0.00')
    pending_fees = max(Decimal('0.00'), total_invoiced - total_collected)
    
    # Today's attendance
    today = timezone.now().date()
    today_attendance = Attendance.objects.filter(student__school=school, date=today)
    att_present = today_attendance.filter(status='Present').count()
    att_absent = today_attendance.filter(status='Absent').count()
    att_total = today_attendance.count()
    att_rate = round((att_present / att_total * 100), 1) if att_total > 0 else 0
    
    recent_students = students_in_session.order_by('-created_at')[:6]
    recent_payments = FeePayment.objects.filter(invoice__school=school).order_by('-payment_date', '-created_at')
    if active_session:
        recent_payments = recent_payments.filter(invoice__academic_session=active_session)
    recent_payments = recent_payments[:5]
    notices = Notice.objects.filter(school=school).order_by('-is_urgent', '-created_at')[:4]
    classes = ClassLevel.objects.filter(school=school).prefetch_related('sections').order_by('order', 'name')
    
    ctx.update({
        'total_students': total_students,
        'total_teachers': total_teachers,
        'total_classes': total_classes,
        'total_invoiced': total_invoiced,
        'total_collected': total_collected,
        'pending_fees': pending_fees,
        'today': today,
        'att_present': att_present,
        'att_absent': att_absent,
        'att_total': att_total,
        'att_rate': att_rate,
        'recent_students': recent_students,
        'recent_payments': recent_payments,
        'notices': notices,
        'classes': classes,
    })
    return render(request, 'school_erp/dashboard.html', ctx)


# ------------------- STUDENTS -------------------

@require_menu_permission('can_manage_students', 'Student Directory')
def student_list_view(request):
    ctx = get_school_context(request)
    school = ctx.get('school')
    active_session = get_active_session(request)

    if school:
        students_qs = Student.objects.filter(school=school, is_active=True).select_related('class_level', 'section', 'academic_session')
    else:
        students_qs = Student.objects.none()
    if active_session:
        students_qs = students_qs.filter(
            Q(academic_session=active_session) | Q(session_enrollments__academic_session=active_session)
        ).distinct()

    # Parent Scoped Access Restriction Enforcement
    if is_parent_user(request.user):
        allowed_students = get_parent_allowed_students(request.user)
        students_qs = students_qs.filter(id__in=allowed_students.values_list('id', flat=True))

    students_qs = students_qs.order_by('class_level__order', 'roll_number', 'first_name')
    
    search_query = request.GET.get('q', '').strip()
    class_id = request.GET.get('class_id', '')
    section_id = request.GET.get('section_id', '')
    
    gender = request.GET.get('gender', '').strip()
    fee_filter = request.GET.get('fee_filter', '').strip()

    if search_query:
        students_qs = students_qs.filter(
            Q(first_name__icontains=search_query) |
            Q(last_name__icontains=search_query) |
            Q(admission_number__icontains=search_query) |
            Q(parent_name__icontains=search_query) |
            Q(parent_phone__icontains=search_query)
        )
    if class_id:
        students_qs = students_qs.filter(class_level_id=class_id)
    if section_id:
        students_qs = students_qs.filter(section_id=section_id)
    if gender in ['Male', 'Female']:
        students_qs = students_qs.filter(gender=gender)
    if fee_filter == 'due':
        due_student_ids = FeeInvoice.objects.filter(status__in=['Unpaid', 'Partially Paid', 'Overdue']).values_list('student_id', flat=True).distinct()
        students_qs = students_qs.filter(id__in=due_student_ids)
    elif fee_filter == 'cleared':
        due_student_ids = FeeInvoice.objects.filter(status__in=['Unpaid', 'Partially Paid', 'Overdue']).values_list('student_id', flat=True).distinct()
        students_qs = students_qs.exclude(id__in=due_student_ids)
        
    if school:
        classes = ClassLevel.objects.filter(school=school).order_by('order', 'name')
        sections = Section.objects.filter(class_level__school=school)
    else:
        classes = ClassLevel.objects.all().order_by('order', 'name')
        sections = Section.objects.all()
    total_count = students_qs.count()
    
    # 20 students per page
    paginator = Paginator(students_qs, 20)
    page_number = request.GET.get('page', 1)
    try:
        page_obj = paginator.get_page(page_number)
    except (EmptyPage, PageNotAnInteger):
        page_obj = paginator.get_page(1)
        
    ctx.update({
        'students': page_obj,
        'page_obj': page_obj,
        'paginator': paginator,
        'is_paginated': page_obj.has_other_pages(),
        'classes': classes,
        'sections': sections,
        'search_query': search_query,
        'selected_class': class_id,
        'selected_section': section_id,
        'selected_gender': gender,
        'selected_fee_filter': fee_filter,
        'total_count': total_count,
        'is_parent_restricted': is_parent_user(request.user),
    })
    return render(request, 'school_erp/student_list.html', ctx)


def student_detail_view(request, pk):
    ctx = get_school_context(request)
    student = get_object_or_404(Student.objects.select_related('class_level', 'section'), pk=pk)

    # Parent Scoped Access Restriction Enforcement
    if is_parent_user(request.user):
        allowed_students = get_parent_allowed_students(request.user)
        if not allowed_students.filter(id=student.id).exists():
            messages.error(
                request,
                "Access Restricted: As a Parent, you can only view profile and academic records for your linked child."
            )
            return redirect('dashboard')
    
    attendance_records = student.attendance_records.order_by('-date')[:30]
    invoices = student.invoices.prefetch_related('payments').all()
    exam_results = student.exam_results.select_related('exam', 'subject').all()
    
    ctx.update({
        'student': student,
        'attendance_records': attendance_records,
        'invoices': invoices,
        'exam_results': exam_results,
    })
    return render(request, 'school_erp/student_detail.html', ctx)


@require_menu_permission('can_manage_students', 'New Admission')
def student_create_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    active_session = get_active_session(request, school=school)

    if request.method == 'POST':
        form = StudentForm(request.POST, request.FILES, school=school)
        if form.is_valid():
            student = form.save(commit=False)
            student.school = school
            student.academic_session = active_session

            # Generate or enforce standardized admission number: [School 3 letters]-[Year]-[Register Sequence]
            adm_no = request.POST.get('admission_number', '').strip()
            if not adm_no or Student.objects.filter(admission_number=adm_no).exclude(pk=student.pk).exists():
                adm_no = generate_student_admission_number(school, active_session)
            student.admission_number = adm_no
            student.save()

            # Record enrollment in active session
            StudentSessionEnrollment.objects.get_or_create(
                student=student,
                academic_session=active_session,
                defaults={
                    'class_level': student.class_level,
                    'section': student.section,
                    'roll_number': student.roll_number,
                    'status': 'Active'
                }
            )
            
            # Automatically create Admission Fee invoice from Class Fee Master
            adm_amt = Decimal('5000.00')
            if student.class_level:
                try:
                    fm = student.class_level.fee_master
                    adm_amt = fm.admission_fee
                except Exception:
                    adm_amt = Decimal('5000.00')

            adm_cat = FeeCategory.objects.filter(school=school, name__icontains='Admission').first()
            if not adm_cat:
                adm_cat, _ = FeeCategory.objects.get_or_create(
                    school=school,
                    name="Admission Fee",
                    defaults={'default_amount': adm_amt, 'description': "One-time student admission and registration fee"}
                )

            inv_adm_no = f"INV-{timezone.now().year}-ADM{student.id:04d}"
            FeeInvoice.objects.create(
                school=school,
                invoice_number=inv_adm_no,
                student=student,
                fee_category=adm_cat,
                academic_session=active_session,
                amount=adm_amt,
                due_date=timezone.now().date() + datetime.timedelta(days=15),
                status='Unpaid'
            )

            # Automatically create initial tuition fee invoice
            tuition_cat = FeeCategory.objects.filter(school=school, name__icontains='Tuition').first() or FeeCategory.objects.filter(school=school, name__icontains='Monthly').first()
            if tuition_cat:
                inv_tui_no = f"INV-{timezone.now().year}-TUI{student.id:04d}"
                FeeInvoice.objects.create(
                    school=school,
                    invoice_number=inv_tui_no,
                    student=student,
                    fee_category=tuition_cat,
                    amount=tuition_cat.default_amount,
                    due_date=timezone.now().date() + datetime.timedelta(days=30),
                    status='Unpaid'
                )

            messages.success(request, f"Student '{student.full_name}' admitted successfully! Admission No: {student.admission_number}. Admission Fee ₹{adm_amt:,.0f} billed.")
            return redirect('fee_counter') + f"?student_id={student.id}"
        else:
            has_error_msg = False
            for err in form.non_field_errors():
                messages.error(request, err)
                has_error_msg = True
            if not has_error_msg:
                for field_name, errs in form.errors.items():
                    for err in errs:
                        messages.error(request, f"{field_name.replace('_', ' ').title()}: {err}")
    else:
        # Generate standardized admission number: [School 3 letters]-[Year]-[Register Sequence]
        default_adm = generate_student_admission_number(school, active_session)
        form = StudentForm(initial={'admission_number': default_adm, 'admission_date': timezone.now().date()}, school=school)
        
    ctx.update({'form': form, 'title': 'New Student Admission'})
    return render(request, 'school_erp/student_form.html', ctx)


def student_edit_view(request, pk):
    ctx = get_school_context(request)
    school = ctx['school']
    student = get_object_or_404(Student, pk=pk, school=school)
    if request.method == 'POST':
        form = StudentForm(request.POST, request.FILES, instance=student, school=school)
        if form.is_valid():
            student = form.save(commit=False)
            student.school = school
            student.save()
            messages.success(request, f"Student '{student.full_name}' details updated successfully!")
            return redirect('student_detail', pk=student.pk)
        else:
            has_error_msg = False
            for err in form.non_field_errors():
                messages.error(request, err)
                has_error_msg = True
            if not has_error_msg:
                for field_name, errs in form.errors.items():
                    for err in errs:
                        messages.error(request, f"{field_name.replace('_', ' ').title()}: {err}")
    else:
        form = StudentForm(instance=student, school=school)
        
    ctx.update({'form': form, 'student': student, 'title': f'Edit Student: {student.full_name}'})
    return render(request, 'school_erp/student_form.html', ctx)


def export_student_excel_template_view(request):
    """
    Exports a styled Excel (.xlsx) or CSV template for bulk student registration.
    Allows pre-filling class if ?class_id= is passed.
    """
    format_type = request.GET.get('format', 'xlsx').lower()
    class_id = request.GET.get('class_id')
    
    target_class = None
    if class_id:
        target_class = ClassLevel.objects.filter(id=class_id).first()
        
    class_name_sample = target_class.name if target_class else "Class 1"
    
    headers = [
        "Roll_No",
        "Class_Name",
        "Section_Name",
        "First_Name",
        "Last_Name",
        "Gender",
        "Date_of_Birth",
        "Blood_Group",
        "Parent_Name",
        "Parent_Relation",
        "Parent_Phone",
        "Student_Aadhaar",
        "Parent_Aadhaar",
        "Residential_Address",
        "Emergency_Contact"
    ]
    
    sample_rows = [
        [
            1, class_name_sample, "A", "Aarav", "Sharma", "Male",
            "2018-04-12", "B+", "Rajesh Sharma", "Father", "9812345670",
            "456789123456", "987654321012", "Sector 62, Noida, UP", "9811223344"
        ],
        [
            2, class_name_sample, "A", "Ananya", "Verma", "Female",
            "2018-08-25", "O+", "Sunil Verma", "Father", "9876501234",
            "345678912345", "876543210987", "Indirapuram, Ghaziabad", "9876543210"
        ]
    ]
    
    if format_type == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="SRS_Student_Upload_Template_{class_name_sample.replace(" ", "_")}.csv"'
        writer = csv.writer(response)
        writer.writerow(headers)
        for row in sample_rows:
            writer.writerow(row)
        return response

    # Excel (.xlsx) Generation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student_Import"
    
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    sample_font = Font(name="Arial", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Write Headers
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        
    # Write Sample Rows
    for row_data in sample_rows:
        ws.append(row_data)
        curr_row = ws.max_row
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.font = sample_font
            cell.border = thin_border
            if col_idx in [1, 3, 6, 7, 8, 11, 12, 13]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Add instructions block at bottom
    ws.append([])
    ws.append(["--- INSTRUCTIONS ---"])
    instructions = [
        "1. Admission Number does not need to be filled in Excel - system will auto-generate it.",
        "2. First_Name, Class_Name, and Parent_Phone are required.",
        "3. Student_Aadhaar & Parent_Aadhaar: Optional 12-digit Aadhaar Card numbers.",
        "4. Date_of_Birth format: YYYY-MM-DD (e.g. 2018-05-15) or DD/MM/YYYY.",
        "5. Section_Name: If not specified, default Section 'A' will be assigned.",
        "6. If the Class or Section does not already exist, system will automatically create it with default Fee Master!"
    ]
    for inst in instructions:
        ws.append([inst])
        ws.cell(row=ws.max_row, column=1).font = Font(name="Arial", size=9, italic=True, color="64748B")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        if col[0].column > len(headers):
            continue
        for cell in col[:10]:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"SRS_Student_Bulk_Upload_Template_{class_name_sample.replace(' ', '_')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@require_menu_permission('can_manage_students', 'Bulk Student Upload')
def student_bulk_upload_view(request):
    """
    Handles bulk student upload via Excel (.xlsx, .xls) or CSV.
    Validates duplicate entries based on Student Name, Father Name, and Date of Birth.
    If a student already exists (in DB or earlier in the file), skips insertion and reports duplicate details.
    """
    ctx = get_school_context(request)
    school = ctx['school']
    active_session = get_active_session(request, school=school)
    classes = ClassLevel.objects.filter(school=school).order_by('order', 'name')
    
    if request.method == 'POST':
        uploaded_file = request.FILES.get('excel_file')
        override_class_id = request.POST.get('target_class_id')
        default_section_name = request.POST.get('default_section', 'A').strip().upper() or 'A'
        
        if not uploaded_file:
            messages.error(request, "Please select an Excel (.xlsx / .xls) or CSV file!")
            return redirect('student_bulk_upload')
            
        file_name = uploaded_file.name.lower()
        rows_data = []
        
        try:
            if file_name.endswith('.csv'):
                decoded_file = uploaded_file.read().decode('utf-8-sig', errors='replace')
                io_string = io.StringIO(decoded_file)
                reader = csv.DictReader(io_string)
                for row in reader:
                    clean_row = {str(k).strip().lower(): (v.strip() if v else '') for k, v in row.items() if k}
                    rows_data.append(clean_row)
            elif file_name.endswith(('.xlsx', '.xls')):
                wb = openpyxl.load_workbook(uploaded_file, data_only=True)
                ws = wb.active
                raw_rows = list(ws.iter_rows(values_only=True))
                if not raw_rows:
                    messages.error(request, "The uploaded file is empty.")
                    return redirect('student_bulk_upload')
                    
                header_row = [str(c).strip().lower() if c is not None else '' for c in raw_rows[0]]
                for r in raw_rows[1:]:
                    if not any(r):
                        continue
                    if str(r[0] or '').startswith('---'):
                        break
                    row_dict = {}
                    for idx, h in enumerate(header_row):
                        if h and idx < len(r):
                            val = r[idx]
                            row_dict[h] = val if val is not None else ''
                    rows_data.append(row_dict)
            else:
                messages.error(request, "Only .xlsx, .xls, or .csv files are supported!")
                return redirect('student_bulk_upload')
        except Exception as e:
            messages.error(request, f"Error reading file: {str(e)}")
            return redirect('student_bulk_upload')
            
        if not rows_data:
            messages.error(request, "No valid student data found in the file.")
            return redirect('student_bulk_upload')
            
        created_count = 0
        duplicate_count = 0
        error_count = 0
        
        seen_in_file = {}
        duplicate_rows = []
        created_students = []
        
        override_class = None
        if override_class_id and override_class_id.isdigit():
            override_class = ClassLevel.objects.filter(id=int(override_class_id)).first()
            
        for idx, row in enumerate(rows_data, start=2):
            first_name = get_row_field(row, [
                'first_name', 'first name', 'name', 'student_name', 'student name',
                'candidate_name', 'candidate name'
            ])
            if not first_name or first_name.startswith('---'):
                continue
                
            last_name = get_row_field(row, ['last_name', 'last name', 'surname'])
            
            # If last_name is empty and first_name has multiple words, split
            if not last_name and ' ' in first_name.strip():
                parts = first_name.strip().split(' ', 1)
                first_name = parts[0]
                last_name = parts[1]
                
            parent_name = get_row_field(row, [
                'parent_name', 'parent name', 'father_name', 'father name',
                "father's name", 'fathers name', 'guardian_name', 'guardian name'
            ])
            if not parent_name:
                parent_name = f"Mr. {last_name or first_name}"
                
            # Date of Birth
            dob_raw = get_row_field(row, ['date_of_birth', 'date of birth', 'dob', 'birth_date', 'birth date'])
            dob_val = parse_date_of_birth(dob_raw)
            if not dob_val:
                dob_val = datetime.date(2015, 1, 1)
                
            # Determine Class
            class_name = get_row_field(row, ['class_name', 'class name', 'class', 'grade', 'standard'])
            cl = None
            if override_class:
                cl = override_class
            elif class_name:
                cl = ClassLevel.objects.filter(school=school).filter(Q(name__iexact=class_name) | Q(code__iexact=class_name)).first()
                if not cl:
                    order_val = (ClassLevel.objects.filter(school=school).order_by('-order').values_list('order', flat=True).first() or 0) + 1
                    clean_code = "".join(ch for ch in class_name if ch.isalnum()).upper()[:8] or f"C{order_val}"
                    cl = ClassLevel.objects.create(school=school, name=class_name, code=clean_code, order=order_val)
                    ClassFeeMaster.objects.get_or_create(class_level=cl)
            else:
                cl = ClassLevel.objects.filter(school=school).first()
                
            if not cl:
                error_count += 1
                continue
                
            # Determine Section
            sec_name = get_row_field(row, ['section_name', 'section name', 'section', 'sec'], default=default_section_name)
            sec_name = sec_name.upper().strip() or 'A'
            sec = Section.objects.filter(class_level=cl, name__iexact=sec_name).first()
            if not sec:
                sec = Section.objects.create(class_level=cl, name=sec_name, capacity=40, room_number='Room 101')
                
            # Admission Number in file (if any)
            adm_no = get_row_field(row, ['admission_no', 'admission no', 'admission_number', 'admission number', 'adm_no', 'adm no'])
            
            student_display_name = f"{first_name} {last_name}".strip()
            dob_display = dob_val.strftime('%d/%m/%Y') if hasattr(dob_val, 'strftime') else str(dob_val)
            
            # ================== DUPLICATE CHECK VALIDATION ==================
            # Duplicate check based on Student Name, Father Name, and Date of Birth
            student_norm = normalize_student_name(first_name, last_name)
            father_norm = normalize_parent_name(parent_name)
            file_key = (student_norm, father_norm, dob_val)
            
            # Check 1: Duplicate within the uploaded spreadsheet
            if file_key in seen_in_file:
                prev_row = seen_in_file[file_key]
                duplicate_rows.append({
                    'row_num': idx,
                    'student_name': student_display_name,
                    'father_name': parent_name,
                    'dob': dob_display,
                    'class_name': cl.name if cl else (class_name or 'N/A'),
                    'section': sec.name if sec else 'A',
                    'reason': f"Duplicate entry in spreadsheet (Matches Row #{prev_row})"
                })
                duplicate_count += 1
                continue
                
            # Check 2: Duplicate against database records
            existing_student = find_duplicate_student(
                first_name=first_name,
                last_name=last_name,
                father_name=parent_name,
                dob=dob_val
            )
            if existing_student:
                existing_adm = existing_student.admission_number
                existing_class = existing_student.class_level.name if existing_student.class_level else 'N/A'
                duplicate_rows.append({
                    'row_num': idx,
                    'student_name': student_display_name,
                    'father_name': parent_name,
                    'dob': dob_display,
                    'class_name': cl.name if cl else (class_name or 'N/A'),
                    'section': sec.name if sec else 'A',
                    'reason': f"Student already exists in database (Adm No: {existing_adm}, Class: {existing_class})"
                })
                duplicate_count += 1
                continue
                
            # Check 3: Explicit Admission Number collision (if provided in file)
            if adm_no and Student.objects.filter(admission_number=adm_no).exists():
                existing_by_adm = Student.objects.filter(admission_number=adm_no).first()
                duplicate_rows.append({
                    'row_num': idx,
                    'student_name': student_display_name,
                    'father_name': parent_name,
                    'dob': dob_display,
                    'class_name': cl.name if cl else (class_name or 'N/A'),
                    'section': sec.name if sec else 'A',
                    'reason': f"Admission number '{adm_no}' is already assigned to student '{existing_by_adm.full_name}'"
                })
                duplicate_count += 1
                continue
                
            # Valid new entry: record key in seen_in_file
            seen_in_file[file_key] = idx
            
            # Auto-generate standardized admission number if not provided
            if not adm_no:
                adm_no = generate_student_admission_number(school, active_session)
                    
            # Roll Number
            roll_str = get_row_field(row, ['roll_no', 'roll no', 'roll_number', 'roll number'], default='1')
            try:
                roll_val = int(float(roll_str))
            except (ValueError, TypeError):
                roll_val = 1
                
            # Gender
            gender_val = get_row_field(row, ['gender', 'sex'], default='Male').capitalize()
            if gender_val not in ['Male', 'Female', 'Other']:
                gender_val = 'Male'
                
            blood_group = get_row_field(row, ['blood_group', 'blood group', 'bloodgroup'], default='B+').upper()
            if blood_group not in ['A+', 'A-', 'B+', 'B-', 'O+', 'O-', 'AB+', 'AB-']:
                blood_group = 'B+'
                
            parent_rel = get_row_field(row, ['parent_relation', 'parent relation', 'relation'], default='Father').capitalize()
            parent_phone = get_row_field(row, ['parent_phone', 'parent phone', 'phone', 'mobile', 'contact'], default='+91 9876543210')
            student_aadhaar = get_row_field(row, ['student_aadhaar', 'student aadhaar', 'student_aadhar', 'student aadhar', 'aadhaar', 'aadhar'])
            parent_aadhaar = get_row_field(row, ['parent_aadhaar', 'parent aadhaar', 'parent_aadhar', 'parent aadhar'])
            address = get_row_field(row, ['residential_address', 'residential address', 'address', 'full_address'], default='New Delhi, India')
            emergency = get_row_field(row, ['emergency_contact', 'emergency contact', 'emergency_phone'], default=parent_phone)
            
            if student_aadhaar:
                student_aadhaar = "".join(ch for ch in str(student_aadhaar) if ch.isdigit())[:14]
            if parent_aadhaar:
                parent_aadhaar = "".join(ch for ch in str(parent_aadhaar) if ch.isdigit())[:14]
                
            # Create Student Record (duplicates are strictly rejected above)
            student = Student.objects.create(
                school=school,
                admission_number=adm_no,
                roll_number=roll_val,
                first_name=first_name,
                last_name=last_name,
                gender=gender_val,
                dob=dob_val,
                blood_group=blood_group,
                class_level=cl,
                section=sec,
                academic_session=active_session,
                admission_date=timezone.now().date(),
                parent_name=parent_name,
                parent_relation=parent_rel,
                parent_phone=parent_phone,
                student_aadhaar=student_aadhaar,
                parent_aadhaar=parent_aadhaar,
                address=address,
                emergency_contact=emergency,
                is_active=True,
            )
            StudentSessionEnrollment.objects.get_or_create(
                student=student,
                academic_session=active_session,
                defaults={
                    'class_level': cl,
                    'section': sec,
                    'roll_number': roll_val,
                    'status': 'Active'
                }
            )
            
            created_count += 1
            created_students.append({
                'row_num': idx,
                'admission_number': student.admission_number,
                'student_name': student.full_name,
                'father_name': student.parent_name,
                'dob': dob_display,
                'class_name': cl.name,
                'section': sec.name
            })
            
            # Automatically create Admission Fee invoice
            try:
                fm = cl.fee_master
                adm_amt = fm.admission_fee
            except Exception:
                adm_amt = Decimal('5000.00')
                
            adm_cat = FeeCategory.objects.filter(name__icontains='Admission').first()
            if not adm_cat:
                adm_cat, _ = FeeCategory.objects.get_or_create(
                    name="Admission Fee",
                    defaults={'default_amount': adm_amt, 'description': "One-time student admission and registration fee"}
                )
            FeeInvoice.objects.get_or_create(
                invoice_number=f"INV-{timezone.now().year}-ADM{student.id:04d}",
                defaults={
                    'student': student,
                    'fee_category': adm_cat,
                    'amount': adm_amt,
                    'due_date': timezone.now().date() + datetime.timedelta(days=15),
                    'status': 'Unpaid'
                }
            )
            
        upload_result = {
            'filename': uploaded_file.name,
            'total_rows': len(rows_data),
            'created_count': created_count,
            'duplicate_count': duplicate_count,
            'error_count': error_count,
            'duplicates': duplicate_rows,
            'created_students': created_students,
        }
        
        # When duplicate entries occur, show them clearly on the upload page with full breakdown
        if duplicate_count > 0:
            if created_count > 0:
                messages.success(request, f"Successfully enrolled {created_count} new students!")
                messages.warning(
                    request,
                    f"Notice: {duplicate_count} duplicate entries were found and skipped (not saved to database). See duplicate details below."
                )
            else:
                messages.warning(
                    request,
                    f"No new students added! All {duplicate_count} entries in the file already exist (duplicate records)."
                )
            ctx.update({
                'classes': classes,
                'upload_result': upload_result,
                'title': 'Bulk Student Excel / CSV Upload'
            })
            return render(request, 'school_erp/student_bulk_upload.html', ctx)
        elif created_count > 0:
            messages.success(
                request,
                f"Successfully enrolled all {created_count} new students!"
            )
            return redirect('student_list')
        else:
            messages.warning(request, "No valid student data found. Please check file format.")
            return redirect('student_bulk_upload')
            
    ctx.update({
        'classes': classes,
        'title': 'Bulk Student Excel / CSV Upload'
    })
    return render(request, 'school_erp/student_bulk_upload.html', ctx)


def student_id_card_view(request, pk):
    ctx = get_school_context(request)
    student = get_object_or_404(Student.objects.select_related('class_level', 'section'), pk=pk)
    school = ctx['school']

    public_slip_url = request.build_absolute_uri(f"/slip/id-card/{student.id}/")
    clean_phone = clean_phone_for_wa(student.parent_phone)

    wa_message = (
        f"*{school.name.upper()}*\n"
        f"*Student Identity Card*\n"
        f"-----------------------------------------\n"
        f"Student: *{student.full_name}*\n"
        f"Adm No: *{student.admission_number}* | Class: *{student.class_level.name}-{student.section.name}*\n"
        f"Father: *{student.parent_name}*\n"
        f"-----------------------------------------\n"
        f"📄 *View / Print Official Student ID Card (No Login Needed):*\n"
        f"{public_slip_url}\n"
        f"-----------------------------------------\n"
        f"Administration Directorate, {school.name}"
    )

    ctx.update({
        'student': student,
        'clean_parent_phone': clean_phone,
        'wa_message': wa_message,
        'public_slip_url': public_slip_url,
    })
    return render(request, 'school_erp/student_id_card.html', ctx)


# ------------------- TEACHERS -------------------

@require_menu_permission('can_manage_users', 'Teachers & Staff')
def teacher_list_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    teachers_qs = Teacher.objects.filter(school=school, is_active=True).prefetch_related('subjects_taught', 'class_teacher_sections__class_level').order_by('employee_id')
    
    dept = request.GET.get('dept', '').strip()
    search = request.GET.get('q', '').strip()
    
    if dept:
        teachers_qs = teachers_qs.filter(department__icontains=dept)
    if search:
        teachers_qs = teachers_qs.filter(
            Q(first_name__icontains=search) |
            Q(last_name__icontains=search) |
            Q(employee_id__icontains=search) |
            Q(designation__icontains=search) |
            Q(qualification__icontains=search) |
            Q(phone__icontains=search) |
            Q(email__icontains=search)
        )
    
    total_count = teachers_qs.count()
    paginator = Paginator(teachers_qs, 20)
    page_number = request.GET.get('page', 1)
    try:
        page_obj = paginator.get_page(page_number)
    except (EmptyPage, PageNotAnInteger):
        page_obj = paginator.get_page(1)
        
    ctx.update({
        'teachers': page_obj,
        'page_obj': page_obj,
        'paginator': paginator,
        'is_paginated': page_obj.has_other_pages(),
        'selected_dept': dept,
        'search_query': search,
        'total_count': total_count,
    })
    return render(request, 'school_erp/teacher_list.html', ctx)


@require_menu_permission('can_manage_users', 'Teachers & Staff')
def teacher_create_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    if request.method == 'POST':
        form = TeacherForm(request.POST)
        if form.is_valid():
            teacher = form.save(commit=False)
            teacher.school = school
            teacher.save()
            form.save_m2m()
            messages.success(request, f"Faculty member '{teacher.full_name}' added successfully!")
            return redirect('teacher_detail', pk=teacher.pk)
    else:
        count = Teacher.objects.filter(school=school).count() + 1
        default_id = f"EMP-2026-{count:02d}"
        form = TeacherForm(initial={'employee_id': default_id, 'joining_date': timezone.now().date()})
        
    ctx.update({'form': form, 'title': 'Register New Faculty / Staff', 'edit_mode': False})
    return render(request, 'school_erp/teacher_form.html', ctx)


@require_menu_permission('can_manage_users', 'Teacher Profile')
def teacher_detail_view(request, pk):
    ctx = get_school_context(request)
    school = ctx['school']
    teacher = get_object_or_404(
        Teacher.objects.prefetch_related('subjects_taught', 'class_teacher_sections__class_level'),
        pk=pk,
        school=school
    )
    
    # Timetable periods for this teacher
    teacher_timetables = Timetable.objects.filter(teacher=teacher).select_related('class_level', 'section', 'subject').order_by('day', 'period_number')
    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']
    schedule_by_day = {day: [] for day in days}
    for item in teacher_timetables:
        if item.day in schedule_by_day:
            schedule_by_day[item.day].append(item)
            
    # Class teacher sections
    managed_sections = teacher.class_teacher_sections.select_related('class_level').all()
    schedule_days = [{'day': day, 'periods': schedule_by_day.get(day, []), 'count': len(schedule_by_day.get(day, []))} for day in days]
    
    ctx.update({
        'teacher': teacher,
        'managed_sections': managed_sections,
        'schedule_by_day': schedule_by_day,
        'schedule_days': schedule_days,
        'total_periods': teacher_timetables.count(),
        'days': days,
        'title': f'Faculty Profile: {teacher.full_name}'
    })
    return render(request, 'school_erp/teacher_detail.html', ctx)


@require_menu_permission('can_manage_users', 'Edit Teacher')
def teacher_edit_view(request, pk):
    ctx = get_school_context(request)
    school = ctx['school']
    teacher = get_object_or_404(Teacher, pk=pk, school=school)
    
    if request.method == 'POST':
        form = TeacherForm(request.POST, instance=teacher)
        if form.is_valid():
            teacher = form.save(commit=False)
            teacher.school = school
            teacher.save()
            form.save_m2m()
            messages.success(request, f"Faculty details for '{teacher.full_name}' updated successfully!")
            return redirect('teacher_detail', pk=teacher.pk)
    else:
        form = TeacherForm(instance=teacher)
        
    ctx.update({
        'form': form,
        'teacher': teacher,
        'title': f'Edit Faculty: {teacher.full_name}',
        'edit_mode': True
    })
    return render(request, 'school_erp/teacher_form.html', ctx)


@require_menu_permission('can_manage_users', 'Delete Faculty')
def teacher_delete_view(request, pk):
    teacher = get_object_or_404(Teacher, pk=pk)
    if request.method == 'POST':
        name = teacher.full_name
        # Clear foreign keys on section before deleting
        Section.objects.filter(class_teacher=teacher).update(class_teacher=None)
        Timetable.objects.filter(teacher=teacher).update(teacher=None)
        teacher.delete()
        messages.success(request, f"Faculty member '{name}' has been deleted from records.")
    return redirect('teacher_list')


def teacher_portal_view(request):
    """
    Dedicated Teacher Portal workspace.
    Shows assigned class section (if class teacher), today's timetable,
    weekly schedule, quick attendance taking shortcut, marks entry shortcut,
    and notices.
    """
    if not request.user.is_authenticated:
        return redirect('user_login')
        
    ctx = get_school_context(request)
    
    # Determine the teacher
    teacher = None
    teacher_id = request.GET.get('teacher_id')
    if teacher_id and request.user.is_superuser:
        teacher = Teacher.objects.filter(id=teacher_id).first()
    
    if not teacher and hasattr(request.user, 'teacher'):
        teacher = request.user.teacher
    elif not teacher:
        # Match by email or first/last name
        teacher = Teacher.objects.filter(email=request.user.email).first()
        if not teacher and request.user.first_name:
            teacher = Teacher.objects.filter(first_name__icontains=request.user.first_name).first()
        if not teacher:
            # Fallback to first active teacher for demonstration
            teacher = Teacher.objects.filter(is_active=True).first()
            
    all_teachers = Teacher.objects.filter(is_active=True).order_by('first_name', 'last_name')
    
    # Days and today's day name
    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']
    today_date = timezone.now().date()
    today_name = today_date.strftime('%A')
    
    schedule_by_day = {day: [] for day in days}
    today_classes = []
    managed_sections = []
    
    if teacher:
        timetables = Timetable.objects.filter(teacher=teacher).select_related('class_level', 'section', 'subject').order_by('period_number')
        for item in timetables:
            if item.day in schedule_by_day:
                schedule_by_day[item.day].append(item)
            if item.day == today_name:
                today_classes.append(item)
                
        managed_sections = Section.objects.filter(class_teacher=teacher).select_related('class_level')
        
    # Check attendance status for managed sections today
    section_attendance_status = []
    for sec in managed_sections:
        total_students = sec.students.filter(is_active=True).count()
        today_att = Attendance.objects.filter(student__section=sec, date=today_date)
        present_count = today_att.filter(status='Present').count()
        absent_count = today_att.filter(status='Absent').count()
        is_taken = today_att.exists()
        section_attendance_status.append({
            'section': sec,
            'total_students': total_students,
            'present_count': present_count,
            'absent_count': absent_count,
            'is_taken': is_taken,
        })
        
    recent_notices = Notice.objects.filter(
        Q(target_audience__in=['All', 'Teachers'])
    ).order_by('-is_urgent', '-created_at')[:4]
    
    schedule_days = [{'day': day, 'periods': schedule_by_day.get(day, []), 'count': len(schedule_by_day.get(day, []))} for day in days]

    ctx.update({
        'teacher': teacher,
        'all_teachers': all_teachers,
        'today_name': today_name,
        'today_date': today_date,
        'today_classes': today_classes,
        'schedule_by_day': schedule_by_day,
        'schedule_days': schedule_days,
        'managed_sections': managed_sections,
        'section_attendance_status': section_attendance_status,
        'recent_notices': recent_notices,
        'days': days,
        'title': 'Faculty & Teacher Workspace Portal',
    })
    return render(request, 'school_erp/teacher_portal.html', ctx)



# ------------------- ACADEMICS & CLASSES -------------------

@require_menu_permission('can_manage_classes', 'Classes & Sections')
def classes_view(request):
    ctx = get_school_context(request)
    school = ctx.get('school')
    
    if request.method == 'POST':
        action = request.POST.get('action')
        
        # 1. Add Single Class
        if action == 'add_class':
            name = request.POST.get('name', '').strip()
            code = request.POST.get('code', '').strip()
            order = request.POST.get('order', '').strip()
            create_section = request.POST.get('create_default_section') == 'on'
            section_name = request.POST.get('default_section_name', 'A').strip().upper() or 'A'
            section_room = request.POST.get('default_section_room', 'Room 101').strip() or 'Room 101'
            section_capacity = request.POST.get('default_section_capacity', '40').strip()
            teacher_id = request.POST.get('default_class_teacher_id')
            
            if not name:
                messages.error(request, "Class Name is required.")
            elif ClassLevel.objects.filter(school=school, name__iexact=name).exists():
                messages.error(request, f"Class '{name}' already exists in your school!")
            else:
                # Auto generate code if blank
                if not code:
                    clean_code = "".join(ch for ch in name if ch.isalnum()).upper()
                    base_code = clean_code[:8] or "CLS"
                    code = base_code
                    suffix = 1
                    while ClassLevel.objects.filter(school=school, code=code).exists():
                        code = f"{base_code[:6]}{suffix}"
                        suffix += 1
                elif ClassLevel.objects.filter(school=school, code__iexact=code).exists():
                    messages.error(request, f"Class code '{code}' is already assigned to another class in your school. Please choose a different code.")
                    return redirect('classes')
                
                try:
                    order_val = int(order) if order else ((ClassLevel.objects.filter(school=school).order_by('-order').values_list('order', flat=True).first() or 0) + 1)
                except ValueError:
                    order_val = (ClassLevel.objects.filter(school=school).order_by('-order').values_list('order', flat=True).first() or 0) + 1
                    
                new_class = ClassLevel.objects.create(school=school, name=name, code=code, order=order_val)
                
                # Create default section
                if create_section:
                    try:
                        cap_val = int(section_capacity) if section_capacity else 40
                    except ValueError:
                        cap_val = 40
                        
                    teacher = None
                    if teacher_id and teacher_id.isdigit():
                        teacher = Teacher.objects.filter(school=school, id=int(teacher_id)).first()
                        
                    Section.objects.create(
                        class_level=new_class,
                        name=section_name,
                        capacity=cap_val,
                        room_number=section_room,
                        class_teacher=teacher
                    )
                
                # Auto create Fee Master for this class
                ClassFeeMaster.objects.get_or_create(class_level=new_class)
                messages.success(request, f"New class '{new_class.name}' (Code: {new_class.code}) added successfully!")
            return redirect('classes')

        # 2. Bulk Add Classes
        elif action == 'bulk_add_classes':
            batch_preset = request.POST.get('batch_preset', '')
            custom_classes = request.POST.get('custom_classes', '').strip()
            
            preset_map = {
                'pre_primary': ['Playgroup', 'Nursery', 'LKG', 'UKG'],
                'primary': ['Class 1', 'Class 2', 'Class 3', 'Class 4', 'Class 5'],
                'middle': ['Class 6', 'Class 7', 'Class 8'],
                'secondary': ['Class 9', 'Class 10'],
                'senior_secondary': ['Class 11', 'Class 12'],
                'k12': ['Nursery', 'LKG', 'UKG'] + [f"Class {i}" for i in range(1, 13)]
            }
            
            names_to_add = []
            if batch_preset in preset_map:
                names_to_add = preset_map[batch_preset]
            elif custom_classes:
                names_to_add = [c.strip() for c in custom_classes.split(',') if c.strip()]
            
            if not names_to_add:
                messages.warning(request, "No classes selected. Please choose a preset or enter class names.")
                return redirect('classes')
                
            created_count = 0
            curr_max_order = (ClassLevel.objects.filter(school=school).order_by('-order').values_list('order', flat=True).first() or 0)
            
            for cname in names_to_add:
                if not ClassLevel.objects.filter(school=school, name__iexact=cname).exists():
                    curr_max_order += 1
                    clean_code = "".join(ch for ch in cname if ch.isalnum()).upper()
                    base_code = clean_code[:8] or f"C{curr_max_order}"
                    code = base_code
                    suffix = 1
                    while ClassLevel.objects.filter(school=school, code=code).exists():
                        code = f"{base_code[:6]}{suffix}"
                        suffix += 1
                        
                    cl = ClassLevel.objects.create(school=school, name=cname, code=code, order=curr_max_order)
                    # Create default Section A
                    Section.objects.create(class_level=cl, name='A', capacity=40, room_number='Room 101')
                    # Create default Fee Master
                    ClassFeeMaster.objects.get_or_create(class_level=cl)
                    created_count += 1
                    
            if created_count > 0:
                messages.success(request, f"Great! {created_count} new classes (with default Section 'A' and Fee Master) were created successfully.")
            else:
                messages.info(request, "All selected classes already exist in the system.")
            return redirect('classes')

        # 3. Edit Class
        elif action == 'edit_class':
            class_id = request.POST.get('class_id')
            class_obj = get_object_or_404(ClassLevel, id=class_id)
            name = request.POST.get('name', '').strip()
            code = request.POST.get('code', '').strip()
            order = request.POST.get('order', '').strip()
            
            if not name:
                messages.error(request, "Class Name cannot be empty.")
                return redirect('classes')
                
            if ClassLevel.objects.filter(school=school, name__iexact=name).exclude(id=class_obj.id).exists():
                messages.error(request, f"Another class with name '{name}' already exists in your school.")
                return redirect('classes')
                
            if code and ClassLevel.objects.filter(school=school, code__iexact=code).exclude(id=class_obj.id).exists():
                messages.error(request, f"Class code '{code}' is already used by another class in your school.")
                return redirect('classes')
                
            class_obj.name = name
            if code:
                class_obj.code = code
            if order:
                try:
                    class_obj.order = int(order)
                except ValueError:
                    pass
            class_obj.save()
            messages.success(request, f"Class '{class_obj.name}' updated successfully.")
            return redirect('classes')

        # 4. Auto-sort classes by standard educational level (Nursery -> 12th)
        elif action == 'auto_sort_classes':
            import re
            def natural_class_sort_key(name):
                n = name.lower().strip()
                if any(k in n for k in ['play', 'creche', 'toddler']):
                    return (1, 0, n)
                if 'nur' in n:
                    return (2, 0, n)
                if any(k in n for k in ['lkg', 'kg1', 'kg-1', 'pp1', 'lower']):
                    return (3, 0, n)
                if any(k in n for k in ['ukg', 'kg2', 'kg-2', 'pp2', 'upper', 'prep']):
                    return (4, 0, n)
                m = re.search(r'\b(1[0-2]|[1-9])\b', n)
                if m:
                    grade = int(m.group(1))
                    stream = 0
                    if 'sci' in n: stream = 1
                    elif 'comm' in n: stream = 2
                    elif 'art' in n or 'hum' in n: stream = 3
                    return (10 + grade, stream, n)
                return (100, 0, n)

            cls_list = list(ClassLevel.objects.filter(school=school))
            cls_list.sort(key=lambda c: natural_class_sort_key(c.name))
            for idx, c in enumerate(cls_list, start=1):
                if c.order != idx:
                    c.order = idx
                    c.save(update_fields=['order'])
            messages.success(request, "All classes have been automatically reordered chronologically (Nursery -> Class 1..12).")
            return redirect('classes')

        # 5. Delete Class
        elif action == 'delete_class':
            class_id = request.POST.get('class_id')
            class_obj = get_object_or_404(ClassLevel, id=class_id)
            student_count = class_obj.student_set.count()
            if student_count > 0:
                messages.error(request, f"Security Warning: Cannot delete class '{class_obj.name}' because it has {student_count} enrolled students. Transfer students to another class first.")
            else:
                c_name = class_obj.name
                class_obj.delete()
                messages.success(request, f"Class '{c_name}' deleted successfully.")
            return redirect('classes')

        # 5. Add Section
        elif action == 'add_section':
            class_id = request.POST.get('class_id')
            class_obj = get_object_or_404(ClassLevel, id=class_id)
            sec_name = request.POST.get('name', '').strip().upper()
            capacity = request.POST.get('capacity', '40').strip()
            room_number = request.POST.get('room_number', 'Room 101').strip() or 'Room 101'
            teacher_id = request.POST.get('class_teacher_id')
            
            if not sec_name:
                messages.error(request, "Section name (e.g. A, B, C) is required.")
                return redirect('classes')
                
            if Section.objects.filter(class_level=class_obj, name__iexact=sec_name).exists():
                messages.error(request, f"Section '{sec_name}' already exists in class {class_obj.name}!")
                return redirect('classes')
                
            try:
                cap_val = int(capacity) if capacity else 40
            except ValueError:
                cap_val = 40
                
            teacher = None
            if teacher_id and teacher_id.isdigit():
                teacher = Teacher.objects.filter(id=int(teacher_id)).first()
                
            sec = Section.objects.create(
                class_level=class_obj,
                name=sec_name,
                capacity=cap_val,
                room_number=room_number,
                class_teacher=teacher
            )
            messages.success(request, f"New section '{sec.name}' added successfully to class {class_obj.name}!")
            return redirect('classes')

        # 6. Edit Section
        elif action == 'edit_section':
            section_id = request.POST.get('section_id')
            sec = get_object_or_404(Section, id=section_id)
            sec_name = request.POST.get('name', '').strip().upper()
            capacity = request.POST.get('capacity', '').strip()
            room_number = request.POST.get('room_number', '').strip()
            teacher_id = request.POST.get('class_teacher_id')
            
            if not sec_name:
                messages.error(request, "Section name cannot be empty.")
                return redirect('classes')
                
            if Section.objects.filter(class_level=sec.class_level, name__iexact=sec_name).exclude(id=sec.id).exists():
                messages.error(request, f"Section '{sec_name}' already exists in class {sec.class_level.name}.")
                return redirect('classes')
                
            sec.name = sec_name
            if capacity:
                try:
                    sec.capacity = int(capacity)
                except ValueError:
                    pass
            if room_number:
                sec.room_number = room_number
                
            if teacher_id is not None:
                if teacher_id == '' or teacher_id == '0':
                    sec.class_teacher = None
                elif teacher_id.isdigit():
                    sec.class_teacher = Teacher.objects.filter(id=int(teacher_id)).first()
                    
            sec.save()
            messages.success(request, f"Section '{sec.name}' ({sec.class_level.name}) updated successfully.")
            return redirect('classes')

        # 7. Delete Section
        elif action == 'delete_section':
            section_id = request.POST.get('section_id')
            sec = get_object_or_404(Section, id=section_id)
            student_count = sec.students.count()
            if student_count > 0:
                messages.error(request, f"Security Warning: Cannot delete section '{sec.name}' because it has {student_count} enrolled students.")
            else:
                s_name = sec.name
                c_name = sec.class_level.name
                sec.delete()
                messages.success(request, f"Section '{s_name}' ({c_name}) deleted successfully.")
            return redirect('classes')

    classes = ClassLevel.objects.filter(school=school).prefetch_related('sections__class_teacher', 'sections__students', 'subjects').order_by('order', 'name')
    teachers = Teacher.objects.filter(school=school, is_active=True).order_by('first_name', 'last_name')
    
    total_classes = classes.count()
    total_sections = Section.objects.filter(class_level__school=school).count()
    total_students = Student.objects.filter(class_level__school=school, is_active=True).count()
    total_capacity = sum(sum(s.capacity for s in c.sections.all()) for c in classes)
    
    ctx.update({
        'classes': classes,
        'teachers': teachers,
        'total_classes': total_classes,
        'total_sections': total_sections,
        'total_students': total_students,
        'total_capacity': total_capacity,
    })
    return render(request, 'school_erp/classes.html', ctx)


@require_menu_permission('can_manage_timetable', 'Timetable Routine')
def timetable_view(request):
    ctx = get_school_context(request)
    school = ctx.get('school')
    classes = ClassLevel.objects.filter(school=school).order_by('order', 'name')
    sections = Section.objects.filter(class_level__school=school)
    
    class_id = request.GET.get('class_id')
    section_id = request.GET.get('section_id')
    
    selected_class = None
    selected_section = None
    timetables = []
    
    if class_id:
        selected_class = ClassLevel.objects.filter(school=school, id=class_id).first()
    elif classes.exists():
        selected_class = classes.first()

    if selected_class:
        if section_id:
            selected_section = Section.objects.filter(class_level=selected_class, id=section_id).first()
        if not selected_section:
            selected_section = selected_class.sections.first()
            
    if selected_class and selected_section:
        timetables = list(Timetable.objects.filter(class_level=selected_class, section=selected_section).select_related('subject', 'teacher').order_by('day', 'period_number', 'start_time'))
            
    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']
    
    # Organize schedule by day
    schedule_by_day = {day: [] for day in days}
    for item in timetables:
        if item.day in schedule_by_day:
            schedule_by_day[item.day].append(item)

    subjects = Subject.objects.all().order_by('name')
    teachers = Teacher.objects.filter(school=school, is_active=True).order_by('first_name', 'last_name') if school else Teacher.objects.filter(is_active=True).order_by('first_name', 'last_name')
    if not teachers.exists():
        teachers = Teacher.objects.filter(is_active=True).order_by('first_name', 'last_name')

    total_periods_count = len(timetables)
            
    ctx.update({
        'classes': classes,
        'sections': sections,
        'selected_class': selected_class,
        'selected_section': selected_section,
        'schedule_by_day': schedule_by_day,
        'days': days,
        'subjects': subjects,
        'teachers': teachers,
        'total_periods_count': total_periods_count,
    })
    return render(request, 'school_erp/timetable.html', ctx)


@require_menu_permission('can_manage_timetable', 'Timetable Routine')
def timetable_create_view(request):
    ctx = get_school_context(request)
    school = ctx.get('school')

    if request.method == 'POST':
        class_id = request.POST.get('class_level') or request.POST.get('class_id')
        section_id = request.POST.get('section') or request.POST.get('section_id')
        day = request.POST.get('day')
        period_number = request.POST.get('period_number', 1)
        subject_id = request.POST.get('subject') or request.POST.get('subject_id')
        teacher_id = request.POST.get('teacher') or request.POST.get('teacher_id')
        start_time = request.POST.get('start_time')
        end_time = request.POST.get('end_time')
        room_number = request.POST.get('room_number', '').strip()

        try:
            class_level = get_object_or_404(ClassLevel, id=class_id)
            section = get_object_or_404(Section, id=section_id)
            subject = get_object_or_404(Subject, id=subject_id)
            teacher = Teacher.objects.filter(id=teacher_id).first() if teacher_id else None
            
            if not room_number:
                room_number = section.room_number or "Room 101"

            tt = Timetable.objects.create(
                class_level=class_level,
                section=section,
                day=day,
                period_number=int(period_number),
                subject=subject,
                teacher=teacher,
                start_time=start_time,
                end_time=end_time,
                room_number=room_number
            )
            messages.success(request, f"Period {tt.period_number} on {tt.day} ({tt.subject.name}) has been added successfully!")
        except Exception as e:
            messages.error(request, f"Failed to add period: {str(e)}")

        return redirect(f"/timetable/?class_id={class_id}&section_id={section_id}")

    # Fallback GET standalone form
    class_id = request.GET.get('class_id')
    section_id = request.GET.get('section_id')
    classes = ClassLevel.objects.filter(school=school).order_by('order', 'name')
    teachers = Teacher.objects.filter(is_active=True).order_by('first_name')
    subjects = Subject.objects.all().order_by('name')

    selected_class = ClassLevel.objects.filter(id=class_id).first() if class_id else classes.first()
    selected_section = Section.objects.filter(id=section_id).first() if section_id else (selected_class.sections.first() if selected_class else None)

    ctx.update({
        'classes': classes,
        'teachers': teachers,
        'subjects': subjects,
        'selected_class': selected_class,
        'selected_section': selected_section,
        'edit_mode': False,
        'days': ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday'],
    })
    return render(request, 'school_erp/timetable_form.html', ctx)


@require_menu_permission('can_manage_timetable', 'Timetable Routine')
def timetable_edit_view(request, pk):
    ctx = get_school_context(request)
    school = ctx.get('school')
    tt = get_object_or_404(Timetable, pk=pk)

    if request.method == 'POST':
        class_id = request.POST.get('class_level') or request.POST.get('class_id') or tt.class_level_id
        section_id = request.POST.get('section') or request.POST.get('section_id') or tt.section_id
        day = request.POST.get('day') or tt.day
        period_number = request.POST.get('period_number') or tt.period_number
        subject_id = request.POST.get('subject') or request.POST.get('subject_id')
        teacher_id = request.POST.get('teacher') or request.POST.get('teacher_id')
        start_time = request.POST.get('start_time') or tt.start_time
        end_time = request.POST.get('end_time') or tt.end_time
        room_number = request.POST.get('room_number', '').strip() or tt.room_number

        try:
            if subject_id:
                tt.subject = get_object_or_404(Subject, id=subject_id)
            tt.teacher = Teacher.objects.filter(id=teacher_id).first() if teacher_id else None
            tt.day = day
            tt.period_number = int(period_number)
            tt.start_time = start_time
            tt.end_time = end_time
            tt.room_number = room_number
            if class_id:
                tt.class_level_id = class_id
            if section_id:
                tt.section_id = section_id
            tt.save()
            messages.success(request, f"Period {tt.period_number} on {tt.day} updated successfully!")
        except Exception as e:
            messages.error(request, f"Failed to update period: {str(e)}")

        return redirect(f"/timetable/?class_id={tt.class_level_id}&section_id={tt.section_id}")

    # JSON response for dynamic AJAX modal fill
    if request.headers.get('x-requested-with') == 'XMLHttpRequest' or request.GET.get('format') == 'json':
        from django.http import JsonResponse
        return JsonResponse({
            'id': tt.id,
            'class_id': tt.class_level_id,
            'class_name': tt.class_level.name,
            'section_id': tt.section_id,
            'section_name': tt.section.name,
            'day': tt.day,
            'period_number': tt.period_number,
            'subject_id': tt.subject_id,
            'teacher_id': tt.teacher_id if tt.teacher_id else '',
            'start_time': tt.start_time.strftime('%H:%M') if tt.start_time else '',
            'end_time': tt.end_time.strftime('%H:%M') if tt.end_time else '',
            'room_number': tt.room_number,
        })

    classes = ClassLevel.objects.filter(school=school).order_by('order', 'name')
    teachers = Teacher.objects.filter(is_active=True).order_by('first_name')
    subjects = Subject.objects.all().order_by('name')

    ctx.update({
        'classes': classes,
        'teachers': teachers,
        'subjects': subjects,
        'timetable': tt,
        'selected_class': tt.class_level,
        'selected_section': tt.section,
        'edit_mode': True,
        'days': ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday'],
    })
    return render(request, 'school_erp/timetable_form.html', ctx)


@require_menu_permission('can_manage_timetable', 'Timetable Routine')
def timetable_delete_view(request, pk):
    tt = get_object_or_404(Timetable, pk=pk)
    class_id = tt.class_level_id
    section_id = tt.section_id
    day = tt.day
    period_number = tt.period_number
    subj_name = tt.subject.name if tt.subject else "Subject"

    tt.delete()
    messages.success(request, f"Period {period_number} ({subj_name}) on {day} deleted successfully.")
    return redirect(f"/timetable/?class_id={class_id}&section_id={section_id}")


@require_menu_permission('can_manage_timetable', 'Timetable Routine')
def timetable_generate_default_view(request):
    ctx = get_school_context(request)
    school = ctx.get('school')
    class_id = request.POST.get('class_id') or request.GET.get('class_id')
    section_id = request.POST.get('section_id') or request.GET.get('section_id')

    if not class_id or not section_id:
        messages.error(request, "Please select a Class and Section first.")
        return redirect('timetable')

    class_level = get_object_or_404(ClassLevel, id=class_id)
    section = get_object_or_404(Section, id=section_id)

    overwrite = request.POST.get('overwrite') == '1' or request.GET.get('overwrite') == '1'
    existing = Timetable.objects.filter(class_level=class_level, section=section)
    if existing.exists() and not overwrite:
        messages.warning(request, f"This section already has {existing.count()} periods. Enable 'Overwrite' to reset.")
        return redirect(f"/timetable/?class_id={class_id}&section_id={section_id}")

    if overwrite:
        existing.delete()

    subjects = list(Subject.objects.filter(Q(class_levels=class_level) | Q(class_levels__isnull=True)).distinct().order_by('id'))
    if not subjects:
        subjects = list(Subject.objects.all().order_by('id'))

    teachers = list(Teacher.objects.filter(school=school, is_active=True).order_by('id')) if school else list(Teacher.objects.filter(is_active=True).order_by('id'))
    if not teachers:
        teachers = list(Teacher.objects.filter(is_active=True).order_by('id'))

    standard_slots = [
        (1, "08:30:00", "09:15:00"),
        (2, "09:15:00", "10:00:00"),
        (3, "10:00:00", "10:45:00"),
        (4, "11:15:00", "12:00:00"),
        (5, "12:00:00", "12:45:00"),
        (6, "12:45:00", "13:30:00"),
    ]

    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']
    room = section.room_number or "Room 101"

    created = 0
    subj_idx = 0
    for day in days:
        for p_num, s_time, e_time in standard_slots:
            if not subjects:
                break
            subj = subjects[subj_idx % len(subjects)]
            teacher = None
            if hasattr(subj, 'teachers') and subj.teachers.exists():
                teacher = subj.teachers.first()
            elif teachers:
                teacher = teachers[subj_idx % len(teachers)]

            Timetable.objects.create(
                class_level=class_level,
                section=section,
                day=day,
                period_number=p_num,
                subject=subj,
                teacher=teacher,
                start_time=s_time,
                end_time=e_time,
                room_number=room
            )
            created += 1
            subj_idx += 1

    messages.success(request, f"Standard dynamic routine generated with {created} periods for {class_level.name} - Section {section.name}! You can now edit any period anytime.")
    return redirect(f"/timetable/?class_id={class_id}&section_id={section_id}")



# ------------------- ATTENDANCE -------------------

@require_menu_permission('can_manage_attendance', 'Daily Attendance')
def attendance_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    classes = ClassLevel.objects.filter(school=school).order_by('order', 'name')
    sections = Section.objects.filter(class_level__school=school)
    
    class_id = request.GET.get('class_id') or request.POST.get('class_id')
    section_id = request.GET.get('section_id') or request.POST.get('section_id')
    date_str = request.GET.get('date') or request.POST.get('date') or str(timezone.now().date())
    
    try:
        current_date = datetime.datetime.strptime(date_str, '%Y-%m-%d').date()
    except Exception:
        current_date = timezone.now().date()
        
    selected_class = ClassLevel.objects.filter(school=school, id=class_id).first() if class_id else classes.first()
    selected_section = None
    if selected_class:
        selected_section = Section.objects.filter(class_level=selected_class, id=section_id).first() if section_id else selected_class.sections.first()
        
    students = []
    attendance_map = {}
    
    if selected_class and selected_section:
        students = Student.objects.filter(school=school, class_level=selected_class, section=selected_section, is_active=True).order_by('roll_number')
        
        # Handle POST attendance submission
        if request.method == 'POST' and 'save_attendance' in request.POST:
            for s in students:
                status_val = request.POST.get(f'status_{s.id}', 'Present')
                remarks_val = request.POST.get(f'remarks_{s.id}', '')
                Attendance.objects.update_or_create(
                    student=s,
                    date=current_date,
                    defaults={'status': status_val, 'remarks': remarks_val}
                )
            messages.success(request, f"Attendance saved successfully for {selected_class.name} - Section {selected_section.name} on {current_date.strftime('%d %b %Y')}!")
            return redirect(f"{request.path}?class_id={selected_class.id}&section_id={selected_section.id}&date={current_date}")
            
        existing_att = Attendance.objects.filter(student__in=students, date=current_date)
        attendance_map = {att.student_id: att.status for att in existing_att}
        
    # Pair students with attendance status
    student_rows = []
    for s in students:
        student_rows.append({
            'student': s,
            'status': attendance_map.get(s.id, 'Present') # default Present
        })
        
    ctx.update({
        'classes': classes,
        'sections': sections,
        'selected_class': selected_class,
        'selected_section': selected_section,
        'current_date': current_date.strftime('%Y-%m-%d'),
        'student_rows': student_rows,
    })
    return render(request, 'school_erp/attendance.html', ctx)


def attendance_report_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    classes = ClassLevel.objects.filter(school=school).order_by('order', 'name')
    class_id = request.GET.get('class_id')
    selected_class = ClassLevel.objects.filter(school=school, id=class_id).first() if class_id else classes.first()
    
    students_report = []
    if selected_class:
        students = Student.objects.filter(school=school, class_level=selected_class, is_active=True).order_by('section__name', 'roll_number')
        for s in students:
            stats = s.attendance_stats
            students_report.append({
                'student': s,
                'total': stats['total'],
                'present': stats['present'],
                'percentage': stats['percentage'],
            })
            
    ctx.update({
        'classes': classes,
        'selected_class': selected_class,
        'students_report': students_report,
    })
    return render(request, 'school_erp/attendance_report.html', ctx)


# ------------------- FEES & FINANCE -------------------

def ensure_fee_structure(school=None):
    try:
        call_command('migrate', interactive=False)
    except Exception:
        pass
    categories_to_ensure = [
        ("Admission Fee", Decimal('5000.00'), "One-time student admission and enrollment fee"),
        ("Monthly Fee", Decimal('2500.00'), "Monthly tuition and academic instruction fee"),
        ("Quarterly Fee", Decimal('7500.00'), "Quarterly composite fee (3 months)"),
        ("Half Yearly Fee", Decimal('15000.00'), "Half-yearly composite fee (6 months)"),
        ("Annual Fee", Decimal('28000.00'), "Full annual session academic composite fee"),
        ("Transport / Bus Fee", Decimal('1500.00'), "School bus & transport safety facility fee"),
        ("Examination Fee", Decimal('1500.00'), "Term examinations, test materials & report cards"),
        ("Book Set", Decimal('3000.00'), "Academic book set and notebooks"),
        ("School Uniform / Dress", Decimal('1800.00'), "School uniform and dress set"),
        ("Tie & Belt", Decimal('250.00'), "School logo tie and belt"),
        ("School Diary", Decimal('150.00'), "School student academic almanac/diary"),
        ("Student ID Card", Decimal('100.00'), "Student smart ID card and lanyard"),
        ("Other / Miscellaneous Fee", Decimal('500.00'), "Laboratory, sports & miscellaneous"),
    ]
    for name, default_amt, desc in categories_to_ensure:
        q = FeeCategory.objects.filter(name=name)
        if school:
            q = q.filter(school=school)
        if not q.exists():
            FeeCategory.objects.create(
                name=name,
                default_amount=default_amt,
                description=desc,
                school=school
            )

    # Seed Default Enterprise Fee Heads
    heads_to_ensure = [
        ("Tuition Fee", "TUITION", "Academic", Decimal('2000.00'), True, False, "Regular monthly academic tuition fee"),
        ("Admission & Registration Fee", "ADM", "OneTime", Decimal('5000.00'), False, False, "One-time enrollment & registration charges"),
        ("Transport / Bus Route Fee", "TRANS", "Transport", Decimal('1200.00'), True, True, "Monthly school bus route & transit facility"),
        ("Hostel & Boarding Fee", "HOSTEL", "Hostel", Decimal('6000.00'), True, True, "Monthly hostel accommodation, mess & care"),
        ("Examination & Evaluation Fee", "EXAM", "Exam", Decimal('1500.00'), True, False, "Terminal examination & marksheet issuance"),
        ("Computer & Science Lab Fee", "LAB", "Optional", Decimal('500.00'), True, True, "Practical laboratory & digital computer lab charges"),
        ("Library & Sports Development", "SPORTS", "Academic", Decimal('400.00'), True, False, "Library access, physical sports amenities"),
        ("Annual Development Charges", "ANNUAL", "Academic", Decimal('3000.00'), False, False, "Annual infrastructure maintenance fee"),
        ("Book & Notebook Set", "BOOKS", "Other", Decimal('2500.00'), False, False, "Complete curriculum book pack"),
        ("School Uniform Set", "UNIFORM", "Other", Decimal('1800.00'), False, False, "Official uniform, tie, belt and emblem"),
        ("Late Payment Fine", "FINE", "Fine", Decimal('50.00'), False, False, "Late payment penalty for dues passed grace period"),
        ("Caution Deposit (Refundable)", "CAUTION", "CautionDeposit", Decimal('2000.00'), False, True, "Refundable security caution deposit at TC"),
    ]
    for h_name, h_code, h_type, h_amt, h_rec, h_opt, h_desc in heads_to_ensure:
        q_head = FeeHead.objects.filter(name=h_name)
        if school:
            q_head = q_head.filter(school=school)
        if not q_head.exists():
            FeeHead.objects.create(
                name=h_name,
                code=h_code,
                fee_type=h_type,
                default_amount=h_amt,
                is_recurring=h_rec,
                is_optional=h_opt,
                description=h_desc,
                school=school
            )

    # Seed Demo Transport & Hostel facilities if empty
    if TransportRoute.objects.count() == 0:
        TransportRoute.objects.create(
            school=school,
            route_name="Route 1 - City Center to Campus",
            vehicle_number="DL 1PB 4589",
            driver_name="Ramesh Singh",
            driver_phone="+91 98765 12345",
            monthly_fare=Decimal('1200.00'),
            pickup_points="City Center, Clock Tower, Gandhi Chowk, Campus Gate"
        )
        TransportRoute.objects.create(
            school=school,
            route_name="Route 2 - Sector 62 / Metro Corridor",
            vehicle_number="UP 16 AB 7721",
            driver_name="Sunil Kumar",
            driver_phone="+91 98765 67890",
            monthly_fare=Decimal('1500.00'),
            pickup_points="Sector 62 Metro, Model Town, Highway Crossing, Campus"
        )

    if HostelFacility.objects.count() == 0:
        HostelFacility.objects.create(
            school=school,
            hostel_name="Tagore Senior Boys Hostel",
            room_type="Double Sharing AC",
            monthly_fee=Decimal('7500.00'),
            warden_name="Dr. S. K. Mishra",
            warden_phone="+91 94150 99887",
            capacity=60
        )
        HostelFacility.objects.create(
            school=school,
            hostel_name="Sarojini Girls Hostel",
            room_type="Double Sharing Non-AC",
            monthly_fee=Decimal('6000.00'),
            warden_name="Mrs. Sunita Verma",
            warden_phone="+91 94150 22334",
            capacity=50
        )

    # Ensure ClassFeeMaster exists for every ClassLevel
    try:
        for cl in ClassLevel.objects.all():
            try:
                fm = cl.fee_master
            except Exception:
                fm = None
            if not fm:
                order = cl.order or 1
                if order <= 3:
                    m_fee, e_fee, t_fee, b_fee, u_fee = Decimal('1600.00'), Decimal('800.00'), Decimal('1000.00'), Decimal('2000.00'), Decimal('1400.00')
                elif order <= 8:
                    m_fee = Decimal('1800.00') + Decimal(str((order - 3) * 100))
                    e_fee, t_fee, b_fee, u_fee = Decimal('1000.00'), Decimal('1200.00'), Decimal('2600.00'), Decimal('1600.00')
                elif order <= 11:
                    m_fee = Decimal('2200.00') + Decimal(str((order - 8) * 100))
                    e_fee, t_fee, b_fee, u_fee = Decimal('1200.00'), Decimal('1400.00'), Decimal('3400.00'), Decimal('1800.00')
                else:
                    m_fee = Decimal('2800.00') + Decimal(str((order - 11) * 200))
                    e_fee, t_fee, b_fee, u_fee = Decimal('1500.00'), Decimal('1600.00'), Decimal('4200.00'), Decimal('2000.00')

                ClassFeeMaster.objects.get_or_create(
                    class_level=cl,
                    defaults={
                        'monthly_tuition_fee': m_fee,
                        'exam_fee': e_fee,
                        'transport_fee': t_fee,
                        'admission_fee': Decimal('5000.00'),
                        'book_set_fee': b_fee,
                        'dress_uniform_fee': u_fee,
                        'tie_belt_fee': Decimal('250.00'),
                        'diary_fee': Decimal('150.00'),
                        'id_card_fee': Decimal('100.00'),
                        'other_activity_fee': Decimal('500.00'),
                    }
                )
    except Exception:
        pass

    # Data fix for accidental 120,000 payment on Shaurya Dubey's invoice
    shaurya_inv = FeeInvoice.objects.filter(invoice_number__icontains='T1', student__admission_number='SRS-2026-1020').first()
    if shaurya_inv:
        huge_pay = shaurya_inv.payments.filter(amount_paid__gte=100000).first()
        if huge_pay:
            huge_pay.amount_paid = Decimal('12000.00')
            huge_pay.save()
            shaurya_inv.update_status()

    # Ensure demo unified combined receipt exists
    if FeeReceipt.objects.count() == 0:
        demo_student = Student.objects.filter(admission_number='SRS-2026-1020').first()
        if demo_student:
            r1 = FeeReceipt.objects.create(
                receipt_number="SRS-REC-2026-0001",
                student=demo_student,
                payment_date=timezone.now().date(),
                payment_method='UPI / QR',
                transaction_reference='UPI-PAY-889021',
                subtotal=Decimal('4000.00'),
                previous_due_cleared=Decimal('3500.00'),
                discount=Decimal('0.00'),
                total_billed=Decimal('7500.00'),
                amount_paid=Decimal('6000.00'),
                remaining_balance_due=Decimal('1500.00')
            )
            FeeReceiptItem.objects.create(
                receipt=r1,
                category_name="Monthly Fee",
                month_name="September 2026",
                description="Monthly Tuition Fee (September 2026)",
                amount=Decimal('2500.00')
            )
            FeeReceiptItem.objects.create(
                receipt=r1,
                category_name="Examination Fee",
                month_name="Term 1",
                description="Mid-Term Examination 2026 Assessment",
                amount=Decimal('1500.00')
            )
            FeeReceiptItem.objects.create(
                receipt=r1,
                category_name="Previous Balance Due",
                month_name="Arrears",
                description="Previous Outstanding Dues Settled",
                amount=Decimal('3500.00')
            )

@require_menu_permission('can_manage_fees', 'Fees & Financial Dashboard')
def fees_dashboard_view(request):
    ctx = get_school_context(request)
    ensure_fee_structure()
    active_session = get_active_session(request)
    today = timezone.now().date()

    # Invoices & General Billing
    invoices = FeeInvoice.objects.all()
    if active_session:
        invoices = invoices.filter(academic_session=active_session)

    total_invoiced = sum((inv.total_payable for inv in invoices), Decimal('0.00'))
    total_discount = sum((inv.discount for inv in invoices), Decimal('0.00'))
    total_fine = sum((inv.fine for inv in invoices), Decimal('0.00'))
    invoices_count = invoices.count()
    paid_invoices_count = invoices.filter(status='Paid').count()
    unpaid_invoices_count = invoices.filter(status__in=['Unpaid', 'Overdue', 'Partially Paid']).count()

    # Payments
    payments = FeePayment.objects.select_related('invoice__student__class_level', 'invoice__fee_category').all()
    if active_session:
        payments = payments.filter(invoice__academic_session=active_session)

    total_collected = payments.aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
    pending_total = max(Decimal('0.00'), total_invoiced - total_collected)
    recovery_rate = round(float(total_collected / total_invoiced * 100), 1) if total_invoiced > 0 else 0

    # Today & Current Month collections
    today_payments = payments.filter(payment_date=today)
    today_collected = today_payments.aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
    today_txns = today_payments.count()

    month_start = today.replace(day=1)
    month_payments = payments.filter(payment_date__gte=month_start, payment_date__lte=today)
    month_collected = month_payments.aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
    month_name = today.strftime('%B %Y')

    # 15-Day Trend Line Data
    trend_15_dates = [today - datetime.timedelta(days=i) for i in range(14, -1, -1)]
    trend_15_labels = [d.strftime('%d %b') for d in trend_15_dates]
    trend_15_data = []
    for d in trend_15_dates:
        amt = payments.filter(payment_date=d).aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
        trend_15_data.append(float(amt))

    # 30-Day Trend Data
    trend_30_dates = [today - datetime.timedelta(days=i) for i in range(29, -1, -1)]
    trend_30_labels = [d.strftime('%d %b') for d in trend_30_dates]
    trend_30_data = []
    for d in trend_30_dates:
        amt = payments.filter(payment_date=d).aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
        trend_30_data.append(float(amt))

    # Payment Mode Breakdown (Donut Chart)
    mode_map = {}
    for p in payments:
        m = p.payment_method or 'Cash'
        mode_map[m] = mode_map.get(m, Decimal('0.00')) + p.amount_paid
    mode_labels = list(mode_map.keys())
    mode_data = [float(v) for v in mode_map.values()]

    # Fee Category Breakdown (Donut / Pie Chart)
    cat_map = {}
    for inv in invoices:
        cname = inv.fee_category.name if inv.fee_category else "General Tuition"
        cat_map[cname] = cat_map.get(cname, Decimal('0.00')) + inv.total_payable
    cat_labels = list(cat_map.keys())
    cat_data = [float(v) for v in cat_map.values()]

    # Class-wise recovery
    class_stats = []
    classes = ClassLevel.objects.all().order_by('order')
    for cl in classes:
        cl_inv = sum((inv.total_payable for inv in invoices if inv.student.class_level_id == cl.id), Decimal('0.00'))
        cl_col = sum((p.amount_paid for p in payments if p.invoice.student.class_level_id == cl.id), Decimal('0.00'))
        cl_due = max(Decimal('0.00'), cl_inv - cl_col)
        cl_rate = round(float(cl_col / cl_inv * 100), 1) if cl_inv > 0 else 0
        class_stats.append({
            'class_name': cl.name,
            'invoiced': float(cl_inv),
            'collected': float(cl_col),
            'due': float(cl_due),
            'rate': cl_rate
        })

    # Recent Transactions Feed (Last 8)
    recent_payments = payments.order_by('-payment_date', '-created_at')[:8]

    # Top Defaulters (Top 6)
    students = Student.objects.filter(is_active=True).select_related('class_level', 'section')
    if active_session:
        students = students.filter(
            Q(academic_session=active_session) | Q(session_enrollments__academic_session=active_session)
        ).distinct()
    defaulters_list = []
    for st in students:
        due = st.fee_due_amount
        if due > 0:
            defaulters_list.append({
                'student': st,
                'due': due
            })
    defaulters_list.sort(key=lambda x: x['due'], reverse=True)
    top_defaulters = defaulters_list[:6]
    defaulters_count = len(defaulters_list)

    ctx.update({
        'active_session': active_session,
        'today': today,
        'total_invoiced': total_invoiced,
        'total_collected': total_collected,
        'total_discount': total_discount,
        'total_fine': total_fine,
        'pending_total': pending_total,
        'recovery_rate': recovery_rate,
        'invoices_count': invoices_count,
        'paid_invoices_count': paid_invoices_count,
        'unpaid_invoices_count': unpaid_invoices_count,
        'today_collected': today_collected,
        'today_txns': today_txns,
        'month_collected': month_collected,
        'month_name': month_name,
        'defaulters_count': defaulters_count,
        # Chart JSONs
        'trend_15_labels_json': json.dumps(trend_15_labels),
        'trend_15_data_json': json.dumps(trend_15_data),
        'trend_30_labels_json': json.dumps(trend_30_labels),
        'trend_30_data_json': json.dumps(trend_30_data),
        'mode_labels_json': json.dumps(mode_labels),
        'mode_data_json': json.dumps(mode_data),
        'cat_labels_json': json.dumps(cat_labels),
        'cat_data_json': json.dumps(cat_data),
        'class_stats': class_stats,
        'recent_payments': recent_payments,
        'top_defaulters': top_defaulters,
    })
    return render(request, 'school_erp/fees_dashboard.html', ctx)


@require_menu_permission('can_manage_fees', 'Fees & Accounting')
def fees_list_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    ensure_fee_structure(school=school)
    active_session = get_active_session(request, school=school)

    active_tab = request.GET.get('tab', 'invoices')
    status_filter = request.GET.get('status', '')
    class_id = request.GET.get('class_id', '')
    category_id = request.GET.get('category_id', '')
    search = request.GET.get('q', '').strip()

    invoices = FeeInvoice.objects.filter(school=school).select_related('student__class_level', 'student__section', 'fee_category').prefetch_related('payments').all()
    if active_session:
        invoices = invoices.filter(academic_session=active_session)

    if status_filter:
        invoices = invoices.filter(status=status_filter)
    if class_id:
        invoices = invoices.filter(student__class_level_id=class_id)
    if category_id:
        invoices = invoices.filter(fee_category_id=category_id)
    if search:
        invoices = invoices.filter(
            Q(invoice_number__icontains=search) |
            Q(student__first_name__icontains=search) |
            Q(student__last_name__icontains=search) |
            Q(student__admission_number__icontains=search) |
            Q(student__parent_phone__icontains=search)
        )

    all_session_invoices = FeeInvoice.objects.filter(school=school, academic_session=active_session) if active_session else FeeInvoice.objects.filter(school=school)
    total_invoiced = sum((inv.total_payable for inv in all_session_invoices), Decimal('0.00'))
    payments_qs = FeePayment.objects.filter(invoice__school=school)
    if active_session:
        payments_qs = payments_qs.filter(invoice__academic_session=active_session)
    total_collected = payments_qs.aggregate(total=Sum('amount_paid'))['total'] or Decimal('0.00')
    pending_total = max(Decimal('0.00'), total_invoiced - total_collected)

    # Payments Transaction History
    payments = FeePayment.objects.filter(invoice__school=school).select_related('invoice__student__class_level', 'invoice__fee_category').order_by('-payment_date', '-created_at')
    if active_session:
        payments = payments.filter(invoice__academic_session=active_session)
    if search:
        payments = payments.filter(
            Q(receipt_number__icontains=search) |
            Q(invoice__student__first_name__icontains=search) |
            Q(invoice__student__last_name__icontains=search) |
            Q(invoice__student__admission_number__icontains=search) |
            Q(transaction_reference__icontains=search)
        )
    if class_id:
        payments = payments.filter(invoice__student__class_level_id=class_id)

    # Student-wise Ledger
    students_ledger = []
    student_qs = Student.objects.filter(school=school, is_active=True).select_related('class_level', 'section').prefetch_related('invoices__payments')
    if active_session:
        student_qs = student_qs.filter(
            Q(academic_session=active_session) | Q(session_enrollments__academic_session=active_session)
        ).distinct()
    if search:
        student_qs = student_qs.filter(
            Q(first_name__icontains=search) |
            Q(last_name__icontains=search) |
            Q(admission_number__icontains=search) |
            Q(parent_phone__icontains=search)
        )
    if class_id:
        student_qs = student_qs.filter(class_level_id=class_id)

    for st in student_qs:
        st_invoices = st.invoices.filter(academic_session=active_session) if active_session else st.invoices.all()
        st_invoiced = sum((inv.total_payable for inv in st_invoices), Decimal('0.00'))
        st_paid = Decimal('0.00')
        st_payments_count = 0
        for inv in st_invoices:
            for p in inv.payments.all():
                st_paid += p.amount_paid
                st_payments_count += 1
        st_due = max(Decimal('0.00'), st_invoiced - st_paid)
        students_ledger.append({
            'student': st,
            'invoiced': st_invoiced,
            'paid': st_paid,
            'due': st_due,
            'payments_count': st_payments_count,
            'invoices_count': st_invoices.count(),
        })

    classes = ClassLevel.objects.filter(school=school).order_by('order', 'name')
    categories = FeeCategory.objects.filter(school=school)

    unified_receipts = FeeReceipt.objects.filter(school=school).select_related('student__class_level', 'student__section').prefetch_related('items').order_by('-payment_date', '-created_at')
    if active_session:
        unified_receipts = unified_receipts.filter(student__in=student_qs)
    if search:
        unified_receipts = unified_receipts.filter(
            Q(receipt_number__icontains=search) |
            Q(student__first_name__icontains=search) |
            Q(student__last_name__icontains=search) |
            Q(student__admission_number__icontains=search)
        )
    if class_id:
        unified_receipts = unified_receipts.filter(student__class_level_id=class_id)

    ctx.update({
        'invoices': invoices,
        'payments': payments,
        'unified_receipts': unified_receipts,
        'students_ledger': students_ledger,
        'classes': classes,
        'categories': categories,
        'active_tab': active_tab,
        'status_filter': status_filter,
        'selected_class': class_id,
        'selected_category': category_id,
        'search': search,
        'total_invoiced': total_invoiced,
        'total_collected': total_collected,
        'pending_total': pending_total,
    })
    return render(request, 'school_erp/fees_list.html', ctx)


def fee_invoice_create_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    ensure_fee_structure(school=school)
    
    if request.method == 'POST':
        form = FeeInvoiceForm(request.POST)
        if form.is_valid():
            invoice = form.save(commit=False)
            invoice.school = school
            invoice.save()
            invoice.update_status()
            messages.success(request, f"New Fee Invoice '{invoice.invoice_number}' created for {invoice.student.full_name} ({invoice.fee_category.name}) - Total Amount: ₹{invoice.total_payable}")
            return redirect('fees_list')
    else:
        count = FeeInvoice.objects.filter(school=school).count() + 1
        inv_no = f"INV-{timezone.now().year}-{count:04d}"
        default_cat = FeeCategory.objects.filter(school=school, name__icontains='Monthly').first() or FeeCategory.objects.filter(school=school).first()
        form = FeeInvoiceForm(initial={
            'invoice_number': inv_no,
            'fee_category': default_cat,
            'amount': default_cat.default_amount if default_cat else Decimal('2500.00'),
            'due_date': timezone.now().date() + datetime.timedelta(days=15),
        })

    ctx.update({
        'form': form,
        'title': 'Issue New Fee Invoice',
    })
    return render(request, 'school_erp/fee_invoice_form.html', ctx)


def fee_collect_view(request, invoice_id):
    ctx = get_school_context(request)
    school = ctx['school']
    invoice = get_object_or_404(FeeInvoice.objects.select_related('student', 'fee_category'), pk=invoice_id, school=school)
    
    if request.method == 'POST':
        form = FeePaymentForm(request.POST)
        if form.is_valid():
            payment = form.save(commit=False)
            payment.invoice = invoice
            
            if payment.amount_paid <= Decimal('0.00'):
                messages.error(request, "Amount paid must be greater than ₹0.")
                return render(request, 'school_erp/fee_collect.html', {'invoice': invoice, 'form': form, **ctx})
            
            if payment.amount_paid > invoice.remaining_due:
                messages.error(request, f"Entered amount (₹{payment.amount_paid}) cannot be greater than remaining balance due (₹{invoice.remaining_due})!")
                return render(request, 'school_erp/fee_collect.html', {'invoice': invoice, 'form': form, **ctx})
                
            payment.save()
            invoice.update_status()
            
            if invoice.remaining_due > Decimal('0.00'):
                messages.success(request, f"Partial installment of ₹{payment.amount_paid} recorded! Remaining balance due: ₹{invoice.remaining_due}. Receipt: {payment.receipt_number}")
            else:
                messages.success(request, f"Full payment of ₹{payment.amount_paid} recorded! Invoice completely settled. Receipt: {payment.receipt_number}")
                
            return redirect('fee_receipt', payment_id=payment.pk)
    else:
        count = FeePayment.objects.filter(invoice__school=school).count() + 1
        rec_no = f"REC-{timezone.now().year}-{count:04d}"
        form = FeePaymentForm(initial={
            'receipt_number': rec_no,
            'amount_paid': invoice.remaining_due,
            'payment_date': timezone.now().date(),
        })
        
    ctx.update({
        'invoice': invoice,
        'form': form,
    })
    return render(request, 'school_erp/fee_collect.html', ctx)


def amount_to_words(n):
    units = ["", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine",
             "Ten", "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen",
             "Seventeen", "Eighteen", "Nineteen"]
    tens = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]
    try:
        val = int(round(float(n)))
    except (ValueError, TypeError):
        return ""
    if val == 0:
        return "Zero Rupees Only"
    def convert(num):
        if num < 20:
            return units[num]
        elif num < 100:
            return tens[num // 10] + (" " + units[num % 10] if num % 10 != 0 else "")
        elif num < 1000:
            return units[num // 100] + " Hundred" + (" and " + convert(num % 100) if num % 100 != 0 else "")
        elif num < 100000:
            return convert(num // 1000) + " Thousand" + (" " + convert(num % 1000) if num % 1000 != 0 else "")
        elif num < 10000000:
            return convert(num // 100000) + " Lakh" + (" " + convert(num % 100000) if num % 100000 != 0 else "")
        else:
            return convert(num // 10000000) + " Crore" + (" " + convert(num % 10000000) if num % 10000000 != 0 else "")
    return convert(val) + " Rupees Only"


def clean_phone_for_wa(phone_val):
    if not phone_val:
        return ""
    digits = "".join(filter(str.isdigit, str(phone_val)))
    if len(digits) == 10:
        return f"91{digits}"
    elif len(digits) == 12 and digits.startswith("91"):
        return digits
    elif len(digits) > 10:
        return digits
    return digits


def fee_receipt_view(request, payment_id):
    ctx = get_school_context(request)
    payment = get_object_or_404(FeePayment.objects.select_related('invoice__student', 'invoice__fee_category'), pk=payment_id)
    invoice = payment.invoice
    student = invoice.student
    school = ctx['school']
    
    amount_in_words = amount_to_words(payment.amount_paid)
    clean_phone = clean_phone_for_wa(student.parent_phone)
    public_slip_url = request.build_absolute_uri(f"/slip/payment/{payment.receipt_number}/")

    wa_message = (
        f"*{school.name.upper()}*\n"
        f"*Official Fee Payment Receipt*\n"
        f"-----------------------------------------\n"
        f"Receipt No: *{payment.receipt_number}*\n"
        f"Date: *{payment.payment_date.strftime('%d-%b-%Y')}*\n"
        f"Student: *{student.full_name}*\n"
        f"Adm No: *{student.admission_number}* | Class: *{student.class_level.name}-{student.section.name}*\n"
        f"Father: *{student.parent_name}*\n"
        f"-----------------------------------------\n"
        f"Category: {invoice.fee_category.name}\n"
        f"Amount Paid: *Rs.{payment.amount_paid:,.0f}*\n"
        f"Remaining Due: *Rs.{invoice.remaining_due:,.0f}*\n"
        f"Payment Mode: {payment.payment_method}\n"
        f"-----------------------------------------\n"
        f"📄 *View / Download Official Slip Online (No Login Needed):*\n"
        f"{public_slip_url}\n"
        f"-----------------------------------------\n"
        f"Thank you for depositing the school fee.\n"
        f"Accounts Office, {school.name}"
    )

    ctx.update({
        'payment': payment,
        'invoice': invoice,
        'student': student,
        'amount_in_words': amount_in_words,
        'clean_parent_phone': clean_phone,
        'wa_message': wa_message,
        'public_slip_url': public_slip_url,
    })
    return render(request, 'school_erp/fee_receipt.html', ctx)


def get_student_fee_details(student):
    academic_months = [
        {"id": "apr", "name": "April 2026", "short": "Apr", "hindi": "", "qtr": "Q1"},
        {"id": "may", "name": "May 2026", "short": "May", "hindi": "", "qtr": "Q1"},
        {"id": "jun", "name": "June 2026", "short": "Jun", "hindi": "", "qtr": "Q1"},
        {"id": "jul", "name": "July 2026", "short": "Jul", "hindi": "", "qtr": "Q2"},
        {"id": "aug", "name": "August 2026", "short": "Aug", "hindi": "", "qtr": "Q2"},
        {"id": "sep", "name": "September 2026", "short": "Sep", "hindi": "", "qtr": "Q2"},
        {"id": "oct", "name": "October 2026", "short": "Oct", "hindi": "", "qtr": "Q3"},
        {"id": "nov", "name": "November 2026", "short": "Nov", "hindi": "", "qtr": "Q3"},
        {"id": "dec", "name": "December 2026", "short": "Dec", "hindi": "", "qtr": "Q3"},
        {"id": "jan", "name": "January 2027", "short": "Jan", "hindi": "", "qtr": "Q4"},
        {"id": "feb", "name": "February 2027", "short": "Feb", "hindi": "", "qtr": "Q4"},
        {"id": "mar", "name": "March 2027", "short": "Mar", "hindi": "", "qtr": "Q4"},
    ]

    receipt_items = FeeReceiptItem.objects.filter(
        receipt__student=student,
        category_name__icontains='Monthly'
    ).select_related('receipt')

    paid_monthly_invoices = student.invoices.filter(
        fee_category__name__icontains='Monthly',
        status='Paid'
    )

    paid_map = {}
    for item in receipt_items:
        m_str = (item.month_name or item.description or '').lower()
        for m in academic_months:
            if m['name'].lower() in m_str or m['short'].lower() in m_str:
                paid_map[m['name']] = {
                    'receipt_no': item.receipt.receipt_number,
                    'paid_date': item.receipt.payment_date.strftime("%d %b %Y"),
                    'amount': item.amount,
                    'receipt_id': item.receipt.id
                }

    for inv in paid_monthly_invoices:
        inv_str = (inv.invoice_number or '').lower()
        for m in academic_months:
            if m['name'].lower() in inv_str or m['short'].lower() in inv_str:
                if m['name'] not in paid_map:
                    paid_map[m['name']] = {
                        'receipt_no': inv.invoice_number,
                        'paid_date': inv.created_at.strftime("%d %b %Y"),
                        'amount': inv.amount,
                        'receipt_id': None
                    }

    months_status = []
    paid_months_count = 0
    unpaid_months_count = 0
    first_unpaid_idx = None

    for idx, m in enumerate(academic_months):
        m_name = m['name']
        if m_name in paid_map:
            p = paid_map[m_name]
            months_status.append({
                'id': m['id'],
                'name': m_name,
                'short': m['short'],
                'hindi': m['hindi'],
                'qtr': m['qtr'],
                'is_paid': True,
                'receipt_no': p['receipt_no'],
                'paid_date': p['paid_date'],
                'paid_amount': float(p['amount']),
                'receipt_id': p.get('receipt_id'),
            })
            paid_months_count += 1
        else:
            if first_unpaid_idx is None:
                first_unpaid_idx = idx
            months_status.append({
                'id': m['id'],
                'name': m_name,
                'short': m['short'],
                'hindi': m['hindi'],
                'qtr': m['qtr'],
                'is_paid': False,
                'receipt_no': None,
                'paid_date': None,
                'paid_amount': None,
                'receipt_id': None,
            })
            unpaid_months_count += 1

    # Check Exam Fees
    exam_items = FeeReceiptItem.objects.filter(
        receipt__student=student,
        category_name__icontains='Exam'
    ).select_related('receipt')

    exam_paid_map = {}
    for item in exam_items:
        desc = (item.description or item.category_name or '').lower()
        if 'mid' in desc or 'term 1' in desc or 'term-1' in desc or 'first' in desc:
            exam_paid_map['Mid-Term Examination 2026'] = {
                'receipt_no': item.receipt.receipt_number,
                'paid_date': item.receipt.payment_date.strftime("%d %b %Y"),
                'amount': float(item.amount),
                'receipt_id': item.receipt.id
            }
        elif 'final' in desc or 'annual' in desc or 'term 2' in desc:
            exam_paid_map['Annual / Final Examination 2027'] = {
                'receipt_no': item.receipt.receipt_number,
                'paid_date': item.receipt.payment_date.strftime("%d %b %Y"),
                'amount': float(item.amount),
                'receipt_id': item.receipt.id
            }
        else:
            exam_paid_map['Mid-Term Examination 2026'] = {
                'receipt_no': item.receipt.receipt_number,
                'paid_date': item.receipt.payment_date.strftime("%d %b %Y"),
                'amount': float(item.amount),
                'receipt_id': item.receipt.id
            }

    exams_status = [
        {
            'name': 'Mid-Term Examination 2026',
            'is_paid': 'Mid-Term Examination 2026' in exam_paid_map,
            'paid_info': exam_paid_map.get('Mid-Term Examination 2026'),
            'default_amount': 1500.0,
        },
        {
            'name': 'Annual / Final Examination 2027',
            'is_paid': 'Annual / Final Examination 2027' in exam_paid_map,
            'paid_info': exam_paid_map.get('Annual / Final Examination 2027'),
            'default_amount': 1500.0,
        }
    ]

    # Check Physical & Optional Item Purchases History
    # Supported: Tie & Belt, Books, Uniform / Dress, School Diary, Student ID Card
    item_keywords = {
        'books': ['book set', 'books', 'notebook'],
        'dress': ['uniform', 'dress'],
        'tie_belt': ['tie & belt', 'tie and belt', 'tie', 'belt'],
        'diary': ['diary', 'almanac'],
        'id_card': ['id card', 'smart id'],
    }

    item_history = {k: [] for k in item_keywords}
    receipt_items = FeeReceiptItem.objects.filter(
        receipt__student=student,
        receipt__status='Active'
    ).select_related('receipt').order_by('receipt__payment_date', 'receipt__id')

    for item in receipt_items:
        c_name = (item.category_name or '').lower()
        d_name = (item.description or '').lower()
        combined = f"{c_name} {d_name}"
        for key, kws in item_keywords.items():
            if any(kw in combined for kw in kws):
                item_history[key].append({
                    'receipt_no': item.receipt.receipt_number,
                    'receipt_id': item.receipt.id,
                    'paid_date': item.receipt.payment_date.strftime("%d %b %Y"),
                    'amount': float(item.amount),
                    'description': item.description,
                    'purchase_num': len(item_history[key]) + 1
                })
                break

    item_purchases = {}
    for k, hist in item_history.items():
        item_purchases[k] = {
            'purchased': len(hist) > 0,
            'count': len(hist),
            'next_num': len(hist) + 1,
            'history': hist,
            'last': hist[-1] if hist else None,
        }

    unpaid_invoices = student.invoices.filter(status__in=['Unpaid', 'Partially Paid', 'Overdue'])
    previous_due = sum((inv.remaining_due for inv in unpaid_invoices), Decimal('0.00'))
    receipts_history = FeeReceipt.objects.filter(student=student).order_by('-payment_date', '-created_at')[:6]

    return {
        'months_status': months_status,
        'paid_months_count': paid_months_count,
        'unpaid_months_count': unpaid_months_count,
        'first_unpaid_idx': first_unpaid_idx,
        'exams_status': exams_status,
        'item_purchases': item_purchases,
        'previous_due': previous_due,
        'receipts_history': receipts_history,
    }


@require_menu_permission('can_manage_fees', 'Fee Collection Counter')
def fee_counter_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    ensure_fee_structure(school=school)
    active_session = get_active_session(request, school=school)

    students = Student.objects.filter(school=school, is_active=True).select_related('class_level', 'section')
    if active_session:
        students = students.filter(
            Q(academic_session=active_session) | Q(session_enrollments__academic_session=active_session)
        ).distinct()
    students = students.order_by('class_level__order', 'roll_number')
    
    # Students data for quick client-side instant search (no clunky dropdown)
    students_data = [
        {
            'id': s.id,
            'name': s.full_name,
            'adm_no': s.admission_number,
            'class_name': f"{s.class_level.name} - Sec {s.section.name}",
            'class_id': s.class_level.id,
            'roll': s.roll_number,
            'father': s.parent_name or '',
            'phone': s.parent_phone or '',
            'due': float(s.fee_due_amount),
        }
        for s in students
    ]
    students_json = json.dumps(students_data)

    selected_student_id = request.GET.get('student_id') or request.POST.get('student_id')
    selected_student = None
    fee_details = None
    previous_due = Decimal('0.00')
    fee_master = None

    if selected_student_id:
        selected_student = Student.objects.filter(id=selected_student_id).first()
        if selected_student:
            fee_details = get_student_fee_details(selected_student)
            previous_due = fee_details['previous_due']
            try:
                fee_master = selected_student.class_level.fee_master
            except Exception:
                fee_master = None

    if fee_master:
        default_monthly_amt = fee_master.monthly_tuition_fee
        default_exam_amt = fee_master.exam_fee
        default_transport_amt = fee_master.transport_fee
        default_admission_amt = fee_master.admission_fee
        default_book_set_amt = fee_master.book_set_fee
        default_dress_amt = fee_master.dress_uniform_fee
        default_tie_belt_amt = fee_master.tie_belt_fee
        default_diary_amt = fee_master.diary_fee
        default_id_card_amt = fee_master.id_card_fee
        default_other_amt = fee_master.other_activity_fee
    else:
        monthly_cat = FeeCategory.objects.filter(name__icontains='Monthly').first()
        default_monthly_amt = monthly_cat.default_amount if monthly_cat else Decimal('2000.00')
        default_exam_amt = Decimal('1200.00')
        default_transport_amt = Decimal('1200.00')
        default_admission_amt = Decimal('5000.00')
        default_book_set_amt = Decimal('2500.00')
        default_dress_amt = Decimal('1800.00')
        default_tie_belt_amt = Decimal('250.00')
        default_diary_amt = Decimal('150.00')
        default_id_card_amt = Decimal('100.00')
        default_other_amt = Decimal('500.00')

    classes = ClassLevel.objects.filter(school=school).order_by('order')

    if request.method == 'POST' and 'submit_payment' in request.POST:
        if not selected_student:
            messages.error(request, "Please select a student first!")
            return redirect('fee_counter')

        payment_method = request.POST.get('payment_method', 'Cash')
        transaction_reference = request.POST.get('transaction_reference', '').strip()
        notes = request.POST.get('notes', '').strip()
        
        items_to_create = []
        subtotal = Decimal('0.00')
        prev_due_cleared = Decimal('0.00')

        # 1. Multi-Month Selection: Consolidate into ONE single receipt row with months listed below
        selected_months = request.POST.getlist('selected_months')
        monthly_rate = Decimal(request.POST.get('monthly_rate', str(default_monthly_amt)) or str(default_monthly_amt))
        if selected_months:
            total_monthly_amount = monthly_rate * len(selected_months)
            months_text = ", ".join(selected_months)
            period_text = f"{selected_months[0]} to {selected_months[-1]}" if len(selected_months) > 1 else selected_months[0]
            items_to_create.append({
                'category_name': 'Monthly Tuition Fee',
                'month_name': period_text,
                'description': f"Months: {months_text} (Total {len(selected_months)} Month{'s' if len(selected_months) > 1 else ''})",
                'amount': total_monthly_amount,
                'is_monthly': True,
                'months_list': selected_months,
            })
            subtotal += total_monthly_amount

        # 2. Examination Fee
        selected_exams = request.POST.getlist('selected_exams')
        for e_name in selected_exams:
            e_amt = Decimal(request.POST.get(f'exam_amt_{e_name}', str(default_exam_amt)) or str(default_exam_amt))
            items_to_create.append({
                'category_name': 'Examination Fee',
                'month_name': e_name,
                'description': f"{e_name} Fee",
                'amount': e_amt
            })
            subtotal += e_amt

        # 3. Book Set (1st Purchase or Buy Again / Replacement)
        if request.POST.get('include_books') == 'on':
            b_amount = Decimal(request.POST.get('book_set_amount', str(default_book_set_amt)) or str(default_book_set_amt))
            is_rep = request.POST.get('is_replacement_books') == '1'
            rep_num = request.POST.get('rep_num_books', '2')
            rep_reason = request.POST.get('reason_books', '').strip()
            if is_rep:
                desc = f"Complete Academic Books Pack (Replacement #{rep_num}{' - ' + rep_reason if rep_reason else ''})"
                m_name = f"Replacement #{rep_num}"
            else:
                desc = "Complete Academic Books & Notebooks Set"
                m_name = "Session 2026-27"
            items_to_create.append({
                'category_name': 'Book Set',
                'month_name': m_name,
                'description': desc,
                'amount': b_amount
            })
            subtotal += b_amount

        # 4. Uniform / Dress (1st Purchase or Buy Again / Replacement)
        if request.POST.get('include_dress') == 'on':
            d_amount = Decimal(request.POST.get('dress_amount', str(default_dress_amt)) or str(default_dress_amt))
            is_rep = request.POST.get('is_replacement_dress') == '1'
            rep_num = request.POST.get('rep_num_dress', '2')
            rep_reason = request.POST.get('reason_dress', '').strip()
            if is_rep:
                desc = f"School Uniform & Dress Set (Replacement #{rep_num}{' - ' + rep_reason if rep_reason else ''})"
                m_name = f"Replacement #{rep_num}"
            else:
                desc = "School Uniform & Dress Set"
                m_name = "Session 2026-27"
            items_to_create.append({
                'category_name': 'School Uniform / Dress',
                'month_name': m_name,
                'description': desc,
                'amount': d_amount
            })
            subtotal += d_amount

        # 5. Tie & Belt (1st Purchase or Buy Again / Replacement)
        if request.POST.get('include_tie_belt') == 'on':
            tb_amount = Decimal(request.POST.get('tie_belt_amount', str(default_tie_belt_amt)) or str(default_tie_belt_amt))
            is_rep = request.POST.get('is_replacement_tie_belt') == '1'
            rep_num = request.POST.get('rep_num_tie_belt', '2')
            rep_reason = request.POST.get('reason_tie_belt', '').strip()
            if is_rep:
                desc = f"School Logo Tie & Belt (Replacement #{rep_num}{' - ' + rep_reason if rep_reason else ''})"
                m_name = f"Replacement #{rep_num}"
            else:
                desc = "School Logo Tie & Belt"
                m_name = "Session 2026-27"
            items_to_create.append({
                'category_name': 'Tie & Belt',
                'month_name': m_name,
                'description': desc,
                'amount': tb_amount
            })
            subtotal += tb_amount

        # 6. School Diary (1st Purchase or Buy Again / Replacement)
        if request.POST.get('include_diary') == 'on':
            dia_amount = Decimal(request.POST.get('diary_amount', str(default_diary_amt)) or str(default_diary_amt))
            is_rep = request.POST.get('is_replacement_diary') == '1'
            rep_num = request.POST.get('rep_num_diary', '2')
            rep_reason = request.POST.get('reason_diary', '').strip()
            if is_rep:
                desc = f"Student Academic Diary / Almanac (Replacement #{rep_num}{' - ' + rep_reason if rep_reason else ''})"
                m_name = f"Replacement #{rep_num}"
            else:
                desc = "Student Academic Diary / Almanac"
                m_name = "Session 2026-27"
            items_to_create.append({
                'category_name': 'School Diary',
                'month_name': m_name,
                'description': desc,
                'amount': dia_amount
            })
            subtotal += dia_amount

        # 7. Student ID Card (1st Purchase or Buy Again / Replacement)
        if request.POST.get('include_id_card') == 'on':
            id_amount = Decimal(request.POST.get('id_card_amount', str(default_id_card_amt)) or str(default_id_card_amt))
            is_rep = request.POST.get('is_replacement_id_card') == '1'
            rep_num = request.POST.get('rep_num_id_card', '2')
            rep_reason = request.POST.get('reason_id_card', '').strip()
            if is_rep:
                desc = f"Student Smart ID Card & Lanyard (Duplicate #{rep_num}{' - ' + rep_reason if rep_reason else ''})"
                m_name = f"Duplicate #{rep_num}"
            else:
                desc = "Student Smart ID Card & Lanyard"
                m_name = "Session 2026-27"
            items_to_create.append({
                'category_name': 'Student ID Card',
                'month_name': m_name,
                'description': desc,
                'amount': id_amount
            })
            subtotal += id_amount

        # 8. Transport Fee
        if request.POST.get('include_transport') == 'on':
            t_amount = Decimal(request.POST.get('transport_amount', str(default_transport_amt)) or str(default_transport_amt))
            items_to_create.append({
                'category_name': 'Transport / Bus Fee',
                'month_name': 'Monthly Facility',
                'description': "School Bus & Safe Transport Facility",
                'amount': t_amount
            })
            subtotal += t_amount

        # 9. Admission Fee
        if request.POST.get('include_admission') == 'on':
            adm_amount = Decimal(request.POST.get('admission_amount', str(default_admission_amt)) or str(default_admission_amt))
            items_to_create.append({
                'category_name': 'Admission Fee',
                'month_name': '',
                'description': "New Student Admission & Registration Fee",
                'amount': adm_amount
            })
            subtotal += adm_amount

        # 10. Other Fee
        if request.POST.get('include_other') == 'on':
            oth_amount = Decimal(request.POST.get('other_amount', str(default_other_amt)) or str(default_other_amt))
            oth_desc = request.POST.get('other_description', 'Other Activity Fee').strip() or 'Other Fee'
            items_to_create.append({
                'category_name': 'Other / Miscellaneous Fee',
                'month_name': '',
                'description': oth_desc,
                'amount': oth_amount
            })
            subtotal += oth_amount

        # 11. Previous Due Balance Cleared Today
        if request.POST.get('include_previous_due') == 'on':
            prev_due_cleared = Decimal(request.POST.get('previous_due_cleared_amount', '0.00') or '0.00')
            if prev_due_cleared > Decimal('0.00'):
                items_to_create.append({
                    'category_name': 'Previous Balance Due',
                    'month_name': 'Arrears',
                    'description': "Previous Outstanding Dues Settled",
                    'amount': prev_due_cleared
                })

        # Duplicate Transaction Reference Detection (Prevents duplicate callbacks / double entry)
        if transaction_reference:
            existing_rec = FeeReceipt.objects.filter(transaction_reference=transaction_reference, status='Active').first()
            if existing_rec:
                messages.error(
                    request,
                    f"Duplicate Transaction Reference Detected! Receipt #{existing_rec.receipt_number} has already been recorded with Reference '{transaction_reference}' for {existing_rec.student.full_name} on {existing_rec.payment_date}."
                )
                return redirect(f"{request.path}?student_id={selected_student.id}")

        discount = Decimal(request.POST.get('discount', '0.00') or '0.00')
        concession_type = request.POST.get('concession_type', '').strip()
        fine_amount = Decimal(request.POST.get('fine_amount', '0.00') or '0.00')
        fine_waived = Decimal(request.POST.get('fine_waived', '0.00') or '0.00')
        
        total_billed = max(Decimal('0.00'), (subtotal + prev_due_cleared + fine_amount) - (discount + fine_waived))
        
        amount_paying = Decimal(request.POST.get('amount_paying_now', '0.00') or '0.00')
        if amount_paying <= Decimal('0.00'):
            messages.error(request, "Amount paying now must be greater than ₹0.")
            return redirect(f"{request.path}?student_id={selected_student.id}")

        if not items_to_create and prev_due_cleared <= Decimal('0.00'):
            messages.error(request, "Please select at least one month, exam fee, school item, or previous due to collect!")
            return redirect(f"{request.path}?student_id={selected_student.id}")

        # Correct Remaining Balance Calculation (Discount is SUBTRACTED, NEVER added to due):
        net_new_charges = max(Decimal('0.00'), (subtotal + fine_amount) - (discount + fine_waived))
        cash_for_old = min(amount_paying, prev_due_cleared) if prev_due_cleared > Decimal('0.00') else Decimal('0.00')
        cash_for_new = max(Decimal('0.00'), amount_paying - cash_for_old)
        
        unpaid_new = max(Decimal('0.00'), net_new_charges - cash_for_new)
        unpaid_old = max(Decimal('0.00'), previous_due - cash_for_old)
        remaining_balance_due = unpaid_old + unpaid_new

        receipt_count = FeeReceipt.objects.filter(school=school).count() + 1
        sch_prefix = "".join([w[0] for w in (school.name or "").split() if w.isalnum()])[:3].upper() or (school.school_code or "SCH")[:3].upper()
        rec_no = f"{sch_prefix}-REC-{timezone.now().year}-{receipt_count:04d}"

        unified_receipt = FeeReceipt.objects.create(
            school=school,
            receipt_number=rec_no,
            student=selected_student,
            payment_date=timezone.now().date(),
            payment_method=payment_method,
            transaction_reference=transaction_reference,
            notes=notes,
            subtotal=subtotal,
            previous_due_cleared=prev_due_cleared,
            discount=discount,
            total_billed=total_billed,
            amount_paid=amount_paying,
            remaining_balance_due=remaining_balance_due,
            concession_amount=discount,
            concession_type=concession_type,
            fine_amount=fine_amount,
            fine_waived=fine_waived,
            cashier=request.user if request.user.is_authenticated else None,
            status='Active'
        )

        for item in items_to_create:
            FeeReceiptItem.objects.create(
                receipt=unified_receipt,
                category_name=item['category_name'],
                month_name=item.get('month_name', ''),
                description=item['description'],
                amount=item['amount']
            )

        # Allocate amount_paying across invoices:
        cash_left = amount_paying

        # 1. Clear or credit student's pending past invoices if previous due was included
        if prev_due_cleared > Decimal('0.00'):
            unpaid_invoices = selected_student.invoices.filter(status__in=['Unpaid', 'Partially Paid', 'Overdue']).order_by('created_at')
            temp_credit = min(prev_due_cleared, cash_left)
            for inv in unpaid_invoices:
                if temp_credit <= Decimal('0.00'):
                    break
                rem = inv.remaining_due
                pay_for_this = min(temp_credit, rem)
                FeePayment.objects.create(
                    receipt_number=f"{rec_no}-OLD{inv.id}",
                    invoice=inv,
                    amount_paid=pay_for_this,
                    payment_date=timezone.now().date(),
                    payment_method=payment_method,
                    transaction_reference=f"Unified Receipt #{rec_no}",
                    notes=f"Settled via Unified Fee Counter (Receipt #{rec_no})"
                )
                inv.update_status()
                temp_credit -= pay_for_this
                cash_left -= pay_for_this

        # 2. Create invoice for each newly billed fee item and allocate remaining cash
        active_session = AcademicSession.objects.filter(is_current=True).first()
        rem_discount = discount

        for item in items_to_create:
            cat_name = item['category_name']
            if cat_name == 'Previous Balance Due':
                continue
            
            clean_cat_name = cat_name.split('/')[0].strip()
            fee_cat = FeeCategory.objects.filter(name__icontains=clean_cat_name).first()
            if not fee_cat:
                fee_cat, _ = FeeCategory.objects.get_or_create(
                    name=cat_name,
                    defaults={'default_amount': item['amount']}
                )
            
            # Apply discount to invoice!
            item_disc = min(rem_discount, item['amount'])
            rem_discount -= item_disc

            inv_count = FeeInvoice.objects.count() + 1
            inv_no = f"INV-{timezone.now().year}-{inv_count:04d}"
            new_inv = FeeInvoice.objects.create(
                invoice_number=inv_no,
                student=selected_student,
                fee_category=fee_cat,
                academic_session=active_session,
                amount=item['amount'],
                discount=item_disc,
                fine=Decimal('0.00'),
                due_date=timezone.now().date() + datetime.timedelta(days=15),
                status='Unpaid'
            )
            
            if cash_left > Decimal('0.00'):
                pay_now = min(cash_left, new_inv.remaining_due)
                FeePayment.objects.create(
                    receipt_number=f"{rec_no}-INV{new_inv.id}",
                    invoice=new_inv,
                    amount_paid=pay_now,
                    payment_date=timezone.now().date(),
                    payment_method=payment_method,
                    transaction_reference=f"Unified Receipt #{rec_no}",
                    notes=f"{item['description']} paid via Unified Counter (Receipt #{rec_no})"
                )
                cash_left -= pay_now
            new_inv.update_status()

        # 3. Post comprehensive entries into StudentFeeLedger
        active_sess = active_session or AcademicSession.objects.filter(is_current=True).first()
        cur_due = remaining_balance_due

        # Debits for newly charged fee items
        for item in items_to_create:
            StudentFeeLedger.objects.create(
                school=ctx.get('school'),
                student=selected_student,
                academic_session=active_sess,
                entry_date=timezone.now().date(),
                entry_type='Fee_Charge',
                voucher_no=rec_no,
                description=f"{item['category_name']} - {item['description']}",
                debit=item['amount'],
                credit=Decimal('0.00'),
                running_balance=cur_due,
                receipt=unified_receipt,
                created_by=request.user if request.user.is_authenticated else None
            )

        # Credit for Concession / Sibling Discount
        if discount > Decimal('0.00'):
            disc_label = f"Concession ({concession_type})" if concession_type else "Fee Concession / Discount"
            StudentFeeLedger.objects.create(
                school=ctx.get('school'),
                student=selected_student,
                academic_session=active_sess,
                entry_date=timezone.now().date(),
                entry_type='Concession',
                voucher_no=rec_no,
                description=disc_label,
                debit=Decimal('0.00'),
                credit=discount,
                running_balance=cur_due,
                receipt=unified_receipt,
                created_by=request.user if request.user.is_authenticated else None
            )

        # Credit for Payment Received
        StudentFeeLedger.objects.create(
            school=ctx.get('school'),
            student=selected_student,
            academic_session=active_sess,
            entry_date=timezone.now().date(),
            entry_type='Fee_Payment',
            voucher_no=rec_no,
            description=f"Fee Payment Received via {payment_method}" + (f" (Ref: {transaction_reference})" if transaction_reference else ""),
            debit=Decimal('0.00'),
            credit=amount_paying,
            running_balance=remaining_balance_due,
            receipt=unified_receipt,
            created_by=request.user if request.user.is_authenticated else None
        )

        # Audit Log Entry for Fee Payment Collection
        log_audit(
            request=request,
            action='CREATE',
            module='FEES',
            object_type='FeeReceipt',
            object_id=str(unified_receipt.id),
            object_repr=f"Receipt #{unified_receipt.receipt_number}",
            description=f"Fee payment of ₹{amount_paying} received for {selected_student.full_name} ({selected_student.admission_number}). Mode: {payment_method}. Remaining Due: ₹{remaining_balance_due}",
            reason=f"Fee Collection: {payment_method}"
        )

        messages.success(request, f"Official Fee Slip #{unified_receipt.receipt_number} generated successfully! Amount Received: ₹{amount_paying}")
        return redirect('unified_receipt', receipt_id=unified_receipt.id)

    # Check for student active concessions and siblings
    student_concessions = []
    sibling_students = []
    if selected_student:
        student_concessions = selected_student.concessions.filter(is_active=True)
        if selected_student.parent_phone:
            sibling_students = Student.objects.filter(
                is_active=True,
                parent_phone=selected_student.parent_phone
            ).exclude(id=selected_student.id).select_related('class_level', 'section')

    ctx.update({
        'students': students,
        'students_json': students_json,
        'classes': classes,
        'selected_student': selected_student,
        'student_concessions': student_concessions,
        'sibling_students': sibling_students,
        'fee_details': fee_details,
        'item_purchases': fee_details.get('item_purchases', {}) if fee_details else {},
        'fee_master': fee_master,
        'previous_due': previous_due,
        'default_monthly_amt': default_monthly_amt,
        'default_exam_amt': default_exam_amt,
        'default_transport_amt': default_transport_amt,
        'default_admission_amt': default_admission_amt,
        'default_book_set_amt': default_book_set_amt,
        'default_dress_amt': default_dress_amt,
        'default_tie_belt_amt': default_tie_belt_amt,
        'default_diary_amt': default_diary_amt,
        'default_id_card_amt': default_id_card_amt,
        'default_other_amt': default_other_amt,
    })
    return render(request, 'school_erp/fee_counter.html', ctx)


def unified_receipt_view(request, receipt_id):
    ctx = get_school_context(request)
    school = ctx['school']
    receipt = get_object_or_404(FeeReceipt.objects.select_related('student__class_level', 'student__section').prefetch_related('items'), pk=receipt_id, school=school)
    student = receipt.student
    
    amount_in_words = amount_to_words(receipt.amount_paid)

    # Format WhatsApp pre-filled receipt confirmation message
    wa_items_list = []
    for it in receipt.items.all():
        period = f" ({it.month_name})" if it.month_name else ""
        wa_items_list.append(f"• {it.category_name}{period}: Rs.{it.amount:,.0f}")
    wa_items_text = "\n".join(wa_items_list)

    public_slip_url = request.build_absolute_uri(f"/slip/receipt/{receipt.receipt_number}/")

    wa_message = (
        f"*{school.name.upper()}*\n"
        f"*Official Fee Payment Receipt*\n"
        f"-----------------------------------------\n"
        f"Receipt No: *{receipt.receipt_number}*\n"
        f"Date: *{receipt.payment_date.strftime('%d-%b-%Y')}*\n"
        f"Student: *{student.full_name}*\n"
        f"Adm No: *{student.admission_number}* | Class: *{student.class_level.name}-{student.section.name}*\n"
        f"Father: *{student.parent_name}*\n"
        f"-----------------------------------------\n"
        f"*Fee Details:*\n"
        f"{wa_items_text}\n"
        f"Total Billed: Rs.{receipt.total_billed:,.0f}\n"
        f"Amount Received: *Rs.{receipt.amount_paid:,.0f}*\n"
        f"Remaining Balance: *Rs.{receipt.remaining_balance_due:,.0f}*\n"
        f"Payment Mode: {receipt.payment_method}\n"
        f"-----------------------------------------\n"
        f"📄 *View / Download Official Slip Online (No Login Needed):*\n"
        f"{public_slip_url}\n"
        f"-----------------------------------------\n"
        f"Thank you for depositing the school fee.\n"
        f"Accounts Office, {school.name}"
    )

    sms_text = (
        f"Fee Received: {school.short_name} thanks you for fee payment of Rs.{receipt.amount_paid:,.0f} "
        f"for {student.full_name} (Adm:{student.admission_number}). "
        f"Rcpt:{receipt.receipt_number}. Balance Due: Rs.{receipt.remaining_balance_due:,.0f}. "
        f"Slip:{public_slip_url}"
    )

    clean_phone = clean_phone_for_wa(student.parent_phone)

    ctx.update({
        'receipt': receipt,
        'student': student,
        'clean_parent_phone': clean_phone,
        'items': receipt.items.all(),
        'amount_in_words': amount_in_words,
        'wa_message': wa_message,
        'sms_text': sms_text,
        'public_slip_url': public_slip_url,
    })
    return render(request, 'school_erp/unified_receipt.html', ctx)


# ------------------- CLASS FEE MASTER & DEFAULTERS REPORT -------------------

def class_fee_master_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    ensure_fee_structure(school=school)

    if request.method == 'POST' and 'update_fee_master' in request.POST:
        class_id = request.POST.get('class_id')
        cl = get_object_or_404(ClassLevel, pk=class_id, school=school)
        fm, _ = ClassFeeMaster.objects.get_or_create(class_level=cl)
        
        monthly_fee = Decimal(request.POST.get('monthly_tuition_fee', '2000.00') or '2000.00')
        exam_fee = Decimal(request.POST.get('exam_fee', '1200.00') or '1200.00')
        transport_fee = Decimal(request.POST.get('transport_fee', '1200.00') or '1200.00')
        admission_fee = Decimal(request.POST.get('admission_fee', '5000.00') or '5000.00')
        book_set_fee = Decimal(request.POST.get('book_set_fee', '2500.00') or '2500.00')
        dress_uniform_fee = Decimal(request.POST.get('dress_uniform_fee', '1800.00') or '1800.00')
        tie_belt_fee = Decimal(request.POST.get('tie_belt_fee', '250.00') or '250.00')
        diary_fee = Decimal(request.POST.get('diary_fee', '150.00') or '150.00')
        id_card_fee = Decimal(request.POST.get('id_card_fee', '100.00') or '100.00')
        other_activity_fee = Decimal(request.POST.get('other_activity_fee', '500.00') or '500.00')

        fm.monthly_tuition_fee = monthly_fee
        fm.exam_fee = exam_fee
        fm.transport_fee = transport_fee
        fm.admission_fee = admission_fee
        fm.book_set_fee = book_set_fee
        fm.dress_uniform_fee = dress_uniform_fee
        fm.tie_belt_fee = tie_belt_fee
        fm.diary_fee = diary_fee
        fm.id_card_fee = id_card_fee
        fm.other_activity_fee = other_activity_fee
        fm.save()

        # Also copy / apply to any additional checked classes
        also_apply_to = request.POST.getlist('also_apply_to')
        copied_names = []
        for other_id in also_apply_to:
            if str(other_id) != str(class_id):
                other_cl = ClassLevel.objects.filter(id=other_id).first()
                if other_cl:
                    other_fm, _ = ClassFeeMaster.objects.get_or_create(class_level=other_cl)
                    other_fm.monthly_tuition_fee = monthly_fee
                    other_fm.exam_fee = exam_fee
                    other_fm.transport_fee = transport_fee
                    other_fm.admission_fee = admission_fee
                    other_fm.book_set_fee = book_set_fee
                    other_fm.dress_uniform_fee = dress_uniform_fee
                    other_fm.tie_belt_fee = tie_belt_fee
                    other_fm.diary_fee = diary_fee
                    other_fm.id_card_fee = id_card_fee
                    other_fm.other_activity_fee = other_activity_fee
                    other_fm.save()
                    copied_names.append(other_cl.name)

        if copied_names:
            messages.success(request, f"Fee Master for '{cl.name}' and {len(copied_names)} other classes ({', '.join(copied_names)}) updated successfully!")
        else:
            messages.success(request, f"Class Fee Master for '{cl.name}' updated successfully! Monthly: ₹{fm.monthly_tuition_fee}, Books: ₹{fm.book_set_fee}, Dress: ₹{fm.dress_uniform_fee}")
        return redirect('class_fee_master')

    # Bulk Multi-Class Fee Structure Update
    if request.method == 'POST' and 'bulk_update_fee_master' in request.POST:
        selected_classes = request.POST.getlist('selected_classes')
        if not selected_classes:
            messages.error(request, "Please select at least one class!")
            return redirect('class_fee_master')

        monthly_fee = Decimal(request.POST.get('monthly_tuition_fee', '2000.00') or '2000.00')
        exam_fee = Decimal(request.POST.get('exam_fee', '1200.00') or '1200.00')
        transport_fee = Decimal(request.POST.get('transport_fee', '1200.00') or '1200.00')
        admission_fee = Decimal(request.POST.get('admission_fee', '5000.00') or '5000.00')
        book_set_fee = Decimal(request.POST.get('book_set_fee', '2500.00') or '2500.00')
        dress_uniform_fee = Decimal(request.POST.get('dress_uniform_fee', '1800.00') or '1800.00')
        tie_belt_fee = Decimal(request.POST.get('tie_belt_fee', '250.00') or '250.00')
        diary_fee = Decimal(request.POST.get('diary_fee', '150.00') or '150.00')
        id_card_fee = Decimal(request.POST.get('id_card_fee', '100.00') or '100.00')
        other_activity_fee = Decimal(request.POST.get('other_activity_fee', '500.00') or '500.00')

        updated_names = []
        for cl_id in selected_classes:
            cl = ClassLevel.objects.filter(id=cl_id).first()
            if cl:
                fm, _ = ClassFeeMaster.objects.get_or_create(class_level=cl)
                fm.monthly_tuition_fee = monthly_fee
                fm.exam_fee = exam_fee
                fm.transport_fee = transport_fee
                fm.admission_fee = admission_fee
                fm.book_set_fee = book_set_fee
                fm.dress_uniform_fee = dress_uniform_fee
                fm.tie_belt_fee = tie_belt_fee
                fm.diary_fee = diary_fee
                fm.id_card_fee = id_card_fee
                fm.other_activity_fee = other_activity_fee
                fm.save()
                updated_names.append(cl.name)

        messages.success(request, f"Fee structure successfully applied across {len(updated_names)} classes ({', '.join(updated_names)})!")
        return redirect('class_fee_master')

    classes = ClassLevel.objects.filter(school=school).order_by('order')
    fee_masters = []
    for cl in classes:
        try:
            fm = cl.fee_master
        except Exception:
            fm, _ = ClassFeeMaster.objects.get_or_create(class_level=cl, defaults={'monthly_tuition_fee': Decimal('2000.00')})
        student_count = Student.objects.filter(school=school, class_level=cl, is_active=True).count()
        fee_masters.append({
            'class': cl,
            'fm': fm,
            'student_count': student_count,
        })

    ctx.update({
        'classes': classes,
        'fee_masters': fee_masters,
        'title': 'Class-Wise Fee Master Setup',
    })
    return render(request, 'school_erp/class_fee_master.html', ctx)


@require_menu_permission('can_manage_fees', 'Fee Defaulters')
def fees_defaulters_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    ensure_fee_structure(school=school)
    active_session = get_active_session(request, school=school)

    class_id = request.GET.get('class_id', '')
    search = request.GET.get('q', '').strip()
    min_due = request.GET.get('min_due', '')

    students_qs = Student.objects.filter(school=school, is_active=True).select_related('class_level', 'section').prefetch_related('invoices__payments')
    if active_session:
        students_qs = students_qs.filter(
            Q(academic_session=active_session) | Q(session_enrollments__academic_session=active_session)
        ).distinct()

    if class_id:
        students_qs = students_qs.filter(class_level_id=class_id)
    if search:
        students_qs = students_qs.filter(
            Q(first_name__icontains=search) |
            Q(last_name__icontains=search) |
            Q(admission_number__icontains=search) |
            Q(parent_name__icontains=search) |
            Q(parent_phone__icontains=search)
        )

    defaulters = []
    total_school_due = Decimal('0.00')
    total_defaulter_count = 0

    for s in students_qs:
        unpaid_invoices = s.invoices.filter(
            academic_session=active_session,
            status__in=['Unpaid', 'Partially Paid', 'Overdue']
        ) if active_session else s.invoices.filter(status__in=['Unpaid', 'Partially Paid', 'Overdue'])
        due = sum((inv.remaining_due for inv in unpaid_invoices), Decimal('0.00'))
        if due > Decimal('0.00'):
            if min_due:
                try:
                    if due < Decimal(min_due):
                        continue
                except ValueError:
                    pass

            unpaid_heads = [inv.fee_category.name for inv in unpaid_invoices[:3]]
            
            # Format clean WhatsApp pre-filled reminder text
            upi_note = f" You can also pay online via College UPI: {ctx['school'].upi_id}." if getattr(ctx['school'], 'upi_id', None) else ""
            wa_text = (
                f"Namaste {s.parent_name or 'Parent'}, "
                f"Fee Reminder from {ctx['school'].name}: "
                f"School fee of Rs.{due:,.0f} is pending for your ward {s.full_name} "
                f"(Adm No: {s.admission_number}, Class: {s.class_level.name}-{s.section.name}).{upi_note} "
                f"Kindly deposit at the fee counter or via UPI at the earliest. Thank you."
            )

            defaulters.append({
                'student': s,
                'due_amount': due,
                'unpaid_count': unpaid_invoices.count(),
                'unpaid_heads': unpaid_heads,
                'whatsapp_text': wa_text,
                'clean_parent_phone': clean_phone_for_wa(s.parent_phone),
            })
            total_school_due += due
            total_defaulter_count += 1

    # Sort by due amount descending (highest defaulters first)
    defaulters.sort(key=lambda x: x['due_amount'], reverse=True)

    classes = ClassLevel.objects.filter(school=school).order_by('order')

    ctx.update({
        'defaulters': defaulters,
        'total_school_due': total_school_due,
        'total_defaulter_count': total_defaulter_count,
        'classes': classes,
        'selected_class': class_id,
        'search': search,
        'min_due': min_due,
    })
    return render(request, 'school_erp/fees_defaulters.html', ctx)


# ==============================================================================
# ENTERPRISE FEE MANAGEMENT SCREENS
# ==============================================================================

@require_menu_permission('can_manage_fees', 'Fee Head Master')
def fee_heads_master_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    ensure_fee_structure(school)

    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'add_head':
            form = FeeHeadForm(request.POST)
            if form.is_valid():
                head = form.save(commit=False)
                head.school = school
                head.save()
                messages.success(request, f"Fee Head '{head.name}' ({head.get_fee_type_display()}) created successfully!")
            else:
                messages.error(request, "Failed to create Fee Head. Please check form inputs.")
            return redirect('fee_heads_master')

        elif action == 'edit_head':
            head_id = request.POST.get('head_id')
            head = get_object_or_404(FeeHead, pk=head_id)
            form = FeeHeadForm(request.POST, instance=head)
            if form.is_valid():
                form.save()
                messages.success(request, f"Fee Head '{head.name}' updated successfully!")
            else:
                messages.error(request, "Failed to update Fee Head.")
            return redirect('fee_heads_master')

        elif action == 'toggle_status':
            head_id = request.POST.get('head_id')
            head = get_object_or_404(FeeHead, pk=head_id)
            head.is_active = not head.is_active
            head.save()
            messages.success(request, f"Fee Head '{head.name}' status changed to {'Active' if head.is_active else 'Inactive'}.")
            return redirect('fee_heads_master')

        elif action == 'delete_head':
            head_id = request.POST.get('head_id')
            head = get_object_or_404(FeeHead, pk=head_id)
            head_name = head.name
            head.delete()
            messages.success(request, f"Fee Head '{head_name}' removed.")
            return redirect('fee_heads_master')

    heads = FeeHead.objects.all().order_by('fee_type', 'name')
    if school:
        heads = heads.filter(Q(school=school) | Q(school__isnull=True))

    type_filter = request.GET.get('type')
    if type_filter:
        heads = heads.filter(fee_type=type_filter)

    search = request.GET.get('search', '').strip()
    if search:
        heads = heads.filter(Q(name__icontains=search) | Q(code__icontains=search))

    stats = {
        'total': FeeHead.objects.count(),
        'academic': FeeHead.objects.filter(fee_type='Academic').count(),
        'transport_hostel': FeeHead.objects.filter(fee_type__in=['Transport', 'Hostel']).count(),
        'optional': FeeHead.objects.filter(fee_type__in=['Optional', 'Other']).count(),
    }

    ctx.update({
        'heads': heads,
        'stats': stats,
        'fee_types': FeeHead.FEE_TYPES,
        'selected_type': type_filter,
        'search': search,
        'form': FeeHeadForm(),
    })
    return render(request, 'school_erp/fee_heads_master.html', ctx)


@require_menu_permission('can_manage_fees', 'Fee Structure & Schedules')
def fee_structure_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    ensure_fee_structure(school)
    active_session = get_active_session(request)

    classes = ClassLevel.objects.all().order_by('order')
    fee_heads = FeeHead.objects.filter(is_active=True).order_by('name')
    routes = TransportRoute.objects.all().order_by('route_name')
    hostels = HostelFacility.objects.all().order_by('hostel_name')
    concessions = StudentConcession.objects.select_related('student', 'academic_session').all().order_by('-created_at')

    # Handle Forms
    if request.method == 'POST':
        action = request.POST.get('action')
        
        # Save Instalment Schedule
        if action == 'save_instalment_schedule':
            class_id = request.POST.get('class_id')
            cl = get_object_or_404(ClassLevel, pk=class_id)
            title = request.POST.get('title')
            due_date = request.POST.get('due_date')
            amount = Decimal(request.POST.get('amount', '0.00') or '0.00')
            grace_days = int(request.POST.get('grace_period_days', '10') or '10')
            late_fine = Decimal(request.POST.get('late_fine_amount', '50.00') or '50.00')
            instalment_num = int(request.POST.get('instalment_number', '1') or '1')

            FeeInstalmentSchedule.objects.create(
                school=school,
                academic_session=active_session or AcademicSession.objects.filter(is_current=True).first(),
                class_level=cl,
                instalment_number=instalment_num,
                title=title,
                due_date=due_date,
                amount=amount,
                grace_period_days=grace_days,
                late_fine_amount=late_fine
            )
            messages.success(request, f"Instalment Schedule '{title}' created for {cl.name} (Due: {due_date})!")
            return redirect('fee_structure')

        # Add Transport Route
        elif action == 'add_route':
            r_form = TransportRouteForm(request.POST)
            if r_form.is_valid():
                r = r_form.save(commit=False)
                r.school = school
                r.save()
                messages.success(request, f"Transport Route '{r.route_name}' added successfully!")
            return redirect('fee_structure')

        # Add Hostel Facility
        elif action == 'add_hostel':
            h_form = HostelFacilityForm(request.POST)
            if h_form.is_valid():
                h = h_form.save(commit=False)
                h.school = school
                h.save()
                messages.success(request, f"Hostel Facility '{h.hostel_name}' added successfully!")
            return redirect('fee_structure')

        # Add Student Concession / Scholarship
        elif action == 'add_concession':
            student_id = request.POST.get('student')
            student = get_object_or_404(Student, pk=student_id)
            c_type = request.POST.get('concession_type', 'Sibling Discount')
            d_type = request.POST.get('discount_type', 'Percentage')
            d_val = Decimal(request.POST.get('discount_value', '10.00') or '10.00')
            reason = request.POST.get('reason', '')
            
            StudentConcession.objects.create(
                student=student,
                academic_session=active_session or AcademicSession.objects.filter(is_current=True).first(),
                concession_type=c_type,
                discount_type=d_type,
                discount_value=d_val,
                reason=reason,
                approved_by=request.user if request.user.is_authenticated else None
            )
            messages.success(request, f"{c_type} approved for {student.full_name} ({d_val}{'%' if d_type == 'Percentage' else '₹'})!")
            return redirect('fee_structure')

    # Get active schedules
    schedules = FeeInstalmentSchedule.objects.select_related('class_level', 'academic_session').all().order_by('class_level__order', 'due_date')
    if active_session:
        schedules = schedules.filter(academic_session=active_session)

    ctx.update({
        'classes': classes,
        'fee_heads': fee_heads,
        'schedules': schedules,
        'routes': routes,
        'hostels': hostels,
        'concessions': concessions,
        'route_form': TransportRouteForm(),
        'hostel_form': HostelFacilityForm(),
        'concession_form': StudentConcessionForm(),
        'active_tab': request.GET.get('tab', 'structures'),
    })
    return render(request, 'school_erp/fee_structure.html', ctx)


@require_menu_permission('can_manage_fees', 'Student Fee Ledger')
def student_ledger_view(request, student_id=None):
    ctx = get_school_context(request)
    school = ctx['school']
    active_session = get_active_session(request)

    if not student_id:
        student_id = request.GET.get('student_id')

    students = Student.objects.filter(is_active=True).select_related('class_level', 'section').order_by('class_level__order', 'roll_number')
    
    # Parent Scoped Access Restriction Enforcement
    if is_parent_user(request.user):
        allowed_students = get_parent_allowed_students(request.user)
        students = allowed_students.order_by('class_level__order', 'roll_number')

    selected_student = None
    ledger_entries = []
    
    total_debit = Decimal('0.00')
    total_credit = Decimal('0.00')
    net_balance_due = Decimal('0.00')

    if student_id:
        selected_student = get_object_or_404(Student.objects.select_related('class_level', 'section'), pk=student_id)
        
        # Verify parent ownership of student
        if is_parent_user(request.user):
            allowed_students = get_parent_allowed_students(request.user)
            if not allowed_students.filter(id=selected_student.id).exists():
                messages.error(
                    request,
                    "Access Restricted: As a Parent, you can only view fee ledger statements for your linked child."
                )
                return redirect('dashboard')
        
        # Pull from StudentFeeLedger
        ledger_qs = StudentFeeLedger.objects.filter(student=selected_student).order_by('entry_date', 'id')
        
        # If ledger entries are empty (e.g. for existing demo students before ledger table), synthesize ledger from invoices and receipts
        if not ledger_qs.exists():
            running = Decimal('0.00')
            # Add Invoices
            invoices = selected_student.invoices.all().order_by('created_at')
            for inv in invoices:
                running += inv.amount
                ledger_entries.append({
                    'date': inv.created_at.date(),
                    'type': 'Fee Charge',
                    'voucher_no': inv.invoice_number,
                    'head': inv.fee_category.name,
                    'description': f"Invoice Generated: {inv.fee_category.name}",
                    'debit': inv.amount,
                    'credit': Decimal('0.00'),
                    'running_balance': running,
                    'badge_class': 'badge-primary'
                })
                if inv.discount > Decimal('0.00'):
                    running -= inv.discount
                    ledger_entries.append({
                        'date': inv.created_at.date(),
                        'type': 'Concession / Discount',
                        'voucher_no': inv.invoice_number,
                        'head': inv.fee_category.name,
                        'description': "Concession / Scholarship Adjusted",
                        'debit': Decimal('0.00'),
                        'credit': inv.discount,
                        'running_balance': running,
                        'badge_class': 'badge-success'
                    })

            # Add Receipts
            receipts = selected_student.unified_receipts.all().order_by('payment_date')
            for r in receipts:
                if r.status == 'Cancelled':
                    continue
                running = max(Decimal('0.00'), running - r.amount_paid)
                ledger_entries.append({
                    'date': r.payment_date,
                    'type': 'Fee Payment',
                    'voucher_no': r.receipt_number,
                    'head': r.payment_method,
                    'description': f"Payment Received ({r.payment_method})" + (f" Ref: {r.transaction_reference}" if r.transaction_reference else ""),
                    'debit': Decimal('0.00'),
                    'credit': r.amount_paid,
                    'running_balance': running,
                    'badge_class': 'badge-info'
                })
        else:
            for entry in ledger_qs:
                ledger_entries.append({
                    'date': entry.entry_date,
                    'type': entry.get_entry_type_display(),
                    'voucher_no': entry.voucher_no,
                    'head': entry.fee_head.name if entry.fee_head else 'General',
                    'description': entry.description,
                    'debit': entry.debit,
                    'credit': entry.credit,
                    'running_balance': entry.running_balance,
                    'badge_class': 'badge-primary' if entry.debit > 0 else 'badge-success'
                })

        total_debit = sum((e['debit'] for e in ledger_entries), Decimal('0.00'))
        total_credit = sum((e['credit'] for e in ledger_entries), Decimal('0.00'))
        net_balance_due = max(Decimal('0.00'), total_debit - total_credit)

    ctx.update({
        'students': students,
        'selected_student': selected_student,
        'ledger_entries': ledger_entries,
        'total_debit': total_debit,
        'total_credit': total_credit,
        'net_balance_due': net_balance_due,
    })
    return render(request, 'school_erp/student_ledger.html', ctx)


@require_menu_permission('can_manage_fees', 'Export Student Ledger')
def export_student_ledger_excel(request, student_id):
    student = get_object_or_404(Student, pk=student_id)
    ctx = get_school_context(request)
    school = ctx['school']

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Ledger - {student.admission_number}"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    th_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")

    ws.merge_cells("A1:G1")
    ws["A1"] = f"{school.name} - OFFICIAL STUDENT FEE LEDGER"
    ws["A1"].font = title_font
    ws["A1"].fill = header_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Student: {student.full_name} | Adm No: {student.admission_number} | Class: {student.class_level.name} - Sec {student.section.name} | Father: {student.parent_name}"
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 20

    headers = ["Date", "Transaction Type", "Voucher / Receipt No", "Description", "Debit (Due)", "Credit (Paid)", "Running Balance"]
    ws.append([])
    ws.append(headers)
    ws.row_dimensions[4].height = 26

    for col_idx in range(1, 8):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = th_fill
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ledger_qs = StudentFeeLedger.objects.filter(student=student).order_by('entry_date', 'id')
    for entry in ledger_qs:
        ws.append([
            entry.entry_date.strftime("%d-%m-%Y"),
            entry.get_entry_type_display(),
            entry.voucher_no,
            entry.description,
            float(entry.debit),
            float(entry.credit),
            float(entry.running_balance)
        ])

    widths = [14, 22, 22, 35, 14, 14, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Ledger_{student.admission_number}.xlsx"'
    wb.save(response)
    return response


@require_menu_permission('can_manage_fees', 'Fee Receipts')
def fee_receipts_list_view(request):
    ctx = get_school_context(request)
    ensure_fee_structure()

    receipts = FeeReceipt.objects.select_related('student__class_level', 'student__section', 'cashier').all()
    
    # Filters
    search = request.GET.get('search', '').strip()
    if search:
        receipts = receipts.filter(
            Q(receipt_number__icontains=search) |
            Q(student__first_name__icontains=search) |
            Q(student__last_name__icontains=search) |
            Q(student__admission_number__icontains=search) |
            Q(transaction_reference__icontains=search)
        )

    status_filter = request.GET.get('status')
    if status_filter:
        receipts = receipts.filter(status=status_filter)

    method_filter = request.GET.get('method')
    if method_filter:
        receipts = receipts.filter(payment_method=method_filter)

    date_from = request.GET.get('date_from')
    if date_from:
        receipts = receipts.filter(payment_date__gte=date_from)

    date_to = request.GET.get('date_to')
    if date_to:
        receipts = receipts.filter(payment_date__lte=date_to)

    total_amount = receipts.filter(status='Active').aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
    total_count = receipts.count()
    cancelled_count = receipts.filter(status='Cancelled').count()
    reconciled_count = receipts.filter(is_reconciled=True).count()

    paginator = Paginator(receipts, 25)
    page = request.GET.get('page')
    try:
        paged_receipts = paginator.page(page)
    except (PageNotAnInteger, EmptyPage):
        paged_receipts = paginator.page(1)

    ctx.update({
        'receipts': paged_receipts,
        'total_amount': total_amount,
        'total_count': total_count,
        'cancelled_count': cancelled_count,
        'reconciled_count': reconciled_count,
        'search': search,
        'status_filter': status_filter,
        'method_filter': method_filter,
        'date_from': date_from,
        'date_to': date_to,
        'payment_methods': FeeReceipt.PAYMENT_MODES,
    })
    return render(request, 'school_erp/fee_receipts_list.html', ctx)


@require_menu_permission('can_manage_fees', 'Cancel Fee Receipt')
def fee_receipt_cancel_view(request, receipt_id):
    if request.method != 'POST':
        messages.error(request, "Invalid cancellation request.")
        return redirect('fee_receipts_list')

    receipt = get_object_or_404(FeeReceipt, pk=receipt_id)
    if receipt.status == 'Cancelled':
        messages.warning(request, f"Receipt #{receipt.receipt_number} is already cancelled.")
        return redirect('fee_receipts_list')

    reason = request.POST.get('cancellation_reason', '').strip()
    if not reason:
        messages.error(request, "Cancellation reason is mandatory for audit trail!")
        return redirect('fee_receipts_list')

    receipt.status = 'Cancelled'
    receipt.cancellation_reason = reason
    receipt.cancelled_by = request.user if request.user.is_authenticated else None
    receipt.cancelled_at = timezone.now()
    receipt.save()

    # Reversal in StudentFeeLedger
    StudentFeeLedger.objects.create(
        school=receipt.school,
        student=receipt.student,
        academic_session=receipt.student.academic_session or AcademicSession.objects.filter(is_current=True).first(),
        entry_date=timezone.now().date(),
        entry_type='Cancellation',
        voucher_no=receipt.receipt_number,
        description=f"Receipt #{receipt.receipt_number} Cancelled: {reason}",
        debit=receipt.amount_paid,
        credit=Decimal('0.00'),
        running_balance=receipt.student.fee_due_amount,
        receipt=receipt,
        created_by=request.user if request.user.is_authenticated else None
    )

    # Immutable Audit Log
    log_audit(
        request=request,
        action='DELETE',
        module='FEES',
        object_type='FeeReceipt',
        object_id=str(receipt.id),
        object_repr=f"Receipt #{receipt.receipt_number}",
        description=f"Fee receipt #{receipt.receipt_number} for {receipt.student.full_name} cancelled. Amount: ₹{receipt.amount_paid}",
        reason=reason,
        requires_approval=False,
        status='SUCCESS'
    )

    messages.success(request, f"Official Fee Receipt #{receipt.receipt_number} has been CANCELLED and logged in the audit trail.")
    return redirect('fee_receipts_list')


@require_menu_permission('can_manage_fees', 'Fee Refunds')
def fee_refunds_view(request):
    ctx = get_school_context(request)
    school = ctx['school']

    if request.method == 'POST' and 'create_refund' in request.POST:
        student_id = request.POST.get('student')
        student = get_object_or_404(Student, pk=student_id)
        amount = Decimal(request.POST.get('amount', '0.00') or '0.00')
        receipt_id = request.POST.get('receipt')
        receipt = FeeReceipt.objects.filter(pk=receipt_id).first() if receipt_id else None
        refund_mode = request.POST.get('refund_mode', 'Bank Transfer')
        reference_no = request.POST.get('reference_number', '').strip()
        reason = request.POST.get('reason', '').strip()

        if amount <= Decimal('0.00'):
            messages.error(request, "Refund amount must be greater than ₹0.")
            return redirect('fee_refunds')

        refund_count = FeeRefund.objects.count() + 1
        ref_no = f"REF-{timezone.now().year}-{refund_count:04d}"

        refund = FeeRefund.objects.create(
            school=school,
            refund_number=ref_no,
            student=student,
            receipt=receipt,
            amount=amount,
            refund_date=timezone.now().date(),
            refund_mode=refund_mode,
            reference_number=reference_no,
            reason=reason,
            status='Pending',
            requested_by=request.user if request.user.is_authenticated else None
        )
        messages.success(request, f"Refund Request #{ref_no} for ₹{amount} submitted for approval!")
        return redirect('fee_refunds')

    refunds = FeeRefund.objects.select_related('student__class_level', 'receipt', 'requested_by', 'approved_by').all().order_by('-created_at')
    
    status_filter = request.GET.get('status')
    if status_filter:
        refunds = refunds.filter(status=status_filter)

    students = Student.objects.filter(is_active=True).select_related('class_level').order_by('first_name')
    receipts = FeeReceipt.objects.filter(status='Active').order_by('-payment_date')[:50]

    stats = {
        'total': FeeRefund.objects.count(),
        'pending': FeeRefund.objects.filter(status='Pending').count(),
        'approved': FeeRefund.objects.filter(status='Approved').count(),
        'processed': FeeRefund.objects.filter(status='Processed').count(),
        'total_disbursed': FeeRefund.objects.filter(status='Processed').aggregate(t=Sum('amount'))['t'] or Decimal('0.00'),
    }

    ctx.update({
        'refunds': refunds,
        'students': students,
        'receipts': receipts,
        'stats': stats,
        'status_filter': status_filter,
    })
    return render(request, 'school_erp/fee_refunds.html', ctx)


@require_menu_permission('can_manage_fees', 'Refund Approval & Disbursement')
def fee_refund_action_view(request, refund_id, action):
    refund = get_object_or_404(FeeRefund, pk=refund_id)
    
    if action == 'approve':
        refund.status = 'Approved'
        refund.approved_by = request.user if request.user.is_authenticated else None
        refund.approved_at = timezone.now()
        refund.save()
        messages.success(request, f"Refund #{refund.refund_number} for {refund.student.full_name} (₹{refund.amount}) has been APPROVED!")

    elif action == 'reject':
        refund.status = 'Rejected'
        refund.remarks = request.POST.get('remarks', 'Rejected by Administrator')
        refund.save()
        messages.warning(request, f"Refund #{refund.refund_number} has been REJECTED.")

    elif action == 'process':
        refund.status = 'Processed'
        refund.reference_number = request.POST.get('reference_number', refund.reference_number)
        refund.save()

        # Record in student fee ledger
        StudentFeeLedger.objects.create(
            school=refund.school,
            student=refund.student,
            academic_session=refund.student.academic_session or AcademicSession.objects.filter(is_current=True).first(),
            entry_date=timezone.now().date(),
            entry_type='Refund',
            voucher_no=refund.refund_number,
            description=f"Refund Disbursed via {refund.refund_mode} (Ref: {refund.reference_number}): {refund.reason}",
            debit=refund.amount,
            credit=Decimal('0.00'),
            running_balance=refund.student.fee_due_amount + refund.amount,
            created_by=request.user if request.user.is_authenticated else None
        )
        messages.success(request, f"Refund #{refund.refund_number} has been DISBURSED & recorded in the ledger!")

    return redirect('fee_refunds')


@require_menu_permission('can_manage_fees', 'Payment Reconciliation')
def fee_reconciliation_view(request):
    ctx = get_school_context(request)
    today = timezone.now().date()
    
    # Handle single or bulk reconciliation toggle
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'reconcile_single':
            receipt_id = request.POST.get('receipt_id')
            rec = get_object_or_404(FeeReceipt, pk=receipt_id)
            rec.is_reconciled = not rec.is_reconciled
            rec.reconciled_at = timezone.now() if rec.is_reconciled else None
            rec.reconciled_by = request.user if rec.is_reconciled else None
            rec.reconciliation_notes = request.POST.get('notes', '')
            rec.save()
            messages.success(request, f"Receipt #{rec.receipt_number} marked as {'Reconciled' if rec.is_reconciled else 'Pending Reconciliation'}.")
            return redirect(request.get_full_path())

        elif action == 'reconcile_all_selected':
            selected_ids = request.POST.getlist('selected_receipts')
            count = 0
            for r_id in selected_ids:
                rec = FeeReceipt.objects.filter(id=r_id).first()
                if rec and not rec.is_reconciled:
                    rec.is_reconciled = True
                    rec.reconciled_at = timezone.now()
                    rec.reconciled_by = request.user
                    rec.save()
                    count += 1
            messages.success(request, f"Successfully verified & reconciled {count} payment transactions!")
            return redirect(request.get_full_path())

    date_filter = request.GET.get('date_range', 'today')
    receipts_qs = FeeReceipt.objects.filter(status='Active').select_related('student__class_level', 'cashier').order_by('-payment_date', '-id')

    if date_filter == 'today':
        receipts_qs = receipts_qs.filter(payment_date=today)
    elif date_filter == 'yesterday':
        receipts_qs = receipts_qs.filter(payment_date=today - datetime.timedelta(days=1))
    elif date_filter == 'this_week':
        week_start = today - datetime.timedelta(days=today.weekday())
        receipts_qs = receipts_qs.filter(payment_date__gte=week_start)
    elif date_filter == 'this_month':
        month_start = today.replace(day=1)
        receipts_qs = receipts_qs.filter(payment_date__gte=month_start)
    elif date_filter == 'custom':
        df = request.GET.get('from_date')
        dt = request.GET.get('to_date')
        if df:
            receipts_qs = receipts_qs.filter(payment_date__gte=df)
        if dt:
            receipts_qs = receipts_qs.filter(payment_date__lte=dt)

    # Breakdown by Payment Modes
    cash_total = receipts_qs.filter(payment_method='Cash').aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
    upi_total = receipts_qs.filter(payment_method__icontains='UPI').aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
    bank_total = receipts_qs.filter(payment_method__in=['Net Banking', 'Cheque / DD', 'Debit/Credit Card']).aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
    total_amount = cash_total + upi_total + bank_total
    reconciled_total = receipts_qs.filter(is_reconciled=True).aggregate(t=Sum('amount_paid'))['t'] or Decimal('0.00')
    unreconciled_total = max(Decimal('0.00'), total_amount - reconciled_total)

    # Check for suspicious duplicates or missing UTRs
    potential_mismatches = []
    seen_refs = set()
    for r in receipts_qs:
        if r.payment_method != 'Cash':
            if not r.transaction_reference:
                potential_mismatches.append(f"Receipt #{r.receipt_number} ({r.payment_method}) is missing Transaction Reference / UTR.")
            elif r.transaction_reference in seen_refs:
                potential_mismatches.append(f"Duplicate Transaction Reference '{r.transaction_reference}' detected on Receipt #{r.receipt_number}.")
            else:
                seen_refs.add(r.transaction_reference)

    ctx.update({
        'receipts': receipts_qs,
        'cash_total': cash_total,
        'upi_total': upi_total,
        'bank_total': bank_total,
        'total_amount': total_amount,
        'reconciled_total': reconciled_total,
        'unreconciled_total': unreconciled_total,
        'date_filter': date_filter,
        'potential_mismatches': potential_mismatches,
    })
    return render(request, 'school_erp/fee_reconciliation.html', ctx)


@require_menu_permission('can_manage_fees', 'Send Fee Reminder')
def fee_send_reminder_view(request, student_id):
    student = get_object_or_404(Student, pk=student_id)
    ctx = get_school_context(request)
    school = ctx['school']
    due = student.fee_due_amount

    upi_str = f" / College UPI ID: {school.upi_id}" if school.upi_id else ""
    msg = (
        f"Namaste {student.parent_name or 'Parent'}, "
        f"Fee Reminder from {school.name}: "
        f"Pending school fee for {student.full_name} (Class {student.class_level.name}-{student.section.name}) is Rs.{due:,.0f}. "
        f"Kindly deposit at the counter or scan College UPI QR code{upi_str}. "
        f"Thank you. - {school.phone}"
    )

    clean_phone = ''.join(c for c in (student.parent_phone or '') if c.isdigit())
    if clean_phone and len(clean_phone) == 10:
        clean_phone = '91' + clean_phone

    FeeReminderLog.objects.create(
        school=school,
        student=student,
        channel='WhatsApp',
        recipient_phone=student.parent_phone or 'N/A',
        due_amount=due,
        message_text=msg,
        sent_by=request.user if request.user.is_authenticated else None
    )

    import urllib.parse
    encoded_msg = urllib.parse.quote(msg)
    wa_url = f"https://wa.me/{clean_phone}?text={encoded_msg}" if clean_phone else "#"

    messages.success(request, f"Due Reminder logged for {student.full_name}! Opening WhatsApp...")
    return redirect(wa_url if clean_phone else 'fees_defaulters')


# ------------------- EXAMINATIONS & REPORT CARDS -------------------

@require_menu_permission('can_manage_exams', 'Examinations')
def exam_list_view(request):
    ctx = get_school_context(request)
    school = ctx.get('school')
    exams = Exam.objects.all().order_by('-start_date')
    classes = ClassLevel.objects.all()
    sections = Section.objects.all()
    form = ExamForm(school=school)
    ctx.update({
        'exams': exams,
        'classes': classes,
        'sections': sections,
        'form': form,
        'title': 'Examinations & Student Assessments'
    })
    return render(request, 'school_erp/exam_list.html', ctx)


@require_menu_permission('can_manage_exams', 'Schedule Examination')
def exam_create_view(request):
    ctx = get_school_context(request)
    school = ctx.get('school')
    if request.method == 'POST':
        form = ExamForm(request.POST, school=school)
        if form.is_valid():
            exam = form.save(commit=False)
            exam.school = school
            exam.save()
            messages.success(request, f"Examination '{exam.name}' scheduled successfully!")
            return redirect('exam_list')
    else:
        form = ExamForm(school=school)
    ctx.update({
        'form': form,
        'title': 'Schedule New Examination'
    })
    return render(request, 'school_erp/exam_form.html', ctx)



@require_menu_permission('can_manage_exams', 'Marks Entry')
def marks_entry_view(request):
    ctx = get_school_context(request)
    exams = Exam.objects.all()
    classes = ClassLevel.objects.all()
    subjects = Subject.objects.all()

    # Teacher Scoped Access Restriction Enforcement
    is_teacher = is_teacher_user(request.user)
    teacher_assigned_classes = None
    teacher_assigned_subjects = None
    if is_teacher:
        teacher_assigned_classes, teacher_assigned_subjects = get_teacher_assigned_classes(request.user)
        classes = teacher_assigned_classes
        if teacher_assigned_subjects.exists():
            subjects = teacher_assigned_subjects

    exam_id = request.GET.get('exam_id') or request.POST.get('exam_id')
    class_id = request.GET.get('class_id') or request.POST.get('class_id')
    subject_id = request.GET.get('subject_id') or request.POST.get('subject_id')

    selected_exam = Exam.objects.filter(id=exam_id).first() if exam_id else exams.first()
    selected_class = classes.filter(id=class_id).first() if class_id else classes.first()
    selected_subject = subjects.filter(id=subject_id).first() if subject_id else subjects.first()

    # If teacher explicitly tries to access a class or subject not assigned to them
    if is_teacher:
        if class_id and not teacher_assigned_classes.filter(id=class_id).exists():
            target_cls = ClassLevel.objects.filter(id=class_id).first()
            cls_name = target_cls.name if target_cls else f"ID {class_id}"
            messages.error(
                request,
                f"Access Denied: You are not authorized to view or enter marks for Class '{cls_name}'. You can only grade your assigned classes."
            )
            return redirect('marks_entry')

    students = []
    existing_marks = {}

    if selected_exam and selected_class and selected_subject:
        students = Student.objects.filter(class_level=selected_class, is_active=True).order_by('roll_number')

        if request.method == 'POST' and 'save_marks' in request.POST:
            # Re-verify teacher permission upon save POST
            if is_teacher and not teacher_assigned_classes.filter(id=selected_class.id).exists():
                messages.error(
                    request,
                    f"Access Denied: You cannot modify marks for Class '{selected_class.name}'. This class is not assigned to you."
                )
                return redirect('marks_entry')

            max_marks_val = float(request.POST.get('max_marks', 100.0))
            updated_count = 0
            for s in students:
                score_str = request.POST.get(f'marks_{s.id}', '').strip()
                remarks_str = request.POST.get(f'remarks_{s.id}', 'Good')
                if score_str != '':
                    try:
                        score = float(score_str)
                        ExamResult.objects.update_or_create(
                            exam=selected_exam,
                            student=s,
                            subject=selected_subject,
                            defaults={
                                'marks_obtained': score,
                                'max_marks': max_marks_val,
                                'remarks': remarks_str,
                            }
                        )
                        updated_count += 1
                    except ValueError:
                        pass

            # Log audit trail for marks update
            log_audit(
                request=request,
                action='UPDATE',
                module='EXAMS',
                object_type='ExamResult',
                object_id=str(selected_exam.id),
                object_repr=f"{selected_exam.name} - {selected_class.name} - {selected_subject.name}",
                description=f"Marks updated for {updated_count} students in {selected_class.name} ({selected_subject.name}). Max Marks: {max_marks_val}",
                reason="Term examination marks assessment update"
            )

            messages.success(request, f"Marks updated successfully for {selected_subject.name} - {selected_class.name} ({updated_count} students)!")
            return redirect(f"{request.path}?exam_id={selected_exam.id}&class_id={selected_class.id}&subject_id={selected_subject.id}")

        results = ExamResult.objects.filter(exam=selected_exam, subject=selected_subject, student__in=students)
        existing_marks = {r.student_id: r for r in results}

    student_rows = []
    for s in students:
        res = existing_marks.get(s.id)
        student_rows.append({
            'student': s,
            'marks': res.marks_obtained if res else '',
            'grade': res.grade if res else '-',
            'remarks': res.remarks if res else '',
        })

    ctx.update({
        'exams': exams,
        'classes': classes,
        'subjects': subjects,
        'selected_exam': selected_exam,
        'selected_class': selected_class,
        'selected_subject': selected_subject,
        'student_rows': student_rows,
        'is_teacher_restricted': is_teacher,
    })
    return render(request, 'school_erp/marks_entry.html', ctx)


def report_card_view(request, student_id, exam_id):
    ctx = get_school_context(request)
    student = get_object_or_404(Student.objects.select_related('class_level', 'section'), pk=student_id)
    exam = get_object_or_404(Exam, pk=exam_id)

    # Parent Scoped Access Restriction Enforcement
    if is_parent_user(request.user):
        allowed_students = get_parent_allowed_students(request.user)
        if not allowed_students.filter(id=student.id).exists():
            messages.error(
                request,
                "Access Restricted: As a Parent, you can only view academic and report card records for your linked child."
            )
            return redirect('dashboard')

    
    results = ExamResult.objects.filter(student=student, exam=exam).select_related('subject')
    
    total_obtained = sum(float(r.marks_obtained) for r in results)
    total_max = sum(float(r.max_marks) for r in results)
    overall_percentage = round((total_obtained / total_max * 100), 2) if total_max > 0 else 0
    
    # Overall Grade
    if overall_percentage >= 90:
        overall_grade = 'A+'
        verdict = 'Outstanding Performance'
    elif overall_percentage >= 80:
        overall_grade = 'A'
        verdict = 'Excellent Performance'
    elif overall_percentage >= 70:
        overall_grade = 'B'
        verdict = 'Very Good'
    elif overall_percentage >= 60:
        overall_grade = 'C'
        verdict = 'Good / Satisfactory'
    elif overall_percentage >= 50:
        overall_grade = 'D'
        verdict = 'Needs Improvement'
    else:
        overall_grade = 'F'
        verdict = 'Failed / Re-appear'
        
    attendance_stats = student.attendance_stats
    school = ctx['school']

    public_slip_url = request.build_absolute_uri(f"/slip/report-card/{student.id}/{exam.id}/")
    clean_phone = clean_phone_for_wa(student.parent_phone)

    wa_message = (
        f"*{school.name.upper()}*\n"
        f"*Official Academic Report Card / Marksheet*\n"
        f"-----------------------------------------\n"
        f"Student: *{student.full_name}*\n"
        f"Adm No: *{student.admission_number}* | Class: *{student.class_level.name}-{student.section.name}* (#{student.roll_number})\n"
        f"Exam: *{exam.name}*\n"
        f"Marks: *{total_obtained:,.1f} / {total_max:,.0f}* ({overall_percentage}%)\n"
        f"Grade: *{overall_grade}* &bull; Result: *{verdict}*\n"
        f"-----------------------------------------\n"
        f"📄 *View / Download Official Marksheet Online (No Login Needed):*\n"
        f"{public_slip_url}\n"
        f"-----------------------------------------\n"
        f"Examination Directorate, {school.name}"
    )
    
    ctx.update({
        'student': student,
        'exam': exam,
        'results': results,
        'total_obtained': total_obtained,
        'total_max': total_max,
        'overall_percentage': overall_percentage,
        'overall_grade': overall_grade,
        'verdict': verdict,
        'attendance_stats': attendance_stats,
        'clean_parent_phone': clean_phone,
        'wa_message': wa_message,
        'public_slip_url': public_slip_url,
    })
    return render(request, 'school_erp/report_card.html', ctx)


@require_menu_permission('can_manage_exams', 'Publish Results')
def exam_toggle_publish_view(request, exam_id):
    exam = get_object_or_404(Exam, pk=exam_id)
    exam.is_published = not exam.is_published
    exam.save()
    status_label = "published for viewing" if exam.is_published else "set back to draft"
    
    # Audit trail for Result Publication
    log_audit(
        request=request,
        action='PUBLISH',
        module='EXAMS',
        object_type='Exam',
        object_id=str(exam.id),
        object_repr=exam.name,
        description=f"Examination results for '{exam.name}' have been {status_label}.",
        reason="Official result publication authorization"
    )
    
    messages.success(request, f"Examination results for '{exam.name}' have been {status_label}.")
    return redirect('exam_list')


@require_menu_permission('can_manage_exams', 'Delete Examination')
def exam_delete_view(request, exam_id):
    exam = get_object_or_404(Exam, pk=exam_id)
    if request.method == 'POST':
        name = exam.name
        exam.results.all().delete()
        exam.delete()
        
        log_audit(
            request=request,
            action='DELETE',
            module='EXAMS',
            object_type='Exam',
            object_id=str(exam_id),
            object_repr=name,
            description=f"Examination '{name}' and all associated marks deleted.",
            reason="Examination deletion"
        )
        
        messages.success(request, f"Examination '{name}' and associated marks have been deleted.")
    return redirect('exam_list')


@require_menu_permission('can_manage_exams', 'Bulk Marksheets')
def bulk_report_cards_view(request, exam_id, class_id):
    ctx = get_school_context(request)
    exam = get_object_or_404(Exam, pk=exam_id)
    class_level = get_object_or_404(ClassLevel, pk=class_id)
    
    section_id = request.GET.get('section_id')
    students = Student.objects.filter(class_level=class_level, is_active=True).select_related('section', 'class_level').order_by('roll_number')
    if section_id:
        students = students.filter(section_id=section_id)
        
    student_cards = []
    for student in students:
        results = ExamResult.objects.filter(student=student, exam=exam).select_related('subject')
        total_obtained = sum(float(r.marks_obtained) for r in results)
        total_max = sum(float(r.max_marks) for r in results)
        overall_percentage = round((total_obtained / total_max * 100), 2) if total_max > 0 else 0
        
        if overall_percentage >= 90:
            overall_grade = 'A+'
            verdict = 'Outstanding Performance'
        elif overall_percentage >= 80:
            overall_grade = 'A'
            verdict = 'Excellent Performance'
        elif overall_percentage >= 70:
            overall_grade = 'B'
            verdict = 'Very Good'
        elif overall_percentage >= 60:
            overall_grade = 'C'
            verdict = 'Good / Satisfactory'
        elif overall_percentage >= 50:
            overall_grade = 'D'
            verdict = 'Needs Improvement'
        else:
            overall_grade = 'F'
            verdict = 'Failed / Re-appear'
            
        student_cards.append({
            'student': student,
            'results': results,
            'total_obtained': total_obtained,
            'total_max': total_max,
            'overall_percentage': overall_percentage,
            'overall_grade': overall_grade,
            'verdict': verdict,
            'attendance_stats': student.attendance_stats,
        })
        
    ctx.update({
        'exam': exam,
        'class_level': class_level,
        'student_cards': student_cards,
        'total_students': len(student_cards),
        'title': f'Bulk Marksheets: {class_level.name} - {exam.name}'
    })
    return render(request, 'school_erp/bulk_report_cards.html', ctx)


# ------------------- NOTICES & CIRCULARS -------------------

def notices_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    notices = Notice.objects.filter(school=school).order_by('-is_urgent', '-created_at')
    ctx.update({'notices': notices})
    return render(request, 'school_erp/notices.html', ctx)


def notice_create_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    if request.method == 'POST':
        form = NoticeForm(request.POST)
        if form.is_valid():
            notice = form.save(commit=False)
            notice.school = school
            notice.save()
            messages.success(request, f"Notice '{notice.title}' published successfully!")
            return redirect('notices')
    else:
        form = NoticeForm()
        
    ctx.update({'form': form, 'title': 'Publish New Circular / Notice'})
    return render(request, 'school_erp/notice_form.html', ctx)


# ==============================================================================
# ACADEMIC SESSIONS & SESSION SWITCHING
# ==============================================================================

def switch_academic_session_view(request):
    """Switch active academic session via POST."""
    current_school = get_current_school(request)
    if request.method == 'POST':
        session_id = request.POST.get('session_id')
        next_url = request.POST.get('next') or request.META.get('HTTP_REFERER') or '/'
        if session_id:
            session_obj = AcademicSession.objects.filter(id=session_id, school=current_school).first()
            if session_obj:
                request.session[f'active_session_{current_school.id}'] = session_obj.id
                request.session['active_academic_session_id'] = session_obj.id
                messages.success(request, f"Active Academic Session changed to {session_obj.name}")
        return redirect(next_url)
    return redirect('dashboard')


def switch_academic_session_get_view(request, session_id):
    """Switch active academic session via GET link."""
    current_school = get_current_school(request)
    session_obj = get_object_or_404(AcademicSession, id=session_id, school=current_school)
    request.session[f'active_session_{current_school.id}'] = session_obj.id
    request.session['active_academic_session_id'] = session_obj.id
    messages.success(request, f"Active Academic Session changed to {session_obj.name}")
    next_url = request.GET.get('next') or request.META.get('HTTP_REFERER') or '/'
    return redirect(next_url)


@require_menu_permission('can_manage_sessions', 'Academic Sessions')
def academic_sessions_list_view(request):
    """Manage and create academic sessions with forward-only validation."""
    current_school = get_current_school(request)
    ctx = get_school_context(request)
    active_session = get_active_session(request, school=current_school)

    if request.method == 'POST':
        action = request.POST.get('action', 'create')
        if action == 'create':
            name = request.POST.get('name', '').strip()
            start_date_str = request.POST.get('start_date', '').strip()
            end_date_str = request.POST.get('end_date', '').strip()
            is_current = request.POST.get('is_current') == 'on'

            if not name or not start_date_str or not end_date_str:
                messages.error(request, "Please fill in all session fields (Name, Start Date, End Date).")
            elif AcademicSession.objects.filter(school=current_school, name__iexact=name).exists():
                messages.error(request, f"Session with name '{name}' already exists in your school.")
            else:
                try:
                    start_date = datetime.datetime.strptime(start_date_str, '%Y-%m-%d').date()
                    end_date = datetime.datetime.strptime(end_date_str, '%Y-%m-%d').date()

                    # STRICT RULE: Cannot add previous or past sessions before school's anchor registration session!
                    if current_school.registered_session_start_date and start_date < current_school.registered_session_start_date:
                        reg_date_str = current_school.registered_session_start_date.strftime('%d %b %Y')
                        messages.error(
                            request,
                            f"Policy Restriction: Your school was registered starting from session '{current_school.registered_session_name}' "
                            f"({reg_date_str}). You cannot add or access any previous or back sessions. Only future or forward sessions are permitted."
                        )
                        return redirect('academic_sessions_manage')

                    if end_date <= start_date:
                        messages.error(request, "Session End Date must be after Start Date.")
                    else:
                        if is_current:
                            AcademicSession.objects.filter(school=current_school, is_current=True).update(is_current=False)
                        new_session = AcademicSession.objects.create(
                            school=current_school,
                            name=name,
                            start_date=start_date,
                            end_date=end_date,
                            is_current=is_current
                        )
                        if is_current:
                            request.session[f'active_session_{current_school.id}'] = new_session.id
                            request.session['active_academic_session_id'] = new_session.id
                        messages.success(request, f"Academic Session '{new_session.name}' created successfully!")
                        return redirect('academic_sessions_manage')
                except ValueError:
                    messages.error(request, "Invalid date format. Use YYYY-MM-DD.")
        elif action == 'set_current':
            session_id = request.POST.get('session_id')
            s_obj = get_object_or_404(AcademicSession, id=session_id, school=current_school)
            AcademicSession.objects.filter(school=current_school, is_current=True).update(is_current=False)
            s_obj.is_current = True
            s_obj.save()
            request.session[f'active_session_{current_school.id}'] = s_obj.id
            request.session['active_academic_session_id'] = s_obj.id
            messages.success(request, f"Session '{s_obj.name}' is now set as the current active session.")
            return redirect('academic_sessions_manage')

    sessions_qs = AcademicSession.objects.filter(school=current_school).order_by('-start_date')
    session_data = []
    for s in sessions_qs:
        students_count = Student.objects.filter(
            school=current_school
        ).filter(
            Q(academic_session=s) | Q(session_enrollments__academic_session=s)
        ).distinct().count()
        invoices_count = FeeInvoice.objects.filter(school=current_school, academic_session=s).count()
        session_data.append({
            'session': s,
            'students_count': students_count,
            'invoices_count': invoices_count,
            'is_active_selected': (active_session and active_session.id == s.id),
        })

    ctx.update({
        'sessions_list': session_data,
        'active_session': active_session,
    })
    return render(request, 'school_erp/academic_sessions.html', ctx)


# ==============================================================================
# STUDENT PROMOTION & MOVE TO NEXT SESSION
# ==============================================================================

@require_menu_permission('can_manage_sessions', 'Student Move / Promotion')
def student_promote_view(request, pk):
    """
    Move / Promote a single student to next session.
    Validation: Target session MUST have start_date > student.academic_session.start_date.
    Moving to past / previous session is strictly prevented!
    """
    ctx = get_school_context(request)
    school = ctx['school']
    student = get_object_or_404(Student, pk=pk, school=school)
    current_session = student.academic_session or get_active_session(request, school=school)

    next_sessions = AcademicSession.objects.filter(
        school=school,
        start_date__gt=current_session.start_date
    ).order_by('start_date')

    classes = ClassLevel.objects.filter(school=school).order_by('order', 'name')
    sections = Section.objects.filter(class_level__school=school)

    current_class = student.class_level
    next_class = ClassLevel.objects.filter(school=school, order__gt=current_class.order).order_by('order').first() or current_class

    if request.method == 'POST':
        target_session_id = request.POST.get('target_session_id')
        target_class_id = request.POST.get('target_class_id')
        target_section_id = request.POST.get('target_section_id')
        target_roll_number = request.POST.get('target_roll_number')
        remarks = request.POST.get('remarks', '').strip()

        target_session = get_object_or_404(AcademicSession, id=target_session_id, school=school)
        target_class = get_object_or_404(ClassLevel, id=target_class_id, school=school)
        target_section = get_object_or_404(Section, id=target_section_id, class_level__school=school)

        # STRICT VALIDATION: Cannot move to past or previous session!
        if target_session.start_date <= current_session.start_date:
            messages.error(
                request,
                f"Validation Error: You cannot move a student to an earlier or past session ({target_session.name})! "
                f"Students can only be moved or promoted to a next or future academic session."
            )
            return redirect('student_promote', pk=student.pk)

        previous_session = student.academic_session
        previous_class = student.class_level

        try:
            roll_num = int(target_roll_number) if target_roll_number else student.roll_number
        except ValueError:
            roll_num = student.roll_number

        student.academic_session = target_session
        student.class_level = target_class
        student.section = target_section
        student.roll_number = roll_num
        student.save()

        # Update or create new session enrollment
        new_enr, created = StudentSessionEnrollment.objects.get_or_create(
            student=student,
            academic_session=target_session,
            defaults={
                'class_level': target_class,
                'section': target_section,
                'roll_number': roll_num,
                'status': 'Promoted',
                'promoted_from_session': previous_session,
                'remarks': remarks or f"Promoted from {previous_class.name} ({previous_session.name if previous_session else ''})"
            }
        )
        if not created:
            new_enr.class_level = target_class
            new_enr.section = target_section
            new_enr.roll_number = roll_num
            new_enr.status = 'Promoted'
            new_enr.save()

        if previous_session:
            StudentSessionEnrollment.objects.filter(
                student=student, academic_session=previous_session
            ).update(status='Promoted')

        messages.success(
            request,
            f"Student {student.full_name} has been successfully moved & promoted to session {target_session.name} "
            f"({target_class.name} - Sec {target_section.name})!"
        )
        return redirect('student_detail', pk=student.pk)

    ctx.update({
        'student': student,
        'current_session': current_session,
        'next_sessions': next_sessions,
        'classes': classes,
        'sections': sections,
        'suggested_class': next_class,
    })
    return render(request, 'school_erp/student_promote.html', ctx)


@require_menu_permission('can_manage_sessions', 'Batch Promotion')
def student_bulk_promote_view(request):
    """
    Bulk Promote an entire class to the next academic session.
    """
    ctx = get_school_context(request)
    school = ctx['school']
    active_session = get_active_session(request, school=school)

    classes = ClassLevel.objects.filter(school=school).order_by('order', 'name')
    sections = Section.objects.filter(class_level__school=school)
    all_sessions = AcademicSession.objects.filter(school=school).order_by('-start_date')

    source_session_id = request.GET.get('source_session_id') or (active_session.id if active_session else '')
    source_class_id = request.GET.get('source_class_id', '')

    source_session = None
    if source_session_id:
        source_session = AcademicSession.objects.filter(id=source_session_id, school=school).first()

    next_sessions = []
    if source_session:
        next_sessions = AcademicSession.objects.filter(
            school=school,
            start_date__gt=source_session.start_date
        ).order_by('start_date')

    students_to_promote = []
    if source_session and source_class_id:
        students_to_promote = Student.objects.filter(
            school=school,
            is_active=True,
            class_level_id=source_class_id
        ).filter(
            Q(academic_session=source_session) | Q(session_enrollments__academic_session=source_session)
        ).distinct().order_by('roll_number')

    if request.method == 'POST':
        target_session_id = request.POST.get('target_session_id')
        target_class_id = request.POST.get('target_class_id')
        target_section_id = request.POST.get('target_section_id')
        selected_student_ids = request.POST.getlist('selected_students')

        if not target_session_id or not target_class_id:
            messages.error(request, "Please select both a Target Session and a Target Class.")
            return redirect('student_bulk_promote')

        target_session = get_object_or_404(AcademicSession, id=target_session_id, school=school)
        target_class = get_object_or_404(ClassLevel, id=target_class_id, school=school)
        target_section = Section.objects.filter(id=target_section_id, class_level__school=school).first() if target_section_id else None

        # Validation: Cannot promote to previous/past session
        if source_session and target_session.start_date <= source_session.start_date:
            messages.error(
                request,
                f"Validation Error: Target session ({target_session.name}) is earlier than or equal to current session ({source_session.name})! "
                f"Students can only be promoted to a next or future academic session."
            )
            return redirect('student_bulk_promote')

        promoted_count = 0
        for sid in selected_student_ids:
            try:
                st = Student.objects.get(id=sid)
                prev_s = st.academic_session
                prev_c = st.class_level

                sec = target_section or st.section
                st.academic_session = target_session
                st.class_level = target_class
                st.section = sec
                st.save()

                StudentSessionEnrollment.objects.update_or_create(
                    student=st,
                    academic_session=target_session,
                    defaults={
                        'class_level': target_class,
                        'section': sec,
                        'roll_number': st.roll_number,
                        'status': 'Promoted',
                        'promoted_from_session': prev_s,
                        'remarks': f"Bulk Promoted from {prev_c.name} ({prev_s.name if prev_s else ''})"
                    }
                )
                if prev_s:
                    StudentSessionEnrollment.objects.filter(
                        student=st, academic_session=prev_s
                    ).update(status='Promoted')
                promoted_count += 1
            except Student.DoesNotExist:
                continue

        messages.success(
            request,
            f"Success! {promoted_count} students successfully promoted to session '{target_session.name}' in {target_class.name}!"
        )
        return redirect('student_list')

    source_class = None
    suggested_target_class = None
    cleared_fee_count = 0
    pending_fee_count = 0

    if source_class_id:
        source_class = ClassLevel.objects.filter(id=source_class_id).first()
        if source_class:
            suggested_target_class = ClassLevel.objects.filter(order__gt=source_class.order).order_by('order').first()
        cleared_fee_count = sum(1 for s in students_to_promote if s.fee_due_amount == Decimal('0.00'))
        pending_fee_count = sum(1 for s in students_to_promote if s.fee_due_amount > Decimal('0.00'))

    ctx.update({
        'classes': classes,
        'sections': sections,
        'all_sessions': all_sessions,
        'source_session': source_session,
        'source_class': source_class,
        'suggested_target_class': suggested_target_class,
        'source_session_id': int(source_session_id) if source_session_id else None,
        'source_class_id': int(source_class_id) if source_class_id else None,
        'next_sessions': next_sessions,
        'students_to_promote': students_to_promote,
        'cleared_fee_count': cleared_fee_count,
        'pending_fee_count': pending_fee_count,
    })
    return render(request, 'school_erp/student_bulk_promote.html', ctx)


# Note: User Management & Menu Permission views are defined in the dedicated section below.


# ==============================================================================
# AUTHENTICATION: LOGIN & LOGOUT
# ==============================================================================

# ==============================================================================
# AUTHENTICATION: LOGIN & LOGOUT
# ==============================================================================

def user_login_view(request):
    """Staff & Admin login view with security monitoring, lockout prevention & MFA support."""
    if request.user.is_authenticated:
        return redirect('dashboard')

    next_url = request.GET.get('next') or request.POST.get('next') or 'dashboard'
    current_school = get_school_context(request)['school']
    client_ip = get_client_ip(request)
    user_agent = request.META.get('HTTP_USER_AGENT', '')[:290]

    sec_config = None
    if current_school:
        sec_config, _ = SecurityConfiguration.objects.get_or_create(school=current_school)

    # Check IP Lockout if security config exists
    if sec_config and current_school and client_ip:
        lockout_time = timezone.now() - datetime.timedelta(minutes=sec_config.lockout_duration_minutes)
        recent_failures = LoginAttemptLog.objects.filter(
            school=current_school,
            ip_address=client_ip,
            status='Failed',
            timestamp__gte=lockout_time
        ).count()
        if recent_failures >= sec_config.max_login_failures:
            LoginAttemptLog.objects.create(
                school=current_school,
                username=request.POST.get('username', 'Unknown').strip(),
                ip_address=client_ip,
                user_agent=user_agent,
                status='Locked_Out',
                failure_reason=f"Account locked out for {sec_config.lockout_duration_minutes} minutes due to excessive failed attempts.",
                is_suspicious=True
            )
            messages.error(
                request,
                f"Security Lockout: Too many failed login attempts. Your IP has been temporarily locked for {sec_config.lockout_duration_minutes} minutes. Please contact IT support."
            )
            return render(request, 'school_erp/login.html', {
                'next': next_url,
                'school': current_school,
                'locked_out': True
            })

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '').strip()
        mfa_pin = request.POST.get('mfa_pin', '').strip()
        user = authenticate(request, username=username, password=password)

        if user is not None:
            if not user.is_active:
                LoginAttemptLog.objects.create(
                    school=current_school,
                    username=username,
                    ip_address=client_ip,
                    user_agent=user_agent,
                    status='Failed',
                    failure_reason="User account is deactivated."
                )
                messages.error(request, "Your individual user account has been disabled. Please contact administrator.")
                return render(request, 'school_erp/login.html', {'next': next_url, 'school': current_school})

            profile_obj = getattr(user, 'profile', None)
            is_superadmin = user.is_superuser or (profile_obj and profile_obj.is_platform_superadmin)

            # Check tenant school subscription & status for non-superadmins
            if not is_superadmin and profile_obj and profile_obj.school:
                user_school = profile_obj.school
                if not user_school.is_active:
                    LoginAttemptLog.objects.create(
                        school=user_school,
                        username=username,
                        ip_address=client_ip,
                        user_agent=user_agent,
                        status='Failed',
                        failure_reason="School account deactivated."
                    )
                    messages.error(
                        request,
                        "Access Denied: Your school account has been deactivated by the Platform Administrator. Please contact support."
                    )
                    return render(request, 'school_erp/login.html', {'next': next_url, 'school': user_school})

                if user_school.days_remaining < 0:
                    LoginAttemptLog.objects.create(
                        school=user_school,
                        username=username,
                        ip_address=client_ip,
                        user_agent=user_agent,
                        status='Failed',
                        failure_reason="School subscription expired."
                    )
                    messages.error(
                        request,
                        f"Subscription Expired: Your school's monthly subscription (₹{user_school.subscription_monthly_fee}/month) "
                        f"expired on {user_school.subscription_end_date}. Please contact the Platform Administrator to renew."
                    )
                    return render(request, 'school_erp/login.html', {'next': next_url, 'school': user_school})

            # Check MFA if enabled
            if sec_config and sec_config.enable_mfa and sec_config.mfa_security_pin:
                if mfa_pin != sec_config.mfa_security_pin:
                    LoginAttemptLog.objects.create(
                        school=current_school,
                        username=username,
                        ip_address=client_ip,
                        user_agent=user_agent,
                        status='Failed',
                        failure_reason="Invalid 2FA / Security PIN provided.",
                        is_suspicious=True
                    )
                    messages.error(request, "Invalid 2FA Security PIN. Please enter your authorized PIN to proceed.")
                    return render(request, 'school_erp/login.html', {
                        'next': next_url,
                        'school': current_school,
                        'require_mfa': True,
                        'username_val': username
                    })

            # Successful Login
            login(request, user)

            # Set session timeout based on SecurityConfiguration
            if sec_config:
                request.session.set_expiry(sec_config.session_expiry_minutes * 60)

            LoginAttemptLog.objects.create(
                school=current_school,
                username=username,
                ip_address=client_ip,
                user_agent=user_agent,
                status='Success',
                failure_reason=""
            )

            messages.success(request, f"Welcome back, {user.first_name or user.username}!")
            if is_superadmin and next_url == 'dashboard':
                return redirect('saas_admin_dashboard')
            return redirect(next_url)
        else:
            LoginAttemptLog.objects.create(
                school=current_school,
                username=username or 'Unknown',
                ip_address=client_ip,
                user_agent=user_agent,
                status='Failed',
                failure_reason="Invalid credentials supplied."
            )
            messages.error(request, "Invalid username or password. Please try again.")

    return render(request, 'school_erp/login.html', {
        'next': next_url,
        'school': current_school,
        'require_mfa': (sec_config.enable_mfa if sec_config else False)
    })


def user_logout_view(request):
    """User logout view."""
    logout(request)
    messages.info(request, "You have been logged out.")
    return redirect('user_login')


# ==============================================================================
# SAAS PLATFORM SUPER ADMIN & MULTI-TENANCY MANAGEMENT
# ==============================================================================

def is_superadmin_user(user):
    """Check if user has Platform Super Admin privileges."""
    if not user or not user.is_authenticated:
        return False
    if user.is_superuser:
        return True
    profile = getattr(user, 'profile', None)
    return bool(profile and profile.is_platform_superadmin)


def require_superadmin(view_func):
    """View decorator restricting access strictly to Platform Super Admins."""
    @wraps(view_func)
    def _wrapped(request, *args, **kwargs):
        if not is_superadmin_user(request.user):
            messages.error(request, "Access Denied: This section is restricted exclusively to SaaS Platform Super Administrators.")
            return redirect('dashboard')
        return view_func(request, *args, **kwargs)
    return _wrapped


@require_superadmin
def saas_admin_dashboard_view(request):
    """SaaS Platform Super Admin Portal to manage tenant schools, subscriptions (₹499/mo), and statuses."""
    schools_qs = SchoolProfile.objects.all().order_by('-created_at')
    
    # Search filter
    q = request.GET.get('q', '').strip()
    if q:
        schools_qs = schools_qs.filter(
            Q(name__icontains=q) | Q(school_code__icontains=q) | Q(email__icontains=q) | Q(phone__icontains=q)
        )

    # Status tab filter
    status_filter = request.GET.get('status', 'all')
    
    # Calculate global metrics
    all_schools = list(SchoolProfile.objects.all())
    total_schools_count = len(all_schools)
    active_schools_count = sum(1 for s in all_schools if s.is_active and s.days_remaining >= 0)
    expiring_soon_count = sum(1 for s in all_schools if s.is_active and 0 <= s.days_remaining <= 7)
    expired_count = sum(1 for s in all_schools if s.days_remaining < 0 or not s.is_active)
    monthly_revenue = sum((s.subscription_monthly_fee for s in all_schools if s.is_active and s.days_remaining >= 0), Decimal('0.00'))

    # Build detailed school items with tenant metrics
    filtered_schools = []
    for s in schools_qs:
        admin_profile = UserProfile.objects.filter(school=s, role='Administrator').first()
        student_count = Student.objects.filter(school=s).count()
        teacher_count = Teacher.objects.filter(school=s).count()
        
        item = {
            'school': s,
            'admin_user': admin_profile.user if admin_profile else None,
            'student_count': student_count,
            'teacher_count': teacher_count,
            'days_remaining': s.days_remaining,
            'status_badge_class': s.status_badge_class,
            'subscription_status': s.subscription_status,
        }

        if status_filter == 'active' and (not s.is_active or s.days_remaining < 0):
            continue
        elif status_filter == 'expiring' and not (s.is_active and 0 <= s.days_remaining <= 7):
            continue
        elif status_filter == 'expired' and not (s.days_remaining < 0 or not s.is_active):
            continue

        filtered_schools.append(item)

    ctx = {
        'schools': filtered_schools,
        'total_schools_count': total_schools_count,
        'active_schools_count': active_schools_count,
        'expiring_soon_count': expiring_soon_count,
        'expired_count': expired_count,
        'monthly_revenue': monthly_revenue,
        'status_filter': status_filter,
        'search_query': q,
        'today': timezone.now().date(),
    }
    return render(request, 'school_erp/saas_admin.html', ctx)


@require_superadmin
def saas_school_create_view(request):
    """Register a new School tenant, configure anchor session, default classes, and admin credentials."""
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        short_name = request.POST.get('short_name', '').strip()
        school_code = request.POST.get('school_code', '').strip().upper()
        email = request.POST.get('email', '').strip()
        phone = request.POST.get('phone', '').strip()
        address = request.POST.get('address', '').strip()
        affiliation_number = request.POST.get('affiliation_number', '').strip()
        board = request.POST.get('board', 'CBSE').strip()
        principal_name = request.POST.get('principal_name', '').strip()

        # Anchor Session
        session_name = request.POST.get('session_name', '2026-2027').strip()
        start_date_str = request.POST.get('session_start_date', '2026-04-01').strip()
        end_date_str = request.POST.get('session_end_date', '2027-03-31').strip()

        # Admin User
        admin_username = request.POST.get('admin_username', '').strip()
        admin_password = request.POST.get('admin_password', '').strip()
        admin_name = request.POST.get('admin_name', '').strip()
        admin_email = request.POST.get('admin_email', '').strip() or email

        # Auto-generate unique school code if not provided or empty
        if not school_code and name:
            school_code = generate_unique_school_code(name)

        # Subscription
        monthly_fee_str = request.POST.get('subscription_fee', '499.00').strip()
        sub_days_str = request.POST.get('subscription_days', '30').strip()

        # Validation
        if not name or not school_code or not session_name or not admin_username or not admin_password:
            messages.error(request, "School Name, Anchor Session, Admin Username, and Admin Password are required.")
            return redirect('saas_school_create')

        if SchoolProfile.objects.filter(school_code__iexact=school_code).exists():
            school_code = generate_unique_school_code(name)

        if User.objects.filter(username__iexact=admin_username).exists():
            messages.error(request, f"User '{admin_username}' already exists. Please choose a different admin username.")
            return redirect('saas_school_create')

        try:
            reg_start_date = datetime.datetime.strptime(start_date_str, '%Y-%m-%d').date()
            reg_end_date = datetime.datetime.strptime(end_date_str, '%Y-%m-%d').date()
        except ValueError:
            reg_start_date = datetime.date(2026, 4, 1)
            reg_end_date = datetime.date(2027, 3, 31)

        try:
            monthly_fee = Decimal(monthly_fee_str)
        except Exception:
            monthly_fee = Decimal('499.00')

        try:
            sub_days = int(sub_days_str)
        except Exception:
            sub_days = 30

        today = timezone.now().date()
        sub_end_date = today + datetime.timedelta(days=sub_days)

        # 1. Create SchoolProfile
        school = SchoolProfile.objects.create(
            name=name,
            short_name=short_name or name[:25],
            school_code=school_code,
            email=email,
            phone=phone,
            address=address,
            affiliation_number=affiliation_number,
            board=board,
            principal_name=principal_name,
            academic_year=session_name,
            registered_session_name=session_name,
            registered_session_start_date=reg_start_date,
            subscription_monthly_fee=monthly_fee,
            subscription_start_date=today,
            subscription_end_date=sub_end_date,
            is_active=True
        )

        # 2. Create Anchor Session
        anchor_session = AcademicSession.objects.create(
            school=school,
            name=session_name,
            start_date=reg_start_date,
            end_date=reg_end_date,
            is_current=True
        )

        # 3. Classes & Sections: Default is empty (0 classes) so new school starts fresh without foreign classes
        seed_classes = request.POST.get('seed_default_classes') == 'on'
        if seed_classes:
            default_classes = [
                ('Nursery', 'NUR', 1), ('LKG', 'LKG', 2), ('UKG', 'UKG', 3),
                ('Class 1', 'CLS1', 4), ('Class 2', 'CLS2', 5), ('Class 3', 'CLS3', 6),
                ('Class 4', 'CLS4', 7), ('Class 5', 'CLS5', 8), ('Class 6', 'CLS6', 9),
                ('Class 7', 'CLS7', 10), ('Class 8', 'CLS8', 11), ('Class 9', 'CLS9', 12),
                ('Class 10', 'CLS10', 13)
            ]
            for c_name, c_code, c_order in default_classes:
                cl = ClassLevel.objects.create(school=school, name=c_name, code=c_code, order=c_order)
                Section.objects.create(class_level=cl, name='A', capacity=40, room_number='Room 101')
                ClassFeeMaster.objects.create(class_level=cl)

        # 4. Create Default Fee Categories
        default_categories = [
            ('Tuition Fee', Decimal('2500.00'), 'Monthly academic tuition fee'),
            ('Admission Fee', Decimal('5000.00'), 'One-time admission and registration fee'),
            ('Examination Fee', Decimal('800.00'), 'Term end examination fee'),
            ('Computer & Lab Fee', Decimal('600.00'), 'IT lab, computer & STEM facility fee'),
            ('Transport Fee', Decimal('1500.00'), 'Monthly school bus transportation fee')
        ]
        for f_name, f_amt, f_desc in default_categories:
            FeeCategory.objects.create(school=school, name=f_name, default_amount=f_amt, description=f_desc)

        # 5. Create School Admin User
        admin_user = User.objects.create_user(
            username=admin_username,
            password=admin_password,
            email=admin_email,
            first_name=admin_name or "School Admin"
        )
        UserProfile.objects.create(
            user=admin_user,
            school=school,
            role='Administrator',
            phone=phone,
            can_manage_students=True,
            can_manage_attendance=True,
            can_manage_fees=True,
            can_manage_exams=True,
            can_manage_timetable=True,
            can_manage_classes=True,
            can_manage_sessions=True,
            can_manage_users=True,
            is_platform_superadmin=False
        )

        messages.success(
            request,
            f"School '{school.name}' ({school.school_code}) successfully onboarded with anchor session '{anchor_session.name}'! "
            f"Admin account '{admin_user.username}' created. ₹{monthly_fee}/month subscription is active for {sub_days} days (until {sub_end_date})."
        )
        return redirect('saas_admin_dashboard')

    return render(request, 'school_erp/saas_school_create.html')


@require_superadmin
def saas_school_renew_view(request, school_id):
    """Extend school subscription by +30 days (or custom days) upon payment and activate account."""
    if request.method == 'POST':
        school = get_object_or_404(SchoolProfile, id=school_id)
        days_to_add = int(request.POST.get('days', 30))
        today = timezone.now().date()

        if school.subscription_end_date and school.subscription_end_date >= today:
            school.subscription_end_date = school.subscription_end_date + datetime.timedelta(days=days_to_add)
        else:
            school.subscription_start_date = today
            school.subscription_end_date = today + datetime.timedelta(days=days_to_add)

        school.is_active = True
        school.save()

        messages.success(
            request,
            f"Payment recorded! Subscription for '{school.name}' successfully extended by {days_to_add} days "
            f"(valid until {school.subscription_end_date}). School account has been activated."
        )
    return redirect('saas_admin_dashboard')


@require_superadmin
def saas_school_toggle_status_view(request, school_id):
    """Deactivate or activate school login access."""
    if request.method == 'POST':
        school = get_object_or_404(SchoolProfile, id=school_id)
        school.is_active = not school.is_active
        school.save()
        status_text = "Activated (Access Restored)" if school.is_active else "Deactivated / Suspended (Login Locked)"
        messages.success(request, f"School '{school.name}' has been {status_text}.")
    return redirect('saas_admin_dashboard')


@require_superadmin
def saas_switch_school_view(request, school_id):
    """Platform Super Admin enters and manages a specific tenant school."""
    school = get_object_or_404(SchoolProfile, id=school_id)
    request.session['saas_active_school_id'] = school.id
    request.session.pop('active_academic_session_id', None)
    messages.success(request, f"Switched to School: '{school.name}' ({school.school_code}). You are now managing this school's ERP.")
    return redirect('dashboard')


@require_superadmin
def saas_school_edit_view(request, school_id):
    """Platform Super Admin edits school profile details, anchor session baseline, and subscription info."""
    school = get_object_or_404(SchoolProfile, id=school_id)
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        short_name = request.POST.get('short_name', '').strip()
        school_code = request.POST.get('school_code', '').strip().upper()
        email = request.POST.get('email', '').strip()
        phone = request.POST.get('phone', '').strip()
        address = request.POST.get('address', '').strip()
        board = request.POST.get('board', '').strip()
        affiliation_number = request.POST.get('affiliation_number', '').strip()
        principal_name = request.POST.get('principal_name', '').strip()
        monthly_fee_str = request.POST.get('subscription_monthly_fee', '499.00').strip()
        sub_end_date_str = request.POST.get('subscription_end_date', '').strip()
        session_name = request.POST.get('registered_session_name', '').strip()
        session_start_date_str = request.POST.get('registered_session_start_date', '').strip()

        if not name or not school_code:
            messages.error(request, "School Name and School Code are required.")
            return redirect('saas_school_edit', school_id=school.id)

        if SchoolProfile.objects.exclude(id=school.id).filter(school_code__iexact=school_code).exists():
            messages.error(request, f"Another school with code '{school_code}' already exists. Please choose a unique code.")
            return redirect('saas_school_edit', school_id=school.id)

        school.name = name
        school.short_name = short_name or name[:50]
        school.school_code = school_code
        school.email = email
        school.phone = phone
        school.address = address
        school.board = board
        school.affiliation_number = affiliation_number
        school.principal_name = principal_name

        try:
            school.subscription_monthly_fee = Decimal(monthly_fee_str)
        except Exception:
            pass

        if sub_end_date_str:
            try:
                school.subscription_end_date = datetime.datetime.strptime(sub_end_date_str, '%Y-%m-%d').date()
            except ValueError:
                pass

        if session_name:
            school.registered_session_name = session_name
            school.academic_year = session_name

        if session_start_date_str:
            try:
                school.registered_session_start_date = datetime.datetime.strptime(session_start_date_str, '%Y-%m-%d').date()
            except ValueError:
                pass

        school.save()
        messages.success(request, f"School '{school.name}' ({school.school_code}) details updated successfully!")
        return redirect('saas_admin_dashboard')

    return render(request, 'school_erp/saas_school_edit.html', {'school': school})


@require_superadmin
def saas_school_reset_password_view(request, school_id):
    """Platform Super Admin resets or sets password for the school's admin account."""
    school = get_object_or_404(SchoolProfile, id=school_id)
    if request.method == 'POST':
        new_password = request.POST.get('new_password', '').strip()
        admin_username = request.POST.get('admin_username', '').strip()

        if not new_password:
            messages.error(request, "New password cannot be empty. Please enter a valid password.")
            return redirect('saas_admin_dashboard')

        if len(new_password) < 6:
            messages.error(request, "Password must be at least 6 characters long.")
            return redirect('saas_admin_dashboard')

        # Find existing admin user for this school
        target_user = None
        if admin_username:
            target_user = User.objects.filter(username__iexact=admin_username).first()

        if not target_user:
            admin_profile = UserProfile.objects.filter(school=school, role='Administrator').first()
            if admin_profile:
                target_user = admin_profile.user

        if target_user:
            target_user.set_password(new_password)
            target_user.save()
            messages.success(
                request,
                f"Password for School Admin '{target_user.username}' (School: '{school.name}') successfully changed to '{new_password}'! "
                f"You can now share this new password with the school."
            )
        else:
            # If school doesn't have an admin user yet, create one
            username_to_create = admin_username or f"admin_{school.school_code.lower().replace('-', '_')}"
            if User.objects.filter(username__iexact=username_to_create).exists():
                username_to_create = f"{username_to_create}_{school.id}"

            new_user = User.objects.create_user(
                username=username_to_create,
                password=new_password,
                email=school.email,
                first_name="School Admin"
            )
            UserProfile.objects.create(
                user=new_user,
                school=school,
                role='Administrator',
                phone=school.phone,
                can_manage_students=True,
                can_manage_attendance=True,
                can_manage_fees=True,
                can_manage_exams=True,
                can_manage_timetable=True,
                can_manage_classes=True,
                can_manage_sessions=True,
                can_manage_users=True,
                is_platform_superadmin=False
            )
            messages.success(
                request,
                f"New School Admin user '{new_user.username}' created for '{school.name}' with password '{new_password}'! "
                f"You can now share these credentials with the school."
            )

    return redirect('saas_admin_dashboard')


# ==============================================================================
# SCHOOL PROFILE & LOGO BRANDING SETTINGS
# ==============================================================================

@require_menu_permission('can_manage_users', 'School Settings')
def school_settings_view(request):
    """Allow School Admin to update school details, logo, and principal seal/signature for slips."""
    school = get_current_school(request)
    
    if request.method == 'POST':
        form = SchoolSettingsForm(request.POST, request.FILES, instance=school)
        if form.is_valid():
            if request.POST.get('clear_logo') == '1':
                school.logo = None
            if request.POST.get('clear_signature') == '1':
                school.principal_signature = None
            form.save()
            messages.success(
                request,
                f"School details and branding for '{school.name}' updated successfully! "
                f"Your logo and details are now dynamically applied to all Fee Slips, ID Cards, and Report Cards."
            )
            return redirect('school_settings')
        else:
            err_msg = ", ".join(f"{k}: {', '.join(v)}" for k, v in form.errors.items())
            messages.error(request, f"Error saving settings: {err_msg}")
    else:
        form = SchoolSettingsForm(instance=school)

    ctx = {
        'school': school,
        'form': form,
    }
    return render(request, 'school_erp/school_settings.html', ctx)


# ==============================================================================
# SMART REPORTS & EXECUTIVE ANALYTICS HUB
# ==============================================================================

def reports_hub_view(request):
    """
    Comprehensive Smart Reports Hub.
    Covers Fee Collection & Revenue, Attendance Register Analytics,
    Class Strength & Demographics, and Exam Merit / Performance.
    """
    if not request.user.is_authenticated:
        return redirect('user_login')

    ctx = get_school_context(request)
    active_session = get_active_session(request)
    report_type = request.GET.get('report', 'fees')
    
    classes = ClassLevel.objects.all().order_by('order', 'name')
    exams = Exam.objects.all().order_by('-start_date')
    
    # Global metrics
    all_students = Student.objects.filter(is_active=True)
    if active_session:
        all_students = all_students.filter(
            Q(academic_session=active_session) | Q(session_enrollments__academic_session=active_session)
        ).distinct()
    total_students = all_students.count()
    boys_count = all_students.filter(gender='Male').count()
    girls_count = all_students.filter(gender='Female').count()

    # 1. Fee Collection Report Logic
    date_filter = request.GET.get('date_filter', 'month')
    payment_mode = request.GET.get('mode', '').strip()
    class_id = request.GET.get('class_id', '').strip()
    custom_start = request.GET.get('start_date', '').strip()
    custom_end = request.GET.get('end_date', '').strip()

    today = timezone.now().date()
    if date_filter == 'today':
        start_d, end_d = today, today
    elif date_filter == 'yesterday':
        start_d, end_d = today - datetime.timedelta(days=1), today - datetime.timedelta(days=1)
    elif date_filter == 'week':
        start_d, end_d = today - datetime.timedelta(days=7), today
    elif date_filter == 'month':
        start_d = today.replace(day=1)
        end_d = today
    elif date_filter == 'custom' and custom_start and custom_end:
        try:
            start_d = datetime.datetime.strptime(custom_start, '%Y-%m-%d').date()
            end_d = datetime.datetime.strptime(custom_end, '%Y-%m-%d').date()
        except Exception:
            start_d, end_d = today.replace(day=1), today
    else:
        start_d, end_d = today.replace(day=1), today

    receipts_qs = FeeReceipt.objects.select_related('student__class_level', 'student__section').prefetch_related('items')
    receipts_qs = receipts_qs.filter(payment_date__range=[start_d, end_d])

    if payment_mode:
        receipts_qs = receipts_qs.filter(payment_method=payment_mode)
    if class_id:
        receipts_qs = receipts_qs.filter(student__class_level_id=class_id)

    total_fee_collected = sum((r.amount_paid for r in receipts_qs), Decimal('0.00'))
    total_receipts_count = receipts_qs.count()

    mode_summary = {}
    for r in receipts_qs:
        mode_summary[r.payment_method] = mode_summary.get(r.payment_method, Decimal('0.00')) + (r.amount_paid or Decimal('0.00'))

    # 2. Attendance Analytics Logic
    att_class_id = request.GET.get('att_class_id') or (classes.first().id if classes.exists() else None)
    selected_att_class = ClassLevel.objects.filter(id=att_class_id).first() if att_class_id else None
    attendance_rows = []
    if selected_att_class:
        students_in_class = Student.objects.filter(class_level=selected_att_class, is_active=True).order_by('section__name', 'roll_number')
        for s in students_in_class:
            stats = s.attendance_stats
            attendance_rows.append({
                'student': s,
                'total': stats['total'],
                'present': stats['present'],
                'late': stats['late'],
                'percentage': stats['percentage'],
                'is_short': stats['percentage'] < 75.0 and stats['total'] > 0
            })

    # 3. Student Strength Demographics
    strength_rows = []
    total_cap = 0
    total_enr = 0
    total_b = 0
    total_g = 0
    for cl in classes:
        for sec in cl.sections.all():
            st_qs = sec.students.filter(is_active=True)
            if active_session:
                st_qs = st_qs.filter(
                    Q(academic_session=active_session) | Q(session_enrollments__academic_session=active_session)
                ).distinct()
            b_cnt = st_qs.filter(gender='Male').count()
            g_cnt = st_qs.filter(gender='Female').count()
            enr = st_qs.count()
            vacant = max(0, sec.capacity - enr)
            total_cap += sec.capacity
            total_enr += enr
            total_b += b_cnt
            total_g += g_cnt
            strength_rows.append({
                'class_level': cl,
                'section': sec,
                'capacity': sec.capacity,
                'enrolled': enr,
                'boys': b_cnt,
                'girls': g_cnt,
                'vacant': vacant,
                'class_teacher': sec.class_teacher,
            })

    # 4. Exam Merit List Logic
    exam_id = request.GET.get('exam_id') or (exams.first().id if exams.exists() else None)
    exam_class_id = request.GET.get('exam_class_id') or (classes.first().id if classes.exists() else None)
    selected_exam = Exam.objects.filter(id=exam_id).first() if exam_id else None
    selected_exam_class = ClassLevel.objects.filter(id=exam_class_id).first() if exam_class_id else None

    merit_list = []
    if selected_exam and selected_exam_class:
        students_in_exam = Student.objects.filter(class_level=selected_exam_class, is_active=True).order_by('roll_number')
        for s in students_in_exam:
            results = ExamResult.objects.filter(student=s, exam=selected_exam)
            if results.exists():
                tot_ob = sum(float(r.marks_obtained) for r in results)
                tot_max = sum(float(r.max_marks) for r in results)
                pct = round((tot_ob / tot_max * 100), 2) if tot_max > 0 else 0
                has_failed = any(r.grade == 'F' for r in results)
                merit_list.append({
                    'student': s,
                    'obtained': tot_ob,
                    'max': tot_max,
                    'percentage': pct,
                    'grade': 'F' if has_failed else ('A+' if pct >= 90 else ('A' if pct >= 80 else ('B' if pct >= 70 else 'C' if pct >= 60 else 'D' if pct >= 50 else 'F'))),
                    'result': 'Fail' if has_failed else 'Pass',
                    'subject_count': results.count()
                })
        merit_list.sort(key=lambda x: x['percentage'], reverse=True)
        for idx, item in enumerate(merit_list, 1):
            item['rank'] = idx

    ctx.update({
        'report_type': report_type,
        'classes': classes,
        'exams': exams,
        'total_students': total_students,
        'boys_count': boys_count,
        'girls_count': girls_count,
        
        # Fees
        'receipts': receipts_qs,
        'total_fee_collected': total_fee_collected,
        'total_receipts_count': total_receipts_count,
        'mode_summary': mode_summary,
        'date_filter': date_filter,
        'start_date': start_d.strftime('%Y-%m-%d'),
        'end_date': end_d.strftime('%Y-%m-%d'),
        'selected_mode': payment_mode,
        'selected_class': class_id,
        
        # Attendance
        'att_class_id': int(att_class_id) if att_class_id else None,
        'selected_att_class': selected_att_class,
        'attendance_rows': attendance_rows,
        
        # Strength
        'strength_rows': strength_rows,
        'total_cap': total_cap,
        'total_enr': total_enr,
        'total_b': total_b,
        'total_g': total_g,
        
        # Exam Merit
        'selected_exam': selected_exam,
        'selected_exam_class': selected_exam_class,
        'merit_list': merit_list,
        'title': 'Smart Reports & Executive Analytics',
    })
    return render(request, 'school_erp/reports_hub.html', ctx)


def export_fees_report_excel(request):
    """Export filtered Fee Collection Register to formatted Excel spreadsheet."""
    start_str = request.GET.get('start_date')
    end_str = request.GET.get('end_date')
    payment_mode = request.GET.get('mode', '').strip()
    class_id = request.GET.get('class_id', '').strip()

    today = timezone.now().date()
    try:
        start_d = datetime.datetime.strptime(start_str, '%Y-%m-%d').date() if start_str else today.replace(day=1)
        end_d = datetime.datetime.strptime(end_str, '%Y-%m-%d').date() if end_str else today
    except Exception:
        start_d, end_d = today.replace(day=1), today

    receipts = FeeReceipt.objects.select_related('student__class_level', 'student__section').filter(
        payment_date__range=[start_d, end_d]
    ).order_by('-payment_date', '-created_at')

    if payment_mode:
        receipts = receipts.filter(payment_method=payment_mode)
    if class_id:
        receipts = receipts.filter(student__class_level_id=class_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fee_Collection"

    school = get_school_context(request)['school']

    # Title block
    ws.merge_cells("A1:H1")
    ws["A1"] = f"{school.name} — Fee Collection Register"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = f"Period: {start_d.strftime('%d-%b-%Y')} to {end_d.strftime('%d-%b-%Y')}  |  Payment Mode: {payment_mode or 'All Modes'}"
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color="64748B")
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["Receipt No", "Date", "Adm No", "Student Name", "Class & Section", "Parent Phone", "Payment Mode", "Amount Paid (Rs.)"]
    ws.append([])
    ws.append(headers)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    total_amount = Decimal('0.00')
    row_font = Font(name="Arial", size=10)

    for r in receipts:
        ws.append([
            r.receipt_number,
            r.payment_date.strftime('%d/%m/%Y'),
            r.student.admission_number,
            r.student.full_name,
            f"{r.student.class_level.name} - Sec {r.student.section.name}",
            r.student.parent_phone,
            r.payment_method,
            float(r.amount_paid)
        ])
        curr_row = ws.max_row
        total_amount += r.amount_paid
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=curr_row, column=col_idx)
            c.font = row_font
            c.border = thin_border
            if col_idx in [1, 2, 3, 5, 6, 7]:
                c.alignment = Alignment(horizontal="center")
            elif col_idx == 8:
                c.alignment = Alignment(horizontal="right")

    # Total row
    ws.append([])
    tot_row = ws.max_row
    ws.merge_cells(f"A{tot_row}:G{tot_row}")
    ws.cell(row=tot_row, column=1).value = "TOTAL AMOUNT COLLECTED:"
    ws.cell(row=tot_row, column=1).font = Font(name="Arial", size=11, bold=True, color="0F172A")
    ws.cell(row=tot_row, column=1).alignment = Alignment(horizontal="right")
    
    tot_cell = ws.cell(row=tot_row, column=8)
    tot_cell.value = float(total_amount)
    tot_cell.font = Font(name="Arial", size=12, bold=True, color="059669")
    tot_cell.alignment = Alignment(horizontal="right")

    # Column widths
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Fee_Collection_Report_{start_d.strftime('%Y%m%d')}_{end_d.strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def export_defaulters_excel_view(request):
    """Export Pending Fee Defaulters list to styled Excel spreadsheet."""
    class_id = request.GET.get('class_id', '').strip()
    search = request.GET.get('q', '').strip()
    min_due = request.GET.get('min_due', '').strip()

    school = get_school_context(request)['school']
    active_session = get_active_session(request)

    students_qs = Student.objects.filter(is_active=True).select_related('class_level', 'section').prefetch_related('invoices__payments')
    if active_session:
        students_qs = students_qs.filter(
            Q(academic_session=active_session) | Q(session_enrollments__academic_session=active_session)
        ).distinct()

    if class_id:
        students_qs = students_qs.filter(class_level_id=class_id)
    if search:
        students_qs = students_qs.filter(
            Q(first_name__icontains=search) | Q(last_name__icontains=search) |
            Q(admission_number__icontains=search) | Q(parent_phone__icontains=search)
        )

    defaulters = []
    total_due = Decimal('0.00')

    for s in students_qs:
        unpaid = s.invoices.filter(status__in=['Unpaid', 'Partially Paid', 'Overdue'])
        if active_session:
            unpaid = unpaid.filter(academic_session=active_session)
        due = sum((inv.remaining_due for inv in unpaid), Decimal('0.00'))
        if due > Decimal('0.00'):
            if min_due:
                try:
                    if due < Decimal(min_due):
                        continue
                except ValueError:
                    pass
            defaulters.append({
                'adm_no': s.admission_number,
                'name': s.full_name,
                'class_sec': f"{s.class_level.name} - Sec {s.section.name}",
                'parent_name': s.parent_name,
                'phone': s.parent_phone,
                'unpaid_count': unpaid.count(),
                'due': due
            })
            total_due += due

    defaulters.sort(key=lambda x: x['due'], reverse=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fee_Defaulters"

    ws.merge_cells("A1:G1")
    ws["A1"] = f"{school.name} — Pending Fee Defaulters Register"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="DC2626")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Generated On: {timezone.now().strftime('%d-%b-%Y %H:%M')}  |  Total Defaulters: {len(defaulters)}  |  Total Pending: Rs.{total_due:,.2f}"
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color="64748B")
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["Adm No", "Student Name", "Class & Section", "Parent / Guardian", "Contact Phone", "Pending Invoices", "Total Dues (Rs.)"]
    ws.append([])
    ws.append(headers)

    header_fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    row_font = Font(name="Arial", size=10)
    for d in defaulters:
        ws.append([
            d['adm_no'],
            d['name'],
            d['class_sec'],
            d['parent_name'],
            d['phone'],
            d['unpaid_count'],
            float(d['due'])
        ])
        curr_row = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=curr_row, column=col_idx)
            c.font = row_font
            c.border = thin_border
            if col_idx in [1, 3, 5, 6]:
                c.alignment = Alignment(horizontal="center")
            elif col_idx == 7:
                c.alignment = Alignment(horizontal="right")
                c.font = Font(name="Arial", size=10, bold=True, color="DC2626")

    # Column widths
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Fee_Defaulters_Report_{timezone.now().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# ==============================================================================
# PUBLIC DIGITAL SLIPS FOR PARENTS (ACCESSIBLE WITHOUT LOGIN VIA WHATSAPP LINKS)
# ==============================================================================

def public_fee_receipt_view(request, receipt_number):
    """
    Public digital fee receipt view accessible by parents via WhatsApp links without login.
    """
    ctx = get_school_context(request)
    school = ctx['school']
    
    receipt = FeeReceipt.objects.filter(receipt_number=receipt_number).select_related('student__class_level', 'student__section').prefetch_related('items').first()
    if not receipt and receipt_number.isdigit():
        receipt = FeeReceipt.objects.filter(pk=int(receipt_number)).select_related('student__class_level', 'student__section').prefetch_related('items').first()
        
    if not receipt:
        return render(request, 'school_erp/public_slip_not_found.html', {'slip_type': 'Fee Slip', 'number': receipt_number, 'school': school}, status=404)

    student = receipt.student
    amount_in_words = amount_to_words(receipt.amount_paid)
    public_slip_url = request.build_absolute_uri(f"/slip/receipt/{receipt.receipt_number}/")
    clean_phone = clean_phone_for_wa(student.parent_phone)

    # Pre-filled WhatsApp confirmation
    wa_items_list = []
    for it in receipt.items.all():
        period = f" ({it.month_name})" if it.month_name else ""
        wa_items_list.append(f"• {it.category_name}{period}: Rs.{it.amount:,.0f}")
    wa_items_text = "\n".join(wa_items_list)

    wa_message = (
        f"*{school.name.upper()}*\n"
        f"*Official Fee Payment Receipt*\n"
        f"-----------------------------------------\n"
        f"Receipt No: *{receipt.receipt_number}*\n"
        f"Date: *{receipt.payment_date.strftime('%d-%b-%Y')}*\n"
        f"Student: *{student.full_name}*\n"
        f"Adm No: *{student.admission_number}* | Class: *{student.class_level.name}-{student.section.name}*\n"
        f"Father: *{student.parent_name}*\n"
        f"-----------------------------------------\n"
        f"*Fee Details:*\n"
        f"{wa_items_text}\n"
        f"Total Billed: Rs.{receipt.total_billed:,.0f}\n"
        f"Amount Received: *Rs.{receipt.amount_paid:,.0f}*\n"
        f"Remaining Balance: *Rs.{receipt.remaining_balance_due:,.0f}*\n"
        f"Payment Mode: {receipt.payment_method}\n"
        f"-----------------------------------------\n"
        f"📄 *View / Download Official Slip Online (No Login Needed):*\n"
        f"{public_slip_url}\n"
        f"-----------------------------------------\n"
        f"Thank you for depositing the school fee.\n"
        f"Accounts Office, {school.name}"
    )

    ctx.update({
        'receipt': receipt,
        'student': student,
        'items': receipt.items.all(),
        'amount_in_words': amount_in_words,
        'clean_parent_phone': clean_phone,
        'public_slip_url': public_slip_url,
        'wa_message': wa_message,
        'is_public_view': True,
    })
    return render(request, 'school_erp/public_slip_receipt.html', ctx)


def public_payment_receipt_view(request, receipt_number):
    """
    Public digital single invoice payment receipt view without login.
    """
    ctx = get_school_context(request)
    school = ctx['school']

    payment = FeePayment.objects.filter(receipt_number=receipt_number).select_related('invoice__student__class_level', 'invoice__student__section', 'invoice__fee_category').first()
    if not payment and receipt_number.isdigit():
        payment = FeePayment.objects.filter(pk=int(receipt_number)).select_related('invoice__student__class_level', 'invoice__student__section', 'invoice__fee_category').first()

    if not payment:
        return render(request, 'school_erp/public_slip_not_found.html', {'slip_type': 'Payment Slip', 'number': receipt_number, 'school': school}, status=404)

    invoice = payment.invoice
    student = invoice.student
    amount_in_words = amount_to_words(payment.amount_paid)
    public_slip_url = request.build_absolute_uri(f"/slip/payment/{payment.receipt_number}/")
    clean_phone = clean_phone_for_wa(student.parent_phone)

    wa_message = (
        f"*{school.name.upper()}*\n"
        f"*Official Fee Payment Receipt*\n"
        f"-----------------------------------------\n"
        f"Receipt No: *{payment.receipt_number}*\n"
        f"Date: *{payment.payment_date.strftime('%d-%b-%Y')}*\n"
        f"Student: *{student.full_name}*\n"
        f"Adm No: *{student.admission_number}* | Class: *{student.class_level.name}-{student.section.name}*\n"
        f"Father: *{student.parent_name}*\n"
        f"-----------------------------------------\n"
        f"Category: {invoice.fee_category.name}\n"
        f"Amount Paid: *Rs.{payment.amount_paid:,.0f}*\n"
        f"Remaining Due: *Rs.{invoice.remaining_due:,.0f}*\n"
        f"Payment Mode: {payment.payment_method}\n"
        f"-----------------------------------------\n"
        f"📄 *View / Download Official Slip Online (No Login Needed):*\n"
        f"{public_slip_url}\n"
        f"-----------------------------------------\n"
        f"Thank you for depositing the school fee.\n"
        f"Accounts Office, {school.name}"
    )

    ctx.update({
        'payment': payment,
        'invoice': invoice,
        'student': student,
        'amount_in_words': amount_in_words,
        'clean_parent_phone': clean_phone,
        'public_slip_url': public_slip_url,
        'wa_message': wa_message,
        'is_public_view': True,
    })
    return render(request, 'school_erp/public_slip_payment.html', ctx)


def public_report_card_view(request, student_id, exam_id):
    """
    Public academic report card view accessible by parents via WhatsApp links without login.
    """
    ctx = get_school_context(request)
    school = ctx['school']
    student = get_object_or_404(Student.objects.select_related('class_level', 'section'), pk=student_id)
    exam = get_object_or_404(Exam, pk=exam_id)

    results = ExamResult.objects.filter(student=student, exam=exam).select_related('subject')
    total_obtained = sum(float(r.marks_obtained) for r in results)
    total_max = sum(float(r.max_marks) for r in results)
    overall_percentage = round((total_obtained / total_max * 100), 2) if total_max > 0 else 0

    if overall_percentage >= 90:
        overall_grade = 'A+'
        verdict = 'Outstanding Performance'
    elif overall_percentage >= 80:
        overall_grade = 'A'
        verdict = 'Excellent Performance'
    elif overall_percentage >= 70:
        overall_grade = 'B'
        verdict = 'Very Good'
    elif overall_percentage >= 60:
        overall_grade = 'C'
        verdict = 'Good / Satisfactory'
    elif overall_percentage >= 50:
        overall_grade = 'D'
        verdict = 'Needs Improvement'
    else:
        overall_grade = 'F'
        verdict = 'Failed / Re-appear'

    public_slip_url = request.build_absolute_uri(f"/slip/report-card/{student.id}/{exam.id}/")
    clean_phone = clean_phone_for_wa(student.parent_phone)

    wa_message = (
        f"*{school.name.upper()}*\n"
        f"*Official Academic Report Card / Marksheet*\n"
        f"-----------------------------------------\n"
        f"Student: *{student.full_name}*\n"
        f"Adm No: *{student.admission_number}* | Class: *{student.class_level.name}-{student.section.name}* (#{student.roll_number})\n"
        f"Exam: *{exam.name}*\n"
        f"Marks: *{total_obtained:,.1f} / {total_max:,.0f}* ({overall_percentage}%)\n"
        f"Grade: *{overall_grade}* &bull; Result: *{verdict}*\n"
        f"-----------------------------------------\n"
        f"📄 *View / Download Official Marksheet Online (No Login Needed):*\n"
        f"{public_slip_url}\n"
        f"-----------------------------------------\n"
        f"Examination Directorate, {school.name}"
    )

    ctx.update({
        'student': student,
        'exam': exam,
        'results': results,
        'total_obtained': total_obtained,
        'total_max': total_max,
        'overall_percentage': overall_percentage,
        'overall_grade': overall_grade,
        'verdict': verdict,
        'attendance_stats': student.attendance_stats,
        'clean_parent_phone': clean_phone,
        'public_slip_url': public_slip_url,
        'wa_message': wa_message,
        'is_public_view': True,
    })
    return render(request, 'school_erp/public_slip_report_card.html', ctx)


def public_id_card_view(request, student_id):
    """
    Public student identity card view accessible without login.
    """
    ctx = get_school_context(request)
    school = ctx['school']
    student = get_object_or_404(Student.objects.select_related('class_level', 'section'), pk=student_id)
    public_slip_url = request.build_absolute_uri(f"/slip/id-card/{student.id}/")
    clean_phone = clean_phone_for_wa(student.parent_phone)

    wa_message = (
        f"*{school.name.upper()}*\n"
        f"*Student Identity Card*\n"
        f"-----------------------------------------\n"
        f"Student: *{student.full_name}*\n"
        f"Adm No: *{student.admission_number}* | Class: *{student.class_level.name}-{student.section.name}*\n"
        f"Father: *{student.parent_name}*\n"
        f"-----------------------------------------\n"
        f"📄 *View / Print Official Student ID Card (No Login Needed):*\n"
        f"{public_slip_url}\n"
        f"-----------------------------------------\n"
        f"Administration Directorate, {school.name}"
    )

    ctx.update({
        'student': student,
        'clean_parent_phone': clean_phone,
        'public_slip_url': public_slip_url,
        'wa_message': wa_message,
        'is_public_view': True,
    })
    return render(request, 'school_erp/public_slip_id_card.html', ctx)


# ==============================================================================
# ROLE PERMISSIONS & AUDIT LOG VIEWS
# ==============================================================================

@require_menu_permission('can_manage_users', 'Role Master')
def roles_master_view(request):
    ctx = get_school_context(request)
    roles = Role.objects.annotate(user_count_annotated=Count('user_profiles')).order_by('id')
    form = RoleForm()

    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'create':
            form = RoleForm(request.POST)
            if form.is_valid():
                role = form.save()
                ensure_roles_and_permissions()
                log_audit(
                    request=request,
                    action='CREATE',
                    module='USERS',
                    object_type='Role',
                    object_id=str(role.id),
                    object_repr=role.name,
                    description=f"Created custom role '{role.name}' with code '{role.code}'",
                    reason="Role Creation"
                )
                messages.success(request, f"Role '{role.name}' has been created successfully!")
                return redirect('roles_master')
            else:
                messages.error(request, "Failed to create role. Please verify that the name and code are unique.")

        elif action == 'edit':
            role_id = request.POST.get('role_id')
            role = get_object_or_404(Role, pk=role_id)
            role.name = request.POST.get('name', role.name).strip()
            role.description = request.POST.get('description', role.description).strip()
            role.save()
            log_audit(
                request=request,
                action='UPDATE',
                module='USERS',
                object_type='Role',
                object_id=str(role.id),
                object_repr=role.name,
                description=f"Updated role details for '{role.name}'",
                reason="Role Update"
            )
            messages.success(request, f"Role '{role.name}' details updated successfully!")
            return redirect('roles_master')

        elif action == 'delete':
            role_id = request.POST.get('role_id')
            role = get_object_or_404(Role, pk=role_id)
            if role.is_system_role:
                messages.error(request, f"System role '{role.name}' is core and cannot be deleted.")
            elif role.user_profiles.exists():
                messages.error(request, f"Cannot delete role '{role.name}' because {role.user_profiles.count()} user(s) are assigned to it.")
            else:
                name = role.name
                role.permissions.all().delete()
                role.delete()
                log_audit(
                    request=request,
                    action='DELETE',
                    module='USERS',
                    object_type='Role',
                    object_id=str(role_id),
                    object_repr=name,
                    description=f"Deleted custom role '{name}'",
                    reason="Role Deletion"
                )
                messages.success(request, f"Role '{name}' has been deleted.")
            return redirect('roles_master')

    ctx.update({
        'roles': roles,
        'form': form,
    })
    return render(request, 'school_erp/roles_master.html', ctx)


@require_menu_permission('can_manage_users', 'Permission Matrix')
def permission_matrix_view(request, role_id=None):
    ctx = get_school_context(request)
    roles = Role.objects.all().order_by('id')
    
    if not role_id:
        default_role = Role.objects.filter(code='TEACHER').first() or roles.first()
        selected_role = default_role
    else:
        selected_role = get_object_or_404(Role, pk=role_id)

    ensure_roles_and_permissions()
    permissions = RolePermission.objects.filter(role=selected_role).order_by('id')

    if request.method == 'POST':
        changes = []
        for perm in permissions:
            old_vals = {
                'can_view': perm.can_view,
                'can_add': perm.can_add,
                'can_edit': perm.can_edit,
                'can_delete': perm.can_delete,
                'can_approve': perm.can_approve,
                'can_export': perm.can_export,
            }
            perm.can_view = f"can_view_{perm.module_code}" in request.POST
            perm.can_add = f"can_add_{perm.module_code}" in request.POST
            perm.can_edit = f"can_edit_{perm.module_code}" in request.POST
            perm.can_delete = f"can_delete_{perm.module_code}" in request.POST
            perm.can_approve = f"can_approve_{perm.module_code}" in request.POST
            perm.can_export = f"can_export_{perm.module_code}" in request.POST
            perm.save()

            new_vals = {
                'can_view': perm.can_view,
                'can_add': perm.can_add,
                'can_edit': perm.can_edit,
                'can_delete': perm.can_delete,
                'can_approve': perm.can_approve,
                'can_export': perm.can_export,
            }
            if old_vals != new_vals:
                changes.append({'module': perm.module_code, 'before': old_vals, 'after': new_vals})

        log_audit(
            request=request,
            action='PERMISSION_CHANGE',
            module='USERS',
            object_type='RolePermission',
            object_id=str(selected_role.id),
            object_repr=f"Permissions for {selected_role.name}",
            description=f"Updated granular module permissions for role '{selected_role.name}' ({len(changes)} modules modified)",
            changes=changes,
            reason=request.POST.get('reason', 'Security policy matrix update')
        )

        messages.success(request, f"Granular access permissions for '{selected_role.name}' saved and logged in audit trail!")
        return redirect('role_permission_matrix', role_id=selected_role.id)

    ctx.update({
        'roles': roles,
        'selected_role': selected_role,
        'permissions': permissions,
    })
    return render(request, 'school_erp/permission_matrix.html', ctx)


@require_menu_permission('can_manage_users', 'User Management')
def user_list_view(request):
    ctx = get_school_context(request)
    school = ctx['school']

    users_qs = User.objects.select_related('profile__role_obj', 'profile__school', 'profile__assigned_teacher').prefetch_related(
        'profile__assigned_classes', 'profile__assigned_sections', 'profile__linked_students'
    ).order_by('-date_joined')

    # Strict multi-tenant isolation: Only display users belonging to the current school
    if not is_superadmin_user(request.user):
        users_qs = users_qs.filter(profile__school=school)
    else:
        # SaaS Platform Superadmin can filter by school or see active switched school
        school_id_filter = request.GET.get('school_id')
        if school_id_filter:
            users_qs = users_qs.filter(profile__school_id=school_id_filter)
        elif request.session.get('saas_active_school_id'):
            users_qs = users_qs.filter(profile__school=school)

    role_filter = request.GET.get('role', '').strip()
    status_filter = request.GET.get('status', '').strip()
    search_query = request.GET.get('q', '').strip()

    if role_filter:
        users_qs = users_qs.filter(Q(profile__role_obj__code=role_filter) | Q(profile__role=role_filter))

    if status_filter == 'active':
        users_qs = users_qs.filter(is_active=True)
    elif status_filter == 'inactive':
        users_qs = users_qs.filter(is_active=False)

    if search_query:
        users_qs = users_qs.filter(
            Q(username__icontains=search_query) |
            Q(first_name__icontains=search_query) |
            Q(last_name__icontains=search_query) |
            Q(email__icontains=search_query) |
            Q(profile__phone__icontains=search_query)
        )

    roles = Role.objects.all().order_by('name')
    if not is_superadmin_user(request.user):
        roles = roles.exclude(code='SUPER_ADMIN')

    paginator = Paginator(users_qs, 20)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    ctx.update({
        'users': page_obj,
        'page_obj': page_obj,
        'paginator': paginator,
        'roles': roles,
        'selected_role': role_filter,
        'selected_status': status_filter,
        'search_query': search_query,
        'total_users': users_qs.count(),
    })
    return render(request, 'school_erp/user_list.html', ctx)


@require_menu_permission('can_manage_users', 'Create User')
def user_create_view(request):
    ctx = get_school_context(request)
    school = ctx['school']
    is_superadmin = is_superadmin_user(request.user)

    if request.method == 'POST':
        form = UserCreateUpdateForm(request.POST, school=school, is_superadmin=is_superadmin)
        if form.is_valid():
            cd = form.cleaned_data
            if User.objects.filter(username=cd['username']).exists():
                messages.error(request, f"Username '{cd['username']}' is already in use. Please select a different username.")
            else:
                user = User.objects.create_user(
                    username=cd['username'],
                    email=cd.get('email', ''),
                    password=cd.get('password') or 'School@123',
                    first_name=cd.get('first_name', ''),
                    last_name=cd.get('last_name', ''),
                    is_active=cd.get('is_active', True)
                )

                role_obj = cd.get('role_obj')
                profile, _ = UserProfile.objects.get_or_create(user=user)
                profile.school = school
                profile.role_obj = role_obj
                profile.role = role_obj.name if role_obj else 'Teacher'
                profile.phone = cd.get('phone', '')
                profile.assigned_teacher = cd.get('assigned_teacher')
                profile.session_restriction = cd.get('session_restriction')
                profile.save()

                if cd.get('assigned_classes'):
                    profile.assigned_classes.set(cd['assigned_classes'])
                if cd.get('assigned_sections'):
                    profile.assigned_sections.set(cd['assigned_sections'])
                if cd.get('linked_students'):
                    profile.linked_students.set(cd['linked_students'])

                log_audit(
                    request=request,
                    action='CREATE',
                    module='USERS',
                    object_type='User',
                    object_id=str(user.id),
                    object_repr=f"{user.username} ({profile.role})",
                    description=f"Created user account '{user.username}' with role '{profile.role}' for school '{school.name}'",
                    reason="New staff/guardian user onboarding"
                )

                messages.success(request, f"User account '{user.username}' ({profile.role}) created successfully!")
                return redirect('user_list')
    else:
        form = UserCreateUpdateForm(school=school, is_superadmin=is_superadmin)

    ctx.update({
        'form': form,
        'title': 'Create New User Account',
        'is_edit': False,
    })
    return render(request, 'school_erp/user_form.html', ctx)


@require_menu_permission('can_manage_users', 'Edit User')
def user_edit_view(request, user_id):
    ctx = get_school_context(request)
    school = ctx['school']
    is_superadmin = is_superadmin_user(request.user)

    if not is_superadmin:
        target_user = get_object_or_404(User, pk=user_id, profile__school=school)
    else:
        target_user = get_object_or_404(User, pk=user_id)
    profile, _ = UserProfile.objects.get_or_create(user=target_user)

    if request.method == 'POST':
        form = UserCreateUpdateForm(request.POST, instance_user=target_user, school=school, is_superadmin=is_superadmin)
        if form.is_valid():
            cd = form.cleaned_data
            if cd['username'] != target_user.username and User.objects.filter(username=cd['username']).exclude(pk=target_user.pk).exists():
                messages.error(request, f"Username '{cd['username']}' is already in use.")
            else:
                target_user.username = cd['username']
                target_user.first_name = cd.get('first_name', '')
                target_user.last_name = cd.get('last_name', '')
                target_user.email = cd.get('email', '')
                target_user.is_active = cd.get('is_active', True)
                if cd.get('password'):
                    target_user.set_password(cd['password'])
                target_user.save()

                role_obj = cd.get('role_obj')
                profile.role_obj = role_obj
                if role_obj:
                    profile.role = role_obj.name
                profile.phone = cd.get('phone', '')
                profile.assigned_teacher = cd.get('assigned_teacher')
                profile.session_restriction = cd.get('session_restriction')
                profile.save()

                profile.assigned_classes.set(cd.get('assigned_classes') or [])
                profile.assigned_sections.set(cd.get('assigned_sections') or [])
                profile.linked_students.set(cd.get('linked_students') or [])

                log_audit(
                    request=request,
                    action='UPDATE',
                    module='USERS',
                    object_type='User',
                    object_id=str(target_user.id),
                    object_repr=f"{target_user.username} ({profile.role})",
                    description=f"Updated profile, role permissions & restrictions for user '{target_user.username}'",
                    reason="User record update"
                )

                messages.success(request, f"User account '{target_user.username}' updated successfully!")
                return redirect('user_list')
    else:
        form = UserCreateUpdateForm(instance_user=target_user, school=school, is_superadmin=is_superadmin)

    ctx.update({
        'form': form,
        'target_user': target_user,
        'title': f"Edit User: {target_user.username}",
        'is_edit': True,
    })
    return render(request, 'school_erp/user_form.html', ctx)


@require_menu_permission('can_manage_users', 'Toggle User Status')
def user_toggle_status_view(request, user_id):
    ctx = get_school_context(request)
    school = ctx['school']
    if not is_superadmin_user(request.user):
        target_user = get_object_or_404(User, pk=user_id, profile__school=school)
    else:
        target_user = get_object_or_404(User, pk=user_id)

    if target_user == request.user:
        messages.error(request, "You cannot deactivate your own logged-in administrator account.")
        return redirect('user_list')

    target_user.is_active = not target_user.is_active
    target_user.save()
    status_str = "activated" if target_user.is_active else "deactivated"

    log_audit(
        request=request,
        action='UPDATE',
        module='USERS',
        object_type='User',
        object_id=str(target_user.id),
        object_repr=target_user.username,
        description=f"User account '{target_user.username}' {status_str}",
        reason=f"Account status toggle to {status_str}"
    )

    messages.success(request, f"User '{target_user.username}' has been {status_str}.")
    return redirect('user_list')


@require_menu_permission('can_manage_users', 'Audit Log Viewer')
def audit_logs_view(request):
    ctx = get_school_context(request)
    school = ctx['school']

    logs_qs = AuditLog.objects.select_related('user', 'school', 'approved_by').all().order_by('-timestamp')

    # Strict multi-tenant isolation: Only display audit logs for the current school
    if not is_superadmin_user(request.user):
        logs_qs = logs_qs.filter(school=school)
    else:
        school_id_filter = request.GET.get('school_id')
        if school_id_filter:
            logs_qs = logs_qs.filter(school_id=school_id_filter)
        elif request.session.get('saas_active_school_id'):
            logs_qs = logs_qs.filter(school=school)

    module_filter = request.GET.get('module', '').strip()
    action_filter = request.GET.get('action', '').strip()
    status_filter = request.GET.get('status', '').strip()
    search_query = request.GET.get('q', '').strip()
    tab = request.GET.get('tab', 'all')

    if tab == 'pending':
        logs_qs = logs_qs.filter(requires_approval=True, status='PENDING_APPROVAL')
    else:
        if status_filter:
            logs_qs = logs_qs.filter(status=status_filter)

    if module_filter:
        logs_qs = logs_qs.filter(module=module_filter)
    if action_filter:
        logs_qs = logs_qs.filter(action=action_filter)
    if search_query:
        logs_qs = logs_qs.filter(
            Q(description__icontains=search_query) |
            Q(object_repr__icontains=search_query) |
            Q(reason__icontains=search_query) |
            Q(user__username__icontains=search_query) |
            Q(user_role__icontains=search_query)
        )

    # Statistics
    today = timezone.now().date()
    stats = {
        'total_logs': AuditLog.objects.count(),
        'today_actions': AuditLog.objects.filter(timestamp__date=today).count(),
        'sensitive_actions': AuditLog.objects.filter(Q(action__in=['DELETE', 'APPROVE', 'REJECT', 'PUBLISH', 'PERMISSION_CHANGE']) | Q(requires_approval=True)).count(),
        'pending_approvals': AuditLog.objects.filter(requires_approval=True, status='PENDING_APPROVAL').count(),
    }

    paginator = Paginator(logs_qs, 25)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    ctx.update({
        'logs': page_obj,
        'page_obj': page_obj,
        'paginator': paginator,
        'stats': stats,
        'selected_module': module_filter,
        'selected_action': action_filter,
        'selected_status': status_filter,
        'search_query': search_query,
        'active_tab': tab,
        'module_choices': AuditLog.MODULE_CHOICES,
        'action_choices': AuditLog.ACTION_CHOICES,
        'status_choices': AuditLog.STATUS_CHOICES,
    })
    return render(request, 'school_erp/audit_logs.html', ctx)


@require_menu_permission('can_manage_users', 'Audit Log Approval Action')
def audit_log_action_view(request, log_id, action):
    audit_log = get_object_or_404(AuditLog, pk=log_id)
    if request.method == 'POST':
        remarks = request.POST.get('remarks', '').strip()
        if action == 'approve':
            audit_log.status = 'APPROVED'
            audit_log.approved_by = request.user
            audit_log.approved_at = timezone.now()
            audit_log.approval_remarks = remarks or 'Approved by administrator'
            audit_log.save()

            log_audit(
                request=request,
                action='APPROVE',
                module=audit_log.module,
                object_type='AuditLog',
                object_id=str(audit_log.id),
                object_repr=f"Approval of #{audit_log.id}",
                description=f"Sensitive operation #{audit_log.id} ({audit_log.object_repr}) approved by {request.user.username}",
                reason=audit_log.approval_remarks
            )
            messages.success(request, f"Operation #{audit_log.id} ({audit_log.object_repr}) has been APPROVED successfully!")

        elif action == 'reject':
            audit_log.status = 'REJECTED'
            audit_log.approved_by = request.user
            audit_log.approved_at = timezone.now()
            audit_log.approval_remarks = remarks or 'Rejected by administrator'
            audit_log.save()

            log_audit(
                request=request,
                action='REJECT',
                module=audit_log.module,
                object_type='AuditLog',
                object_id=str(audit_log.id),
                object_repr=f"Rejection of #{audit_log.id}",
                description=f"Sensitive operation #{audit_log.id} ({audit_log.object_repr}) rejected by {request.user.username}",
                reason=audit_log.approval_remarks
            )
            messages.warning(request, f"Operation #{audit_log.id} ({audit_log.object_repr}) has been REJECTED.")

    return redirect('audit_logs')


@require_menu_permission('can_manage_users', 'Export Audit Logs')
def export_audit_logs_excel(request):
    ctx = get_school_context(request)
    school = ctx['school']
    logs_qs = AuditLog.objects.select_related('user', 'school', 'approved_by').all().order_by('-timestamp')

    if not is_superadmin_user(request.user):
        logs_qs = logs_qs.filter(school=school)
    else:
        school_id_filter = request.GET.get('school_id')
        if school_id_filter:
            logs_qs = logs_qs.filter(school_id=school_id_filter)
        elif request.session.get('saas_active_school_id'):
            logs_qs = logs_qs.filter(school=school)

    module_filter = request.GET.get('module', '').strip()
    action_filter = request.GET.get('action', '').strip()
    status_filter = request.GET.get('status', '').strip()
    search_query = request.GET.get('q', '').strip()

    if module_filter:
        logs_qs = logs_qs.filter(module=module_filter)
    if action_filter:
        logs_qs = logs_qs.filter(action=action_filter)
    if status_filter:
        logs_qs = logs_qs.filter(status=status_filter)
    if search_query:
        logs_qs = logs_qs.filter(
            Q(description__icontains=search_query) |
            Q(object_repr__icontains=search_query) |
            Q(reason__icontains=search_query) |
            Q(user__username__icontains=search_query)
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Audit Trail Logs"

    headers = ["ID", "Timestamp (IST)", "User", "Role", "Action", "Module", "Target Object", "Description", "Reason / Justification", "Status", "Approved By", "IP Address"]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for log in logs_qs[:2000]:
        approved_by_str = log.approved_by.username if log.approved_by else '-'
        ws.append([
            log.id,
            log.timestamp.strftime("%Y-%m-%d %H:%M:%S") if log.timestamp else '',
            log.user.username if log.user else 'System',
            log.user_role or '-',
            log.get_action_display() if hasattr(log, 'get_action_display') else log.action,
            log.get_module_display() if hasattr(log, 'get_module_display') else log.module,
            log.object_repr or log.object_type,
            log.description,
            log.reason or '-',
            log.status,
            approved_by_str,
            log.ip_address or '-'
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="school_audit_logs_{timezone.now().strftime("%Y%m%d_%H%M")}.xlsx"'
    wb.save(response)

    log_audit(
        request=request,
        action='EXPORT',
        module='USERS',
        object_type='AuditLog',
        description=f"Exported {logs_qs.count()} audit trail records to Excel",
        reason="Security compliance & audit export"
    )
    return response


# ==============================================================================
# SAAS PLATFORM SUPER ADMIN: BACKUP & RESTORE CONSOLE
# ==============================================================================

@require_superadmin
def backup_config_view(request):
    """SaaS Admin: Backup Configuration screen: scheduled backups, encryption, retention."""
    schools = SchoolProfile.objects.all().order_by('name')
    selected_school_id = request.GET.get('school_id', '') or request.POST.get('school_id', '')
    if selected_school_id:
        school = SchoolProfile.objects.filter(id=selected_school_id).first() or schools.first()
    else:
        ctx = get_school_context(request)
        school = ctx.get('school') or schools.first()

    config, _ = BackupConfiguration.objects.get_or_create(school=school)

    if request.method == 'POST':
        form = BackupConfigurationForm(request.POST, instance=config)
        if form.is_valid():
            form.save()
            log_audit(
                request=request,
                action='UPDATE',
                module='USERS',
                object_type='BackupConfiguration',
                object_id=str(config.id),
                object_repr=f"Backup policy for {school.name}",
                description=f"Updated backup schedule ({config.backup_frequency}), retention ({config.retention_days} days) for {school.name}",
                reason="SaaS infrastructure policy update"
            )
            messages.success(request, f"Backup configuration for '{school.name}' saved successfully.")
            return redirect(f"/backup/config/?school_id={school.id}")
    else:
        form = BackupConfigurationForm(instance=config)

    ctx = {
        'config': config,
        'form': form,
        'school': school,
        'schools': schools,
        'selected_school_id': str(school.id) if school else '',
        'title': 'SaaS Backup Policy Configuration'
    }
    return render(request, 'school_erp/backup_config.html', ctx)


@require_superadmin
def backup_history_view(request):
    """SaaS Admin: Backup History screen showing all generated archives with checksums & verification status."""
    schools = SchoolProfile.objects.all().order_by('name')
    selected_school_id = request.GET.get('school_id', '').strip()

    backups = BackupLog.objects.select_related('school', 'created_by').all().order_by('-created_at')
    if selected_school_id:
        try:
            backups = backups.filter(school_id=int(selected_school_id))
        except (ValueError, TypeError):
            pass

    total_backups = backups.count()
    verified_backups = backups.filter(is_verified=True).count()
    total_size = sum(b.file_size_mb for b in backups)
    last_backup = backups.first()

    ctx = {
        'backups': backups,
        'schools': schools,
        'selected_school_id': selected_school_id,
        'total_backups': total_backups,
        'verified_backups': verified_backups,
        'total_size': total_size,
        'last_backup': last_backup,
        'title': 'SaaS Backup History & Archives'
    }
    return render(request, 'school_erp/backup_history.html', ctx)


@require_superadmin
def backup_run_view(request):
    """SaaS Admin: Trigger manual backup generation on demand."""
    if request.method == 'POST':
        school_id = request.POST.get('school_id', '').strip()
        if school_id:
            school = SchoolProfile.objects.filter(id=school_id).first()
        else:
            ctx = get_school_context(request)
            school = ctx.get('school') or SchoolProfile.objects.first()

        if not school:
            messages.error(request, "No school selected to generate backup.")
            return redirect('backup_history')

        backup_type = request.POST.get('backup_type', 'Full')
        try:
            backup_log = generate_system_backup(
                school=school,
                backup_type=backup_type,
                created_by=request.user
            )
            messages.success(
                request,
                f"Backup '{backup_log.backup_name}' for '{school.name}' generated successfully! "
                f"Size: {backup_log.file_size_mb} MB | SHA-256: {backup_log.checksum_sha256[:16]}..."
            )
        except Exception as e:
            messages.error(request, f"Error generating backup: {str(e)}")

    redirect_url = 'backup_history'
    if request.POST.get('school_id'):
        redirect_url += f"?school_id={request.POST.get('school_id')}"
    return redirect(redirect_url)


@require_superadmin
def backup_download_view(request, backup_id):
    """SaaS Admin: Secure download of generated backup archive."""
    backup = get_object_or_404(BackupLog, pk=backup_id)

    if not os.path.exists(backup.file_path):
        messages.error(request, "Backup file does not exist on disk.")
        return redirect('backup_history')

    log_audit(
        request=request,
        action='EXPORT',
        module='USERS',
        object_type='BackupLog',
        object_id=str(backup.id),
        object_repr=backup.backup_name,
        description=f"Downloaded backup archive {backup.backup_name} for {backup.school.name}",
        reason="Disaster recovery archive retrieval"
    )

    response = FileResponse(open(backup.file_path, 'rb'), content_type='application/zip')
    response['Content-Disposition'] = f'attachment; filename="{backup.backup_name}"'
    return response


@require_superadmin
def backup_restore_view(request):
    """SaaS Admin: Restore Console: isolated verification testing & staging environment restores."""
    schools = SchoolProfile.objects.all().order_by('name')
    selected_school_id = request.GET.get('school_id', '').strip()

    backups = BackupLog.objects.select_related('school').all().order_by('-created_at')
    if selected_school_id:
        try:
            backups = backups.filter(school_id=int(selected_school_id))
        except (ValueError, TypeError):
            pass

    ctx = {
        'backups': backups,
        'schools': schools,
        'selected_school_id': selected_school_id,
        'title': 'SaaS Restore Console & Sandbox Verification'
    }
    return render(request, 'school_erp/backup_restore.html', ctx)


@require_superadmin
def backup_restore_test_view(request, backup_id):
    """SaaS Admin: Execute isolated sandbox verification or staging test restore."""
    backup = get_object_or_404(BackupLog, pk=backup_id)

    action = request.POST.get('test_action', 'verify')
    if action == 'staging':
        result = execute_test_environment_restore(backup.id, user=request.user)
        if result['success']:
            messages.success(request, f"Staging Restore Successful! {result['message']}")
        else:
            messages.error(request, f"Staging Restore Failed: {result['message']}")
    else:
        result = verify_restore_archive(backup.id, user=request.user)
        if result['success']:
            messages.success(
                request,
                f"Isolated Verification Passed! Database Tables Verified: {result.get('verified_tables', 0)}, "
                f"Students Verified: {result.get('student_records', 0)}. Production data is completely intact."
            )
        else:
            messages.error(request, f"Verification Failed: {result['message']}")
    return redirect('backup_restore')


# ==============================================================================
# SECURITY SETTINGS & AUDIT COMPLIANCE
# ==============================================================================

@require_menu_permission('can_manage_users', 'Security Settings')
def security_settings_view(request):
    """Security Settings, Environment Mode Switcher, OWASP/GIGW Compliance & Login Failure Monitor."""
    ctx = get_school_context(request)
    school = ctx['school']
    sec_config, _ = SecurityConfiguration.objects.get_or_create(school=school)

    if request.method == 'POST':
        form = SecurityConfigurationForm(request.POST, instance=sec_config)
        if form.is_valid():
            sec_config.last_security_audit_at = timezone.now()
            form.save()
            AuditLog.log(
                school=school,
                user=request.user,
                module='Backup & Security',
                action='Update Security Policy',
                target=f"Environment: {sec_config.environment_mode}",
                details=f"Session Expiry: {sec_config.session_expiry_minutes}m, MFA: {sec_config.enable_mfa}, Lockout: {sec_config.max_login_failures} attempts"
            )
            messages.success(request, f"Security policy updated for {sec_config.environment_mode} environment.")
            return redirect('security_settings')
    else:
        form = SecurityConfigurationForm(instance=sec_config)

    # Filter Login Attempt Logs
    status_filter = request.GET.get('status', '').strip()
    login_logs = LoginAttemptLog.objects.filter(school=school)
    if status_filter:
        login_logs = login_logs.filter(status=status_filter)
    login_logs = login_logs[:60]

    # Compliance Scorecard (Government of India GIGW & OWASP Top 10)
    compliance_checks = [
        {"title": "HTTPS & Transport Security", "desc": "TLS 1.3 / HSTS enforced with secure cookie flags", "passed": sec_config.enforce_https},
        {"title": "Brute Force Protection", "desc": "Account lockout after threshold failures with IP cooldown", "passed": sec_config.max_login_failures <= 5},
        {"title": "Session Inactivity Timeout", "desc": "Session expires automatically after idle duration (<= 60 mins)", "passed": sec_config.session_expiry_minutes <= 60},
        {"title": "Password Complexity Policy", "desc": "Requires uppercase, digits, and special characters (min 8 chars)", "passed": (sec_config.password_min_length >= 8 and sec_config.password_require_uppercase and sec_config.password_require_digits)},
        {"title": "Multi-Factor Authentication (MFA)", "desc": "Two-factor authentication pin verification on login", "passed": sec_config.enable_mfa},
        {"title": "Suspicious Activity Monitoring", "desc": "Real-time anomaly detection and IP lockout logging", "passed": sec_config.enable_suspicious_activity_alerts},
        {"title": "Data Loss & Encrypted Backups", "desc": "Scheduled AES-encrypted backups with SHA-256 integrity hash", "passed": True},
        {"title": "Role-Based Access Control (RBAC)", "desc": "Granular permissions preventing cross-account access", "passed": True},
    ]
    passed_count = sum(1 for c in compliance_checks if c["passed"])
    compliance_score = int((passed_count / len(compliance_checks)) * 100)

    ctx.update({
        'sec_config': sec_config,
        'form': form,
        'login_logs': login_logs,
        'compliance_checks': compliance_checks,
        'compliance_score': compliance_score,
        'status_filter': status_filter,
        'title': 'Security Settings & Compliance'
    })
    return render(request, 'school_erp/security_settings.html', ctx)


# ==============================================================================
# MULTI-BRANCH MASTER & SCHOOL SETUP
# ==============================================================================

@require_menu_permission('can_manage_classes', 'Branch Master')
def branch_master_view(request):
    """Branch Master screen: view, manage, and configure school branches."""
    ctx = get_school_context(request)
    school = ctx['school']
    branches = Branch.objects.filter(school=school).annotate(
        student_count=Count('students', distinct=True),
        class_count=Count('class_levels', distinct=True)
    )

    total_branches = branches.count()
    active_branches = branches.filter(is_active=True).count()

    ctx.update({
        'branches': branches,
        'total_branches': total_branches,
        'active_branches': active_branches,
        'title': 'Branch Master'
    })
    return render(request, 'school_erp/branch_master.html', ctx)


@require_menu_permission('can_manage_classes', 'Create Branch')
def branch_create_view(request):
    """Add a new school branch with custom numbering prefixes & branding."""
    ctx = get_school_context(request)
    school = ctx['school']

    if request.method == 'POST':
        form = BranchForm(request.POST, request.FILES)
        if form.is_valid():
            branch = form.save(commit=False)
            branch.school = school
            if branch.is_main_branch:
                Branch.objects.filter(school=school, is_main_branch=True).update(is_main_branch=False)
            branch.save()

            AuditLog.log(
                school=school,
                user=request.user,
                module='School Setup',
                action='Create Branch',
                target=f"{branch.name} ({branch.code})",
                details=f"Receipt Prefix: {branch.receipt_prefix}, Student Prefix: {branch.student_id_prefix}"
            )
            messages.success(request, f"Branch '{branch.name}' created successfully.")
            return redirect('branch_master')
    else:
        form = BranchForm()

    ctx.update({
        'form': form,
        'title': 'Add New Branch'
    })
    return render(request, 'school_erp/branch_form.html', ctx)


@require_menu_permission('can_manage_classes', 'Edit Branch')
def branch_edit_view(request, branch_id):
    """Edit branch details, prefix, and settings."""
    ctx = get_school_context(request)
    school = ctx['school']
    branch = get_object_or_404(Branch, pk=branch_id, school=school)

    if request.method == 'POST':
        form = BranchForm(request.POST, request.FILES, instance=branch)
        if form.is_valid():
            b = form.save(commit=False)
            if b.is_main_branch:
                Branch.objects.filter(school=school, is_main_branch=True).exclude(pk=b.pk).update(is_main_branch=False)
            b.save()

            AuditLog.log(
                school=school,
                user=request.user,
                module='School Setup',
                action='Update Branch',
                target=f"{branch.name} ({branch.code})",
                details="Updated branch information and settings"
            )
            messages.success(request, f"Branch '{branch.name}' updated successfully.")
            return redirect('branch_master')
    else:
        form = BranchForm(instance=branch)

    ctx.update({
        'form': form,
        'branch': branch,
        'title': f'Edit Branch: {branch.name}'
    })
    return render(request, 'school_erp/branch_form.html', ctx)


@require_menu_permission('can_manage_classes', 'Toggle Branch')
def branch_toggle_view(request, branch_id):
    """Activate or deactivate a branch."""
    if request.method == 'POST':
        ctx = get_school_context(request)
        school = ctx['school']
        branch = get_object_or_404(Branch, pk=branch_id, school=school)
        branch.is_active = not branch.is_active
        branch.save()

        AuditLog.log(
            school=school,
            user=request.user,
            module='School Setup',
            action='Toggle Branch Status',
            target=branch.name,
            details=f"Status set to {'Active' if branch.is_active else 'Inactive'}"
        )
        messages.info(request, f"Branch '{branch.name}' status set to {'Active' if branch.is_active else 'Inactive'}.")
    return redirect('branch_master')


# ==============================================================================
# PROMOTION WIZARD (REPEAT-YEAR, SECTION-CHANGE, SESSION PRESERVATION)
# ==============================================================================

@require_menu_permission('can_manage_students', 'Promotion Wizard')
def promotion_wizard_view(request):
    """
    Comprehensive Promotion Wizard:
    - Supports Promote, Section Change, and Repeat Year (Retain) workflows.
    - Guarantees that historical marks, receipts, and attendance are NOT overwritten.
    - Updates/creates StudentSessionEnrollment records preserving past academic session history.
    """
    ctx = get_school_context(request)
    school = ctx['school']

    sessions = AcademicSession.objects.filter(school=school).order_by('-start_date')
    classes = ClassLevel.objects.filter(school=school).order_by('order', 'name')
    branches = Branch.objects.filter(school=school, is_active=True).order_by('name')

    source_session_id = request.GET.get('source_session') or request.POST.get('source_session')
    source_class_id = request.GET.get('source_class') or request.POST.get('source_class')
    source_section_id = request.GET.get('source_section') or request.POST.get('source_section')

    source_session = None
    source_class = None
    source_section = None
    students = []

    if source_session_id:
        source_session = AcademicSession.objects.filter(id=source_session_id, school=school).first()
    if not source_session and sessions.exists():
        source_session = sessions.first()
        source_session_id = str(source_session.id)

    if source_class_id:
        source_class = ClassLevel.objects.filter(id=source_class_id, school=school).first()
    if not source_class and classes.exists():
        source_class = classes.first()
        source_class_id = str(source_class.id)

    if source_section_id:
        source_section = Section.objects.filter(id=source_section_id, class_level=source_class).first()

    # Load students matching source session and class
    if source_session and source_class:
        students_qs = Student.objects.filter(
            school=school,
            is_active=True,
            academic_session=source_session,
            class_level=source_class
        )
        if source_section:
            students_qs = students_qs.filter(section=source_section)
        students = students_qs.select_related('class_level', 'section', 'academic_session').order_by('roll_number', 'first_name')

    # Available sections for source class
    source_sections = Section.objects.filter(class_level=source_class).order_by('name') if source_class else []

    # Handle Promotion / Retention / Section Change Submission
    if request.method == 'POST' and request.POST.get('action') == 'process_promotion':
        target_session_id = request.POST.get('target_session')
        target_class_id = request.POST.get('target_class')
        target_section_id = request.POST.get('target_section')
        selected_student_ids = request.POST.getlist('selected_students')

        target_session = AcademicSession.objects.filter(id=target_session_id, school=school).first()
        target_class = ClassLevel.objects.filter(id=target_class_id, school=school).first()
        target_section = Section.objects.filter(id=target_section_id, class_level=target_class).first()

        if not selected_student_ids:
            messages.warning(request, "No students were selected for processing.")
            return redirect(f"{request.path}?source_session={source_session_id}&source_class={source_class_id}")

        processed_count = 0
        promoted_count = 0
        retained_count = 0
        section_changed_count = 0

        for student_id in selected_student_ids:
            student = Student.objects.filter(id=student_id, school=school).first()
            if not student:
                continue

            student_action = request.POST.get(f'student_action_{student.id}', 'Promote')

            if student_action == 'Promote':
                if not target_session or not target_class or not target_section:
                    messages.error(request, "Target Session, Target Class, and Target Section are required for Promotion.")
                    return redirect(f"{request.path}?source_session={source_session_id}&source_class={source_class_id}")

                # 1. Update source session enrollment as 'Promoted' (preserving past history)
                StudentSessionEnrollment.objects.update_or_create(
                    student=student,
                    academic_session=source_session,
                    defaults={
                        'class_level': student.class_level,
                        'section': student.section,
                        'roll_number': student.roll_number,
                        'status': 'Promoted'
                    }
                )

                # 2. Create target session enrollment as 'Active'
                StudentSessionEnrollment.objects.update_or_create(
                    student=student,
                    academic_session=target_session,
                    defaults={
                        'class_level': target_class,
                        'section': target_section,
                        'roll_number': student.roll_number,
                        'status': 'Active',
                        'promoted_from_session': source_session
                    }
                )

                # 3. Update current student active pointer
                student.academic_session = target_session
                student.class_level = target_class
                student.section = target_section
                student.save(update_fields=['academic_session', 'class_level', 'section'])

                promoted_count += 1
                processed_count += 1

            elif student_action == 'Repeat_Year':
                if not target_session:
                    messages.error(request, "Target Session is required for Repeat Year.")
                    return redirect(f"{request.path}?source_session={source_session_id}&source_class={source_class_id}")

                # 1. Mark source session enrollment as 'Retained'
                StudentSessionEnrollment.objects.update_or_create(
                    student=student,
                    academic_session=source_session,
                    defaults={
                        'class_level': student.class_level,
                        'section': student.section,
                        'roll_number': student.roll_number,
                        'status': 'Retained'
                    }
                )

                # 2. Create target session enrollment in SAME class as 'Active'
                StudentSessionEnrollment.objects.update_or_create(
                    student=student,
                    academic_session=target_session,
                    defaults={
                        'class_level': student.class_level,
                        'section': student.section,
                        'roll_number': student.roll_number,
                        'status': 'Active',
                        'promoted_from_session': source_session
                    }
                )

                # 3. Update student pointer to new session (same class)
                student.academic_session = target_session
                student.save(update_fields=['academic_session'])

                retained_count += 1
                processed_count += 1

            elif student_action == 'Section_Change':
                new_sec_id = request.POST.get(f'new_section_{student.id}')
                new_sec = Section.objects.filter(id=new_sec_id, class_level=student.class_level).first()
                if new_sec:
                    student.section = new_sec
                    student.save(update_fields=['section'])
                    # Update active session enrollment
                    StudentSessionEnrollment.objects.filter(
                        student=student,
                        academic_session=source_session
                    ).update(section=new_sec)
                    section_changed_count += 1
                    processed_count += 1

        AuditLog.log(
            school=school,
            user=request.user,
            module='Promotion Wizard',
            action='Batch Student Promotion',
            target=f"From {source_session.name if source_session else '-'} {source_class.name if source_class else '-'}",
            details=f"Processed {processed_count} students: {promoted_count} Promoted, {retained_count} Repeat-Year, {section_changed_count} Section Change. Historical marks and receipts preserved."
        )

        messages.success(
            request,
            f"Promotion Wizard completed successfully! Processed {processed_count} student records "
            f"({promoted_count} Promoted, {retained_count} Repeat-Year, {section_changed_count} Section Changes). "
            f"Historical marks, receipts, and attendance remain fully preserved."
        )
        return redirect(f"{request.path}?source_session={source_session_id}&source_class={source_class_id}")

    ctx.update({
        'sessions': sessions,
        'classes': classes,
        'branches': branches,
        'source_session': source_session,
        'source_class': source_class,
        'source_section': source_section,
        'source_sections': source_sections,
        'students': students,
        'source_session_id': source_session_id,
        'source_class_id': source_class_id,
        'source_section_id': source_section_id,
        'title': 'Promotion Wizard'
    })
    return render(request, 'school_erp/promotion_wizard.html', ctx)


# ==============================================================================
# CONSOLIDATED MANAGEMENT REPORT
# ==============================================================================

@require_menu_permission('can_manage_fees', 'Consolidated Management Report')
def consolidated_report_view(request):
    """Cross-branch & cross-session executive summary report."""
    ctx = get_school_context(request)
    school = ctx['school']

    branches = Branch.objects.filter(school=school)
    sessions = AcademicSession.objects.filter(school=school).order_by('-start_date')

    # Branch-wise metrics
    branch_metrics = []
    total_students_all = 0
    total_collections_all = Decimal('0.00')

    for b in branches:
        student_cnt = Student.objects.filter(school=school, branch=b, is_active=True).count()
        receipt_total = FeeReceipt.objects.filter(
            school=school, branch=b, status='Active'
        ).aggregate(total=Sum('amount_paid'))['total'] or Decimal('0.00')

        total_students_all += student_cnt
        total_collections_all += receipt_total

        branch_metrics.append({
            'branch': b,
            'students_count': student_cnt,
            'receipts_total': receipt_total,
            'is_main': b.is_main_branch,
            'branch_head': b.branch_head or '-'
        })

    # Unassigned branch count
    unassigned_students = Student.objects.filter(school=school, branch__isnull=True, is_active=True).count()
    unassigned_collections = FeeReceipt.objects.filter(
        school=school, branch__isnull=True, status='Active'
    ).aggregate(total=Sum('amount_paid'))['total'] or Decimal('0.00')

    if unassigned_students > 0 or unassigned_collections > 0 or not branches.exists():
        branch_metrics.append({
            'branch': {'name': 'Main / Head Office Campus', 'code': 'HQ-MAIN', 'is_main_branch': True},
            'students_count': unassigned_students,
            'receipts_total': unassigned_collections,
            'is_main': True,
            'branch_head': school.principal_name or '-'
        })
        total_students_all += unassigned_students
        total_collections_all += unassigned_collections

    # Session-wise comparison
    session_metrics = []
    for s in sessions[:5]:
        s_students = Student.objects.filter(school=school, academic_session=s).count()
        s_invoices = FeeInvoice.objects.filter(school=school, academic_session=s).aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
        s_paid = FeeReceipt.objects.filter(school=school, student__academic_session=s, status='Active').aggregate(total=Sum('amount_paid'))['total'] or Decimal('0.00')
        s_due = s_invoices - s_paid
        session_metrics.append({
            'session': s,
            'student_count': s_students,
            'invoiced_total': s_invoices,
            'collected_total': s_paid,
            'due_total': s_due if s_due > 0 else Decimal('0.00'),
            'is_current': s.is_current
        })

    ctx.update({
        'branches': branches,
        'branch_metrics': branch_metrics,
        'session_metrics': session_metrics,
        'total_students_all': total_students_all,
        'total_collections_all': total_collections_all,
        'title': 'Consolidated Management Report'
    })
    return render(request, 'school_erp/consolidated_report.html', ctx)
